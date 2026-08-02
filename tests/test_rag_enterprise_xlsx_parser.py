from pathlib import Path

import openpyxl

from knowledge_pipeline.enterprise.orchestrator import parse_document
from knowledge_pipeline.enterprise.parsed_contracts import ParseStatus
from knowledge_pipeline.enterprise.parsers.base import ParserLimits


def test_xlsx_streams_every_sheet_without_500_row_truncation(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "rows.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "明细"
    for index in range(1, 503):
        sheet.append((index, f"row-{index}"))
    workbook.create_sheet("第二页").append(("状态", "有效"))
    workbook.save(path)
    called = {}
    original = openpyxl.load_workbook

    def tracked(*args, **kwargs):
        called.update(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", tracked)
    result = parse_document(path, source_id="src_xlsx", detected_format="xlsx")
    repeated = parse_document(path, source_id="src_xlsx", detected_format="xlsx")

    assert result.status is ParseStatus.READY
    assert called["read_only"] is True and called["data_only"] is True
    assert [table.sheet_name for table in result.tables] == ["明细", "第二页"]
    assert len(result.tables[0].rows) == 502
    assert result.tables[0].rows[-1] == ("502", "row-502")
    assert result.to_dict() == repeated.to_dict()


def test_xlsx_resource_limit_is_quarantined(tmp_path: Path) -> None:
    path = tmp_path / "limited.xlsx"
    workbook = openpyxl.Workbook()
    for index in range(3):
        workbook.active.append((index,))
    workbook.save(path)
    result = parse_document(path, source_id="src_limit", detected_format="xlsx", limits=ParserLimits(max_table_rows=2))
    assert result.status is ParseStatus.QUARANTINED
    assert result.isolation_reason == "resource_limit:table_rows"
