from pathlib import Path

from docx import Document

from knowledge_pipeline.enterprise.orchestrator import parse_document
from knowledge_pipeline.enterprise.parsed_contracts import ParseStatus


def test_docx_preserves_paragraph_table_order_and_is_deterministic(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("表格之前")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "名称", "数值"
    table.cell(1, 0).text, table.cell(1, 1).text = "电价", "100"
    document.add_paragraph("表格之后")
    document.save(path)

    first = parse_document(path, source_id="src_docx", detected_format="docx")
    second = parse_document(path, source_id="src_docx", detected_format="docx")

    assert first.status is ParseStatus.READY
    assert [block.kind for block in first.blocks] == ["heading", "paragraph", "table", "paragraph"]
    assert [block.text for block in first.blocks if block.text] == ["第一章", "表格之前", "表格之后"]
    assert first.tables[0].rows[1] == ("电价", "100")
    assert first.to_dict() == second.to_dict()


def test_corrupt_docx_is_quarantined(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not-an-ooxml-container")
    result = parse_document(path, source_id="src_broken", detected_format="docx")
    assert result.status is ParseStatus.QUARANTINED
    assert result.isolation_reason == "corrupt_docx"
