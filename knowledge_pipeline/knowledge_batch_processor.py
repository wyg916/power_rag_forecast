from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any


TEXT_ENCODINGS = ["utf-8", "utf-8-sig", "gbk", "gb2312", "gb18030"]
SUPPORTED_TEXT_TYPES = {"txt", "md", "html"}
SUPPORTED_DIRECT_TYPES = {"pdf", "docx", "xlsx", *SUPPORTED_TEXT_TYPES}
CONVERTIBLE_TYPES = {"doc", "xls", "wps", "et"}
UNSUPPORTED_TYPES = {"zip", "jpg", "jpeg", "png", "ofd", "rar", "7z"}
MAX_SHEET_ROWS = 500
MIN_CHUNK_LENGTH = 300
SOFT_MAX_CHUNK_LENGTH = 1500
HARD_MAX_CHUNK_LENGTH = 1800


CATEGORY_KEYWORDS: dict[str, list[str]] = {
    "电力市场规则": ["市场规则", "实施细则", "交易规则", "管理办法", "信息披露", "公平开放"],
    "电力现货交易": ["日前市场", "实时市场", "现货交易", "中长期交易", "出清", "零售市场"],
    "绿证交易": ["绿证", "绿色电力", "可再生能源证书", "GEC", "绿电"],
    "新能源政策": ["新能源", "风电", "光伏", "储能", "可再生能源", "消纳"],
    "电价政策": ["电价", "分时电价", "峰谷电价", "输配电价", "容量电价", "上网电价"],
    "电网代理购电": ["代理购电", "电网企业", "居民农业", "工商业用户"],
    "项目申报与建设": ["项目申报", "备案", "并网", "建设方案", "可研", "申报材料", "承诺书"],
    "交易策略": ["套利", "价差", "峰谷", "风险", "交易策略", "组合优化", "输电权"],
    "市场机制解释": ["节点电价", "边际价格", "拥堵", "损耗", "市场机制"],
    "PJM/DOM/LMP机制": ["PJM", "DOM", "Dominion", "LMP", "Day-Ahead", "Real-Time"],
    "模型与预测方法": ["预测", "模型", "误差", "RMSE", "MAE", "负荷", "天气"],
}

BUSINESS_KEYWORDS = [
    "电力市场",
    "现货交易",
    "日前市场",
    "实时市场",
    "中长期交易",
    "绿证",
    "绿色电力",
    "新能源",
    "光伏",
    "风电",
    "储能",
    "电价",
    "分时电价",
    "峰谷电价",
    "输配电价",
    "代理购电",
    "PJM",
    "DOM",
    "LMP",
    "节点电价",
    "边际价格",
    "拥堵",
    "损耗",
    "负荷",
    "天气",
    "尖峰",
    "风险",
    "预测",
    "RMSE",
    "MAE",
]


@dataclass
class FileResult:
    source_file: str
    source_path: str
    file_type: str
    status: str
    message: str = ""
    output_markdown: str = ""
    chunk_count: int = 0
    category: str = "其他"
    file_size: int = 0
    file_hash: str = ""
    text_hash: str = ""
    truncated_sheets: list[str] = field(default_factory=list)


@dataclass
class PipelineStats:
    total_files: int = 0
    success_files: int = 0
    failed_files: int = 0
    duplicate_files: int = 0
    duplicate_texts: int = 0
    total_chunks: int = 0
    scanned_pdf_count: int = 0
    format_counts: Counter[str] = field(default_factory=Counter)
    category_counts: Counter[str] = field(default_factory=Counter)
    longest_file: tuple[str, int] = ("", 0)
    max_chunk_file: tuple[str, int] = ("", 0)
    truncated_sheets: list[dict[str, str]] = field(default_factory=list)


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def safe_filename(name: str, max_length: int = 110) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip().strip(".")
    value = re.sub(r"\s+", " ", value)
    if not value:
        value = "untitled"
    reserved = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
    if value.upper() in reserved:
        value = f"{value}_file"
    return value[:max_length].rstrip(" .")


def unique_path(directory: Path, stem: str, suffix: str) -> Path:
    base = safe_filename(stem)
    candidate = directory / f"{base}{suffix}"
    index = 1
    while candidate.exists():
        candidate = directory / f"{base}_{index}{suffix}"
        index += 1
    return candidate


def read_text_with_fallback(path: Path) -> tuple[str, str]:
    data = path.read_bytes()
    for encoding in TEXT_ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    try:
        import chardet

        detected = chardet.detect(data)
        encoding = detected.get("encoding") or ""
        if encoding:
            return data.decode(encoding), encoding
    except Exception:
        pass
    raise UnicodeDecodeError("unknown", data[:1], 0, 1, "encoding detection failed")


def detect_file_type(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    name_lower = path.name.lower()
    if suffix:
        if suffix in {"pdf", "docx", "doc", "xlsx", "xls", "wps", "et", "html", "htm", "txt", "md", "zip", "jpg", "jpeg", "png", "ofd", "rar", "7z"}:
            return "html" if suffix == "htm" else suffix
        if suffix in {"cn%2f", "dot", "do"} or "treetempurl" in name_lower:
            return sniff_unknown_type(path)
    return sniff_unknown_type(path)


def sniff_unknown_type(path: Path) -> str:
    head = path.read_bytes()[:4096]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
                if any(name.startswith("word/") for name in names):
                    return "docx"
                if any(name.startswith("xl/") for name in names):
                    return "xlsx"
        except Exception:
            return "unknown"
    lower_head = head[:1024].lower()
    if b"<html" in lower_head or b"<!doctype html" in lower_head:
        return "html"
    for encoding in TEXT_ENCODINGS:
        try:
            text = head.decode(encoding)
            if re.search(r"[\u4e00-\u9fa5A-Za-z]{8,}", text):
                return "txt"
        except UnicodeDecodeError:
            continue
    return "unknown"


def find_soffice() -> Path | None:
    found = shutil.which("soffice")
    if found:
        return Path(found)
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    return next((path for path in candidates if path.exists()), None)


def convert_with_soffice(path: Path, outdir: Path, target_ext: str) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice soffice not found")
    outdir.mkdir(parents=True, exist_ok=True)
    command = [
        str(soffice),
        "--headless",
        "--convert-to",
        target_ext,
        "--outdir",
        str(outdir),
        str(path),
    ]
    proc = subprocess.run(command, capture_output=True, text=True, timeout=180)
    if proc.returncode != 0:
        raise RuntimeError((proc.stderr or proc.stdout or "conversion failed").strip()[:300])
    converted = outdir / f"{path.stem}.{target_ext}"
    if converted.exists():
        return converted
    matches = sorted(outdir.glob(f"{path.stem}*.{target_ext}"))
    if matches:
        return matches[0]
    raise RuntimeError("converted file not found")


def table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = []
    for row in rows:
        values = [(str(cell or "").replace("\n", " ").strip()) for cell in row]
        values.extend([""] * (width - len(values)))
        normalized.append(values)
    header = normalized[0]
    if not any(header):
        header = [f"列{i + 1}" for i in range(width)]
        body = normalized
    else:
        header = [value or f"列{i + 1}" for i, value in enumerate(header)]
        body = normalized[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def extract_pdf(path: Path) -> str:
    try:
        import fitz
    except Exception as exc:
        raise RuntimeError(f"PyMuPDF不可用：{exc}") from exc
    parts: list[str] = []
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                parts.append(f"\n\n## 第 {page_index} 页\n\n{text}")
    text = "\n".join(parts).strip()
    plain_length = len(re.sub(r"\s+", "", text))
    if plain_length < 80:
        raise ValueError("扫描版PDF或文本不可提取，需要OCR")
    return text


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError(f"python-docx不可用：{exc}") from exc
    document = Document(path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        if "heading" in style_name or "标题" in style_name:
            match = re.search(r"(\d+)", style_name)
            level = min(int(match.group(1)) if match else 2, 6)
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md_table = table_to_markdown(rows)
        if md_table:
            lines.append(f"\n\n### 表格 {table_index}\n\n{md_table}")
    return "\n\n".join(lines).strip()


def extract_xlsx(path: Path, max_sheet_rows: int = MAX_SHEET_ROWS) -> tuple[str, list[str]]:
    try:
        import openpyxl
    except Exception as exc:
        raise RuntimeError(f"openpyxl不可用：{exc}") from exc
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    truncated: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            if not any(value.strip() for value in values):
                continue
            if len(rows) >= max_sheet_rows:
                truncated.append(sheet.title)
                break
            rows.append(values)
        if not rows:
            continue
        parts.append(f"\n\n## Sheet：{sheet.title}\n\n{table_to_markdown(rows)}")
    return "\n".join(parts).strip(), truncated


def extract_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup
    except Exception as exc:
        raise RuntimeError(f"BeautifulSoup不可用：{exc}") from exc
    raw, _ = read_text_with_fallback(path)
    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else path.stem
    parts = [f"# {title}"]
    for table in soup.find_all("table"):
        rows = []
        for tr in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
            if cells:
                rows.append(cells)
        md_table = table_to_markdown(rows)
        if md_table:
            table.replace_with(soup.new_string("\n" + md_table + "\n"))
    body = soup.get_text("\n", strip=True)
    parts.append(body)
    return "\n\n".join(parts).strip()


def clean_text(text: str) -> str:
    value = (text or "").replace("\ufeff", "").replace("\u200b", "")
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n[ \t]+", "\n", value)
    lines = [line.rstrip() for line in value.splitlines()]
    cleaned_lines: list[str] = []
    previous = ""
    repeat_count = 0
    line_counter = Counter(line.strip() for line in lines if line.strip())
    noisy_lines = {
        line
        for line, count in line_counter.items()
        if count >= 4 and len(line) <= 40 and not re.search(r"\d|第|章|条|电价|市场|规则|办法|通知", line)
    }
    for line in lines:
        stripped = line.strip()
        if stripped in noisy_lines:
            continue
        if stripped == previous:
            repeat_count += 1
            if repeat_count >= 1:
                continue
        else:
            repeat_count = 0
        cleaned_lines.append(line)
        previous = stripped
    value = "\n".join(cleaned_lines)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def classify_document(filename: str, text: str) -> str:
    haystack = f"{filename}\n{text[:5000]}".lower()
    scores: dict[str, int] = {}
    for category, keywords in CATEGORY_KEYWORDS.items():
        score = 0
        for keyword in keywords:
            score += haystack.count(keyword.lower()) * (3 if keyword.lower() in filename.lower() else 1)
        scores[category] = score
    best_category, best_score = max(scores.items(), key=lambda item: item[1])
    return best_category if best_score > 0 else "其他"


def extract_keywords(text: str, filename: str, category: str) -> list[str]:
    haystack = f"{filename}\n{text[:10000]}"
    found: list[str] = []
    for keyword in BUSINESS_KEYWORDS:
        if keyword in haystack and keyword not in found:
            found.append(keyword)
    for keyword in CATEGORY_KEYWORDS.get(category, []):
        if keyword in haystack and keyword not in found:
            found.append(keyword)
    chinese_terms = re.findall(r"[\u4e00-\u9fa5]{2,12}", filename)
    for term in chinese_terms:
        if term not in found and len(term) >= 2:
            found.append(term)
    if category not in found and category != "其他":
        found.insert(0, category)
    fallback = ["电力资料", "知识库", category]
    for item in fallback:
        if item and item not in found:
            found.append(item)
    return found[:12]


def choose_cut_position(text: str, hard_max: int) -> int:
    if len(text) <= hard_max:
        return len(text)
    window = text[:hard_max]
    min_cut = min(max(MIN_CHUNK_LENGTH, int(hard_max * 0.45)), max(1, hard_max - 1))
    boundary_patterns = [
        r"\n#{1,6}\s+",
        r"\n第[一二三四五六七八九十百千万\d]+[章节条款项]",
        r"[。！？!?；;]\s*",
        r"[：:]\s*",
        r"\n+",
        r"[，,]\s*",
        r"\s+",
    ]
    for pattern in boundary_patterns:
        positions = [match.end() for match in re.finditer(pattern, window) if match.end() >= min_cut]
        if positions:
            return max(positions)
    return hard_max


def force_split_long_text(text: str, hard_max: int = HARD_MAX_CHUNK_LENGTH, overlap: int = 180) -> list[str]:
    value = (text or "").strip()
    if not value:
        return []
    if len(value) <= hard_max:
        return [value]
    safe_overlap = max(0, min(overlap, hard_max // 3))
    pieces: list[str] = []
    start = 0
    text_length = len(value)
    while start < text_length:
        remaining = value[start:]
        if len(remaining) <= hard_max:
            piece = remaining.strip()
            if piece:
                pieces.append(piece)
            break
        cut = choose_cut_position(remaining, hard_max)
        if cut <= 0:
            cut = hard_max
        end = min(start + cut, text_length)
        if end <= start:
            end = min(start + hard_max, text_length)
        piece = value[start:end].strip()
        if piece:
            pieces.append(piece)
        next_start = end - safe_overlap if safe_overlap else end
        if next_start <= start:
            next_start = end
        start = next_start
    final: list[str] = []
    for piece in pieces:
        if len(piece) <= hard_max:
            final.append(piece)
            continue
        for offset in range(0, len(piece), hard_max):
            sliced = piece[offset : offset + hard_max].strip()
            if sliced:
                final.append(sliced)
    return final


def looks_like_markdown_table(text: str) -> bool:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) < 3:
        return False
    table_lines = [line for line in lines if line.startswith("|") and line.endswith("|")]
    return len(table_lines) >= 3 and len(table_lines) / len(lines) >= 0.35


def split_table_markdown(table_text: str, hard_max: int = HARD_MAX_CHUNK_LENGTH) -> list[str]:
    text = (table_text or "").strip()
    if not text:
        return []
    if len(text) <= hard_max:
        return [text]
    lines = [line.rstrip() for line in text.splitlines()]
    prefix_lines: list[str] = []
    table_lines: list[str] = []
    suffix_lines: list[str] = []
    seen_table = False
    for line in lines:
        stripped = line.strip()
        is_table_line = stripped.startswith("|") and stripped.endswith("|")
        if is_table_line:
            seen_table = True
            table_lines.append(line)
        elif not seen_table:
            prefix_lines.append(line)
        else:
            suffix_lines.append(line)
    if len(table_lines) < 3:
        return force_split_long_text(text, hard_max=hard_max, overlap=180)

    header_lines = table_lines[:2]
    body_lines = table_lines[2:]
    groups: list[list[str]] = []
    current_rows: list[str] = []
    base_lines = [line for line in prefix_lines if line.strip()]

    for row in body_lines:
        candidate_lines = base_lines + ["表格分片"] + header_lines + current_rows + [row]
        candidate = "\n".join(candidate_lines).strip()
        if current_rows and len(candidate) > hard_max:
            groups.append(current_rows)
            current_rows = [row]
            continue
        current_rows.append(row)
    if current_rows:
        groups.append(current_rows)

    pieces: list[str] = []
    total = max(len(groups), 1)
    for index, rows in enumerate(groups, start=1):
        block_lines = base_lines + [f"表格分片 {index}/{total}"] + header_lines + rows
        if suffix_lines and index == total:
            block_lines.extend(suffix_lines)
        block = "\n".join(block_lines).strip()
        if len(block) <= hard_max:
            pieces.append(block)
        else:
            pieces.extend(force_split_long_text(block, hard_max=hard_max, overlap=0))
    return [piece for piece in pieces if piece.strip()]


def normalize_chunk_length(
    chunks: list[str],
    chunk_size: int,
    chunk_overlap: int,
    hard_max: int = HARD_MAX_CHUNK_LENGTH,
) -> list[str]:
    normalized: list[str] = []
    for chunk in chunks:
        value = chunk.strip()
        if not value:
            continue
        if len(value) <= hard_max:
            normalized.append(value)
            continue
        if looks_like_markdown_table(value):
            pieces = split_table_markdown(value, hard_max=hard_max)
        else:
            pieces = force_split_long_text(value, hard_max=hard_max, overlap=chunk_overlap)
        for piece in pieces:
            if len(piece) > hard_max:
                normalized.extend(force_split_long_text(piece, hard_max=hard_max, overlap=chunk_overlap))
            elif piece.strip():
                normalized.append(piece.strip())

    merged: list[str] = []
    merge_limit = min(SOFT_MAX_CHUNK_LENGTH, hard_max)
    for piece in normalized:
        if merged and len(piece) < MIN_CHUNK_LENGTH and len(merged[-1]) + len(piece) + 2 <= merge_limit:
            merged[-1] = f"{merged[-1]}\n\n{piece}".strip()
        else:
            merged.append(piece)

    final: list[str] = []
    for piece in merged:
        if len(piece) <= hard_max:
            final.append(piece)
        elif looks_like_markdown_table(piece):
            final.extend(split_table_markdown(piece, hard_max=hard_max))
        else:
            final.extend(force_split_long_text(piece, hard_max=hard_max, overlap=chunk_overlap))
    return [piece for piece in final if piece.strip()]


def split_long_unit(unit: str, chunk_size: int) -> list[str]:
    limit = min(max(chunk_size, MIN_CHUNK_LENGTH), HARD_MAX_CHUNK_LENGTH)
    return force_split_long_text(unit, hard_max=limit, overlap=0)


def split_chunks(text: str, chunk_size: int = 1000, chunk_overlap: int = 180) -> list[str]:
    normalized = clean_text(text)
    if not normalized:
        return []
    units = re.split(r"(?=\n#{1,6}\s+)|(?=\n第[一二三四五六七八九十百千万\d]+[章节条]\s*)|\n{2,}", normalized)
    expanded_units: list[str] = []
    for unit in units:
        unit = unit.strip()
        if not unit:
            continue
        expanded_units.extend(split_long_unit(unit, chunk_size))
    chunks: list[str] = []
    current = ""
    for unit in expanded_units:
        if not current:
            current = unit
            continue
        if len(current) + len(unit) + 2 <= chunk_size:
            current = f"{current}\n\n{unit}"
            continue
        if len(current) < MIN_CHUNK_LENGTH and chunks and len(chunks[-1]) + len(current) + 2 <= SOFT_MAX_CHUNK_LENGTH:
            chunks[-1] = f"{chunks[-1]}\n\n{current}".strip()
        else:
            chunks.append(current.strip())
        overlap = current[-chunk_overlap:].strip() if chunk_overlap > 0 and len(current) > chunk_overlap else ""
        current = f"{overlap}\n\n{unit}".strip() if overlap else unit
    if current:
        if len(current) < MIN_CHUNK_LENGTH and chunks and len(chunks[-1]) + len(current) + 2 <= SOFT_MAX_CHUNK_LENGTH:
            chunks[-1] = f"{chunks[-1]}\n\n{current}".strip()
        else:
            chunks.append(current.strip())
    if not chunks and normalized:
        chunks = [normalized]
    return normalize_chunk_length(chunks, chunk_size, chunk_overlap, HARD_MAX_CHUNK_LENGTH)


def build_markdown_header(
    *,
    title: str,
    source_path: Path,
    file_type: str,
    file_size: int,
    category: str,
    text_hash: str,
) -> str:
    size_kb = round(file_size / 1024, 2)
    return (
        f"# {title}\n\n"
        f"来源文件：{source_path.name}  \n"
        f"来源路径：{source_path}  \n"
        f"文件类型：{file_type}  \n"
        f"文件大小：{size_kb} KB  \n"
        f"处理时间：{now_text()}  \n"
        f"所属类别：{category}  \n"
        f"文本hash：{text_hash}  \n\n"
        "## 正文\n\n"
    )


def validate_chunks(jsonl_path: str | Path) -> dict[str, Any]:
    path = Path(jsonl_path)
    result: dict[str, Any] = {
        "available": path.exists(),
        "total_lines": 0,
        "valid_lines": 0,
        "errors": [],
        "duplicate_chunk_ids": [],
        "empty_text": 0,
        "missing_source_file": 0,
        "missing_source_path": 0,
        "source_path_not_exists": 0,
        "missing_category": 0,
        "missing_keywords": 0,
        "invalid_index": 0,
        "missing_text_hash": 0,
        "too_short_chunks": 0,
        "too_long_chunks": 0,
        "over_hard_limit_chunks": 0,
        "max_chunk_length": 0,
        "average_chunk_length": 0,
        "median_chunk_length": 0,
        "longest_chunks": [],
    }
    if not path.exists():
        result["errors"].append("chunks.jsonl 不存在")
        return result
    seen: set[str] = set()
    lengths: list[int] = []
    longest_chunks: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as fh:
        for line_no, line in enumerate(fh, start=1):
            stripped = line.strip()
            if not stripped:
                continue
            result["total_lines"] += 1
            try:
                item = json.loads(stripped)
            except json.JSONDecodeError as exc:
                result["errors"].append(f"第 {line_no} 行 JSON 解析失败：{exc}")
                continue
            result["valid_lines"] += 1
            chunk_id = str(item.get("chunk_id") or "")
            if chunk_id in seen:
                result["duplicate_chunk_ids"].append(chunk_id)
            elif chunk_id:
                seen.add(chunk_id)
            text = str(item.get("text") or "")
            text_length = len(text)
            lengths.append(text_length)
            longest_chunks.append(
                {
                    "line_no": line_no,
                    "chunk_id": chunk_id,
                    "source_file": item.get("source_file") or "",
                    "source_path": item.get("source_path") or "",
                    "category": item.get("category") or "",
                    "length": text_length,
                }
            )
            if not text.strip():
                result["empty_text"] += 1
            if not item.get("source_file"):
                result["missing_source_file"] += 1
            source_path = item.get("source_path")
            if not source_path:
                result["missing_source_path"] += 1
            elif not Path(str(source_path)).exists():
                result["source_path_not_exists"] += 1
            if not item.get("category"):
                result["missing_category"] += 1
            keywords = item.get("keywords")
            if not isinstance(keywords, list) or not keywords:
                result["missing_keywords"] += 1
            if not item.get("text_hash"):
                result["missing_text_hash"] += 1
            chunk_index = int(item.get("chunk_index") or 0)
            chunk_total = int(item.get("chunk_total") or 0)
            if chunk_index < 1 or chunk_total < chunk_index:
                result["invalid_index"] += 1
            if text_length < MIN_CHUNK_LENGTH:
                result["too_short_chunks"] += 1
            if text_length > SOFT_MAX_CHUNK_LENGTH:
                result["too_long_chunks"] += 1
            if text_length > HARD_MAX_CHUNK_LENGTH:
                result["over_hard_limit_chunks"] += 1
    if lengths:
        sorted_lengths = sorted(lengths)
        mid = len(sorted_lengths) // 2
        if len(sorted_lengths) % 2:
            median_length = sorted_lengths[mid]
        else:
            median_length = round((sorted_lengths[mid - 1] + sorted_lengths[mid]) / 2, 2)
        result["max_chunk_length"] = max(lengths)
        result["average_chunk_length"] = round(sum(lengths) / len(lengths), 2)
        result["median_chunk_length"] = median_length
    result["longest_chunks"] = sorted(longest_chunks, key=lambda item: item["length"], reverse=True)[:20]
    result["ok"] = (
        result["available"]
        and result["total_lines"] == result["valid_lines"]
        and not result["errors"]
        and not result["duplicate_chunk_ids"]
        and result["empty_text"] == 0
        and result["missing_source_file"] == 0
        and result["missing_source_path"] == 0
        and result["source_path_not_exists"] == 0
        and result["missing_category"] == 0
        and result["missing_keywords"] == 0
        and result["invalid_index"] == 0
        and result["missing_text_hash"] == 0
        and result["over_hard_limit_chunks"] == 0
    )
    return result


def create_sample_files(raw_dir: Path) -> None:
    raw_dir.mkdir(parents=True, exist_ok=True)
    samples = {
        "样例_电力市场规则.md": "# 电力市场规则样例\n\n电力市场规则通常包括日前市场、实时市场、中长期交易、信息披露和风险控制要求。该样例用于在真实资料目录不存在时验证批量整理工具。",
        "样例_电价政策.txt": "分时电价、峰谷电价、输配电价和代理购电政策会影响售电公司交易策略。整理知识库时应保留政策名称、发布日期、适用范围、价格单位和执行条件。",
        "样例_PJM_LMP机制.html": "<html><head><title>PJM LMP机制样例</title></head><body><h1>PJM LMP</h1><p>LMP 包含能源价格、拥堵价格和损耗价格，Day-Ahead 与 Real-Time 市场价格可能存在偏差。</p></body></html>",
    }
    for name, content in samples.items():
        path = raw_dir / name
        if not path.exists():
            path.write_text(content, encoding="utf-8")


class KnowledgeBatchProcessor:
    def __init__(
        self,
        input_dir: Path,
        output_dir: Path,
        chunk_size: int = 1000,
        chunk_overlap: int = 180,
        dry_run: bool = False,
        resume: bool = False,
        rebuild_chunks: bool = False,
        max_files: int = 0,
        verbose: bool = False,
    ) -> None:
        self.input_dir = input_dir
        self.output_dir = output_dir
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.dry_run = dry_run
        self.resume = resume
        self.rebuild_chunks = rebuild_chunks
        self.max_files = max_files
        self.verbose = verbose
        self.raw_dir = output_dir / "raw"
        self.converted_docx_dir = output_dir / "converted" / "docx"
        self.converted_xlsx_dir = output_dir / "converted" / "xlsx"
        self.converted_pdf_dir = output_dir / "converted" / "pdf"
        self.markdown_dir = output_dir / "markdown"
        self.chunks_dir = output_dir / "chunks"
        self.output_files_dir = output_dir / "output"
        self.chunks_jsonl = self.output_files_dir / "chunks.jsonl"
        self.failed_csv = self.output_files_dir / "failed_files.csv"
        self.log_path = self.output_files_dir / "processing_log.txt"
        self.report_path = self.output_files_dir / "processing_report.md"
        self.validation_path = self.output_files_dir / "validation_report.json"
        self.manifest_path = self.output_files_dir / "processed_manifest.json"
        self.stats = PipelineStats()
        self.failures: list[dict[str, str]] = []
        self.results: list[FileResult] = []
        self.file_hash_seen: dict[tuple[int, str], Path] = {}
        self.text_hash_seen: dict[str, Path] = {}
        self.processed_hashes: set[str] = set()

    def ensure_dirs(self) -> None:
        for directory in [
            self.raw_dir,
            self.converted_docx_dir,
            self.converted_xlsx_dir,
            self.converted_pdf_dir,
            self.markdown_dir,
            self.chunks_dir,
            self.output_files_dir,
        ]:
            directory.mkdir(parents=True, exist_ok=True)

    def load_manifest(self) -> None:
        if self.rebuild_chunks:
            self.processed_hashes = set()
            return
        if not self.resume or not self.manifest_path.exists():
            return
        try:
            data = json.loads(self.manifest_path.read_text(encoding="utf-8"))
            self.processed_hashes = set(data.get("processed_hashes") or [])
        except Exception:
            self.processed_hashes = set()

    def save_manifest(self) -> None:
        data = {"processed_hashes": sorted(self.processed_hashes), "updated_at": now_text()}
        self.manifest_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    def log(self, status: str, file_type: str, path: Path, output_file: str = "", chunk_count: int = 0, note: str = "") -> None:
        line = f"{now_text()} | {status} | {file_type} | {path} | {output_file} | {chunk_count} | {note}\n"
        with self.log_path.open("a", encoding="utf-8") as fh:
            fh.write(line)
        if self.verbose:
            try:
                print(line.rstrip())
            except UnicodeEncodeError:
                encoding = sys.stdout.encoding or "utf-8"
                safe_line = line.rstrip().encode(encoding, errors="replace").decode(encoding, errors="replace")
                print(safe_line)

    def add_failure(self, path: Path, file_type: str, reason: str, action: str) -> None:
        self.failures.append(
            {
                "source_file": path.name,
                "source_path": str(path),
                "file_type": file_type,
                "failed_reason": reason,
                "suggested_action": action,
            }
        )
        if "扫描版PDF" in reason:
            self.stats.scanned_pdf_count += 1
        self.log("提取失败", file_type, path, "", 0, reason)

    def scan_files(self) -> list[Path]:
        if not self.input_dir.exists():
            return []
        files = [path for path in self.input_dir.rglob("*") if path.is_file()]
        files = sorted(files, key=lambda item: str(item).lower())
        if self.max_files:
            files = files[: self.max_files]
        return files

    def next_chunk_id_from_existing(self) -> int:
        if self.rebuild_chunks or not self.resume or not self.chunks_jsonl.exists():
            return 1
        max_id = 0
        try:
            with self.chunks_jsonl.open("r", encoding="utf-8") as fh:
                for line in fh:
                    if not line.strip():
                        continue
                    item = json.loads(line)
                    chunk_id = str(item.get("chunk_id") or "")
                    match = re.search(r"(\d+)$", chunk_id)
                    if match:
                        max_id = max(max_id, int(match.group(1)))
        except Exception:
            return 1
        return max_id + 1

    def dry_run_report(self, files: list[Path]) -> None:
        counter = Counter(detect_file_type(path) for path in files)
        self.output_files_dir.mkdir(parents=True, exist_ok=True)
        lines = [
            "# 知识库批量整理 Dry Run 报告",
            "",
            f"- 输入目录：`{self.input_dir}`",
            f"- 扫描文件数量：{len(files)}",
            "",
            "## 格式分布",
            "",
            "| 文件类型 | 数量 |",
            "|---|---:|",
        ]
        for file_type, count in sorted(counter.items()):
            lines.append(f"| {file_type} | {count} |")
        self.report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        print(f"Dry run scanned {len(files)} files.")
        print(dict(counter))

    def extract_content(self, path: Path, file_type: str) -> tuple[str, list[str], str]:
        if file_type == "pdf":
            return extract_pdf(path), [], file_type
        if file_type == "docx":
            return extract_docx(path), [], file_type
        if file_type == "xlsx":
            text, truncated = extract_xlsx(path)
            return text, truncated, file_type
        if file_type == "html":
            return extract_html(path), [], file_type
        if file_type in {"txt", "md"}:
            text, _ = read_text_with_fallback(path)
            return text, [], file_type
        if file_type == "doc":
            converted = convert_with_soffice(path, self.converted_docx_dir, "docx")
            return extract_docx(converted), [], "docx"
        if file_type == "xls":
            converted = convert_with_soffice(path, self.converted_xlsx_dir, "xlsx")
            text, truncated = extract_xlsx(converted)
            return text, truncated, "xlsx"
        if file_type == "wps":
            converted = convert_with_soffice(path, self.converted_docx_dir, "docx")
            return extract_docx(converted), [], "docx"
        if file_type == "et":
            converted = convert_with_soffice(path, self.converted_xlsx_dir, "xlsx")
            text, truncated = extract_xlsx(converted)
            return text, truncated, "xlsx"
        raise ValueError("异常后缀或未知格式，需人工判断")

    def write_failure_csv(self) -> None:
        with self.failed_csv.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=["source_file", "source_path", "file_type", "failed_reason", "suggested_action"])
            writer.writeheader()
            writer.writerows(self.failures)

    def write_chunks(self, chunks: list[dict[str, Any]]) -> None:
        mode = "a" if self.resume and not self.rebuild_chunks and self.chunks_jsonl.exists() else "w"
        with self.chunks_jsonl.open(mode, encoding="utf-8") as fh:
            for item in chunks:
                fh.write(json.dumps(item, ensure_ascii=False) + "\n")

    def process_file(self, path: Path, next_chunk_id: int) -> tuple[FileResult, list[dict[str, Any]], int]:
        stat = path.stat()
        file_type = detect_file_type(path)
        self.stats.format_counts[file_type] += 1
        f_hash = file_hash(path)
        file_key = (stat.st_size, f_hash)
        if f_hash in self.processed_hashes:
            self.stats.duplicate_files += 1
            message = "resume 已处理" if self.resume and not self.rebuild_chunks else "已在本次运行处理"
            result = FileResult(path.name, str(path), file_type, "跳过重复文件", message, file_size=stat.st_size, file_hash=f_hash)
            self.log("跳过重复文件", file_type, path, "", 0, message)
            return result, [], next_chunk_id
        if file_key in self.file_hash_seen:
            self.stats.duplicate_files += 1
            result = FileResult(path.name, str(path), file_type, "跳过重复文件", f"与 {self.file_hash_seen[file_key].name} 完全一致", file_size=stat.st_size, file_hash=f_hash)
            self.log("跳过重复文件", file_type, path, "", 0, result.message)
            return result, [], next_chunk_id
        self.file_hash_seen[file_key] = path
        if file_type in UNSUPPORTED_TYPES:
            reason = "格式不支持，需人工判断"
            self.add_failure(path, file_type, reason, "请手动转换为 PDF、DOCX、XLSX、HTML、TXT 或 MD")
            return FileResult(path.name, str(path), file_type, "格式不支持", reason, file_size=stat.st_size, file_hash=f_hash), [], next_chunk_id
        try:
            text, truncated_sheets, actual_type = self.extract_content(path, file_type)
        except ValueError as exc:
            reason = str(exc)
            if file_type == "pdf" and "扫描版PDF" in reason:
                action = "后续使用 OCR 后再导入"
            elif file_type == "doc":
                reason, action = "DOC转换DOCX失败，需手动另存为DOCX", "手动另存为 DOCX"
            elif file_type == "xls":
                reason, action = "XLS转换XLSX失败，需手动另存为XLSX", "手动另存为 XLSX"
            elif file_type == "wps":
                reason, action = "WPS文字文档需手动另存为DOCX", "手动另存为 DOCX"
            elif file_type == "et":
                reason, action = "WPS表格需手动另存为XLSX", "手动另存为 XLSX"
            else:
                reason, action = "异常后缀或未知格式，需人工判断", "人工判断文件真实格式"
            self.add_failure(path, file_type, reason, action)
            return FileResult(path.name, str(path), file_type, "提取失败", reason, file_size=stat.st_size, file_hash=f_hash), [], next_chunk_id
        except UnicodeDecodeError:
            reason = "编码识别失败，需手动转换为UTF-8文本"
            self.add_failure(path, file_type, reason, "手动转换为 UTF-8 文本")
            return FileResult(path.name, str(path), file_type, "提取失败", reason, file_size=stat.st_size, file_hash=f_hash), [], next_chunk_id
        except Exception as exc:
            if file_type == "doc":
                reason, action = "DOC转换DOCX失败，需手动另存为DOCX", "手动另存为 DOCX"
            elif file_type == "xls":
                reason, action = "XLS转换XLSX失败，需手动另存为XLSX", "手动另存为 XLSX"
            elif file_type == "wps":
                reason, action = "WPS文字文档需手动另存为DOCX", "手动另存为 DOCX"
            elif file_type == "et":
                reason, action = "WPS表格需手动另存为XLSX", "手动另存为 XLSX"
            else:
                reason, action = str(exc)[:300], "检查依赖或人工转换格式"
            self.add_failure(path, file_type, reason, action)
            return FileResult(path.name, str(path), file_type, "提取失败", reason, file_size=stat.st_size, file_hash=f_hash), [], next_chunk_id
        cleaned = clean_text(text)
        if not cleaned:
            reason = "内容为空"
            self.add_failure(path, file_type, reason, "检查文件内容或手动转换")
            return FileResult(path.name, str(path), file_type, "内容为空", reason, file_size=stat.st_size, file_hash=f_hash), [], next_chunk_id
        t_hash = sha256_text(cleaned)
        if t_hash in self.text_hash_seen:
            self.stats.duplicate_texts += 1
            result = FileResult(path.name, str(path), file_type, "跳过重复文件", f"文本与 {self.text_hash_seen[t_hash].name} 完全一致", file_size=stat.st_size, file_hash=f_hash, text_hash=t_hash)
            self.log("跳过重复文件", file_type, path, "", 0, result.message)
            return result, [], next_chunk_id
        self.text_hash_seen[t_hash] = path
        category = classify_document(path.name, cleaned)
        title = path.stem.strip() or path.name
        md_body = build_markdown_header(
            title=title,
            source_path=path,
            file_type=file_type,
            file_size=stat.st_size,
            category=category,
            text_hash=t_hash,
        ) + cleaned
        md_path = unique_path(self.markdown_dir, path.stem, ".md")
        md_path.write_text(md_body, encoding="utf-8")
        text_chunks = split_chunks(cleaned, chunk_size=self.chunk_size, chunk_overlap=self.chunk_overlap)
        text_chunks = normalize_chunk_length(text_chunks, self.chunk_size, self.chunk_overlap, HARD_MAX_CHUNK_LENGTH)
        keywords = extract_keywords(cleaned, path.name, category)
        chunk_records: list[dict[str, Any]] = []
        chunk_total = len(text_chunks)
        for index, chunk in enumerate(text_chunks, start=1):
            chunk_records.append(
                {
                    "chunk_id": f"knowledge_{next_chunk_id:06d}",
                    "source_file": path.name,
                    "source_path": str(path),
                    "file_type": file_type,
                    "category": category,
                    "title": title,
                    "chunk_index": index,
                    "chunk_total": chunk_total,
                    "text": chunk,
                    "keywords": keywords,
                    "text_hash": sha256_text(chunk),
                    "created_at": now_text(),
                }
            )
            next_chunk_id += 1
        chunk_file = unique_path(self.chunks_dir, path.stem, ".jsonl")
        with chunk_file.open("w", encoding="utf-8") as fh:
            for item in chunk_records:
                fh.write(json.dumps(item, ensure_ascii=False) + "\n")
        self.processed_hashes.add(f_hash)
        if truncated_sheets:
            for sheet in truncated_sheets:
                self.stats.truncated_sheets.append({"source_file": path.name, "sheet": sheet})
        result = FileResult(
            source_file=path.name,
            source_path=str(path),
            file_type=file_type,
            status="成功",
            output_markdown=str(md_path),
            chunk_count=len(chunk_records),
            category=category,
            file_size=stat.st_size,
            file_hash=f_hash,
            text_hash=t_hash,
            truncated_sheets=truncated_sheets,
        )
        self.log("成功", file_type, path, str(md_path), len(chunk_records), f"category={category}")
        return result, chunk_records, next_chunk_id

    def write_report(self, validation: dict[str, Any]) -> None:
        avg_chunks = round(self.stats.total_chunks / self.stats.success_files, 2) if self.stats.success_files else 0
        lines = [
            "# 知识库批量整理处理报告",
            "",
            f"- 处理时间：{now_text()}",
            f"- 输入目录：`{self.input_dir}`",
            f"- 输出目录：`{self.output_dir}`",
            f"- 总文件数量：{self.stats.total_files}",
            f"- 成功处理文件数量：{self.stats.success_files}",
            f"- 失败文件数量：{self.stats.failed_files}",
            f"- 跳过重复文件数量：{self.stats.duplicate_files}",
            f"- 文本级重复数量：{self.stats.duplicate_texts}",
            f"- 总 chunk 数：{self.stats.total_chunks}",
            f"- 平均每个文件 chunk 数：{avg_chunks}",
            f"- 最长文件：{self.stats.longest_file[0]}（{self.stats.longest_file[1]} 字符）",
            f"- 最多 chunk 文件：{self.stats.max_chunk_file[0]}（{self.stats.max_chunk_file[1]} chunks）",
            f"- 扫描版 PDF 数量：{self.stats.scanned_pdf_count}",
            "",
            "## 各文件格式数量",
            "",
            "| 文件类型 | 数量 |",
            "|---|---:|",
        ]
        for file_type, count in sorted(self.stats.format_counts.items()):
            lines.append(f"| {file_type} | {count} |")
        lines.extend(["", "## 各类别文件数量", "", "| 类别 | 数量 |", "|---|---:|"])
        for category, count in sorted(self.stats.category_counts.items()):
            lines.append(f"| {category} | {count} |")
        lines.extend(
            [
                "",
                "## JSONL 验证结果",
                "",
                f"- 校验是否通过：{'是' if validation.get('ok') else '否'}",
                f"- 总行数：{validation.get('total_lines', 0)}",
                f"- 合法 JSON 行数：{validation.get('valid_lines', 0)}",
                f"- 最大 chunk 长度：{validation.get('max_chunk_length', 0)}",
                f"- 平均 chunk 长度：{validation.get('average_chunk_length', 0)}",
                f"- 中位数 chunk 长度：{validation.get('median_chunk_length', 0)}",
                f"- 超过软限制 {SOFT_MAX_CHUNK_LENGTH} 的 chunk 数：{validation.get('too_long_chunks', 0)}",
                f"- 超过硬限制 {HARD_MAX_CHUNK_LENGTH} 的 chunk 数：{validation.get('over_hard_limit_chunks', 0)}",
                f"- 过短 chunk 数：{validation.get('too_short_chunks', 0)}",
                f"- 重复 chunk_id 数：{len(validation.get('duplicate_chunk_ids') or [])}",
                f"- 缺失 source_file 数：{validation.get('missing_source_file', 0)}",
                f"- 缺失 source_path 数：{validation.get('missing_source_path', 0)}",
                f"- source_path 不存在数：{validation.get('source_path_not_exists', 0)}",
                f"- 缺失 category 数：{validation.get('missing_category', 0)}",
                f"- 缺失 keywords 数：{validation.get('missing_keywords', 0)}",
                f"- chunk_index/chunk_total 异常数：{validation.get('invalid_index', 0)}",
                f"- 缺失 text_hash 数：{validation.get('missing_text_hash', 0)}",
            ]
        )
        longest_chunks = validation.get("longest_chunks") or []
        lines.extend(["", "### 最长 Top 20 chunk", "", "| 行号 | chunk_id | source_file | 类别 | 长度 |", "|---:|---|---|---|---:|"])
        if longest_chunks:
            for item in longest_chunks:
                lines.append(
                    f"| {item.get('line_no', '')} | {item.get('chunk_id', '')} | "
                    f"{item.get('source_file', '')} | {item.get('category', '')} | {item.get('length', 0)} |"
                )
        else:
            lines.append("| - | - | - | - | - |")
        if validation.get("errors"):
            lines.extend(["", "### 验证错误", ""])
            lines.extend(f"- {item}" for item in validation["errors"][:50])
        lines.extend(["", "## 被截断的 Sheet", ""])
        if self.stats.truncated_sheets:
            lines.append("| 文件 | Sheet |")
            lines.append("|---|---|")
            for item in self.stats.truncated_sheets:
                lines.append(f"| {item['source_file']} | {item['sheet']} |")
        else:
            lines.append("无。")
        lines.extend(["", "## 需要人工处理的文件列表", ""])
        if self.failures:
            lines.append("| 文件 | 类型 | 失败原因 | 建议动作 |")
            lines.append("|---|---|---|---|")
            for item in self.failures[:80]:
                lines.append(f"| {item['source_file']} | {item['file_type']} | {item['failed_reason']} | {item['suggested_action']} |")
        else:
            lines.append("无。")
        lines.extend(
            [
                "",
                "## 下一步导入 RAG 建议",
                "",
                "1. 优先人工处理 `failed_files.csv` 中的扫描版 PDF、WPS/ET、异常后缀和图片/OFD 文件。",
                "2. 将 `markdown/` 目录作为可追溯文本库纳入现有 `index_local_knowledge` 的扫描根目录，或使用 `output/chunks.jsonl` 批量写入数据库。",
                "3. 导入向量库前先运行 `validate_chunks()`，确认 chunk_id 唯一、source_path 可追溯、text 非空。",
                "4. 对大表格被截断的文件，可提高 `MAX_SHEET_ROWS` 或拆分为专用结构化表后再入库。",
            ]
        )
        self.report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    def run(self) -> dict[str, Any]:
        self.ensure_dirs()
        self.load_manifest()
        if self.rebuild_chunks or not self.resume:
            self.log_path.write_text("", encoding="utf-8")
            self.failed_csv.write_text("", encoding="utf-8")
            self.chunks_jsonl.write_text("", encoding="utf-8")
            self.validation_path.write_text("", encoding="utf-8")
            self.report_path.write_text("", encoding="utf-8")
            self.manifest_path.write_text(json.dumps({"processed_hashes": [], "updated_at": now_text()}, ensure_ascii=False, indent=2), encoding="utf-8")
        files = self.scan_files()
        self.stats.total_files = len(files)
        if self.dry_run:
            self.dry_run_report(files)
            return {"dry_run": True, "total_files": len(files)}
        next_chunk_id = self.next_chunk_id_from_existing()
        all_chunks: list[dict[str, Any]] = []
        for path in files:
            result, chunks, next_chunk_id = self.process_file(path, next_chunk_id)
            self.results.append(result)
            if result.status == "成功":
                self.stats.success_files += 1
                self.stats.total_chunks += result.chunk_count
                self.stats.category_counts[result.category] += 1
                text_len = 0
                try:
                    text_len = len(Path(result.output_markdown).read_text(encoding="utf-8"))
                except Exception:
                    pass
                if text_len > self.stats.longest_file[1]:
                    self.stats.longest_file = (result.source_file, text_len)
                if result.chunk_count > self.stats.max_chunk_file[1]:
                    self.stats.max_chunk_file = (result.source_file, result.chunk_count)
            elif result.status in {"提取失败", "格式不支持", "内容为空"}:
                self.stats.failed_files += 1
            all_chunks.extend(chunks)
        self.write_chunks(all_chunks)
        self.write_failure_csv()
        validation = validate_chunks(self.chunks_jsonl)
        self.validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding="utf-8")
        self.write_report(validation)
        self.save_manifest()
        return {
            "dry_run": False,
            "total_files": self.stats.total_files,
            "success_files": self.stats.success_files,
            "failed_files": self.stats.failed_files,
            "duplicate_files": self.stats.duplicate_files,
            "duplicate_texts": self.stats.duplicate_texts,
            "total_chunks": self.stats.total_chunks,
            "validation": validation,
        }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="批量整理多格式电力资料为 Markdown 和 RAG chunks.jsonl。")
    default_input = Path(__file__).resolve().parent / "raw"
    default_output = Path(__file__).resolve().parent
    parser.add_argument("--input", type=Path, default=default_input, help="原始资料目录")
    parser.add_argument("--output", type=Path, default=default_output, help="输出目录")
    parser.add_argument("--chunk-size", type=int, default=1000)
    parser.add_argument("--chunk-overlap", type=int, default=180)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--rebuild-chunks", action="store_true", help="忽略旧 manifest 并重新生成 chunks.jsonl、报告和 chunk 编号")
    parser.add_argument("--max-files", type=int, default=0)
    parser.add_argument("--verbose", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    input_dir = args.input
    output_dir = args.output
    if not input_dir.exists():
        print(f"提示：原始资料目录不存在：{input_dir}")
        fallback = Path(__file__).resolve().parent / "raw"
        fallback.mkdir(parents=True, exist_ok=True)
        create_sample_files(fallback)
        input_dir = fallback
        print(f"已使用默认 input：{input_dir}")
    processor = KnowledgeBatchProcessor(
        input_dir=input_dir,
        output_dir=output_dir,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        dry_run=args.dry_run,
        resume=args.resume,
        rebuild_chunks=args.rebuild_chunks,
        max_files=args.max_files,
        verbose=args.verbose,
    )
    result = processor.run()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
