from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import os
import re
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from backend.app.ai.identity_context import IdentityContext


ROOT = Path(__file__).resolve().parents[2] / "data" / "assistant_uploads"
MAX_BYTES = int(20 * 1024 * 1024)
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
SESSION_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:-]{0,127}$")
ATTACHMENT_ID_RE = re.compile(r"^att_[0-9a-f]{32}$")
PROMPT_INJECTION = re.compile(
    r"ignore\s+(all\s+)?previous|system\s*prompt|忽略.{0,8}(指令|规则)|绕过.{0,8}(权限|安全)", re.I
)

MEDIA_TYPES = {
    ".png": {"image/png"},
    ".jpg": {"image/jpeg", "image/jpg"},
    ".jpeg": {"image/jpeg", "image/jpg"},
    ".webp": {"image/webp"},
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".csv": {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"},
}
TERMINAL_STATUSES = {"failed", "cancelled", "deleted"}


class AttachmentError(RuntimeError):
    def __init__(self, code: str, message: str, *, status_code: int, attachment_id: str = "") -> None:
        self.code, self.message, self.status_code, self.attachment_id = code, message, status_code, attachment_id
        super().__init__(message)


def _safe_name(value: str) -> str:
    name = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", Path(value or "attachment").name, flags=re.UNICODE).strip("._")
    return (name or "attachment")[:180]


def _scope(identity: IdentityContext) -> Path:
    identity.require_valid(require_session=False)
    digest = hashlib.sha256(f"{identity.tenant_id}\0{identity.user_id}".encode()).hexdigest()[:32]
    return ROOT / digest


def _paths(identity: IdentityContext, attachment_id: str) -> tuple[Path, Path]:
    if not ATTACHMENT_ID_RE.fullmatch(attachment_id or ""):
        raise AttachmentError("ATTACHMENT_NOT_FOUND", "附件不存在。", status_code=404)
    base = _scope(identity) / attachment_id
    return base.with_suffix(".bin"), base.with_suffix(".json")


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _public(meta: dict[str, Any]) -> dict[str, Any]:
    return {key: meta.get(key) for key in (
        "attachment_id", "session_id", "file_name", "media_type", "size_bytes", "sha256",
        "status", "parser", "created_at", "expires_at", "error", "prompt_injection_detected",
    )}


def _discard_attachment_payload(
    meta: dict[str, Any],
    data_path: Path,
    *,
    status: str,
    error: dict[str, str] | None = None,
) -> None:
    """Remove the temporary object and every parsed-content derivative."""
    if data_path.is_file():
        data_path.unlink()
    meta.update(
        status=status,
        chunks=[],
        properties={},
        prompt_injection_detected=False,
        error=error,
        deleted_at=datetime.now(timezone.utc).isoformat(),
    )


def _check_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = [item.filename.lower() for item in archive.infolist()]
            total = sum(item.file_size for item in archive.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise AttachmentError("ATTACHMENT_PARSE_FAILED", "压缩文档结构无效。", status_code=424) from exc
    forbidden = ("vbaproject.bin", "embeddings/", ".exe", ".js", ".vbs", ".ps1", ".cmd", ".bat")
    if total > MAX_UNCOMPRESSED_BYTES or any(any(token in name for token in forbidden) for name in names):
        raise AttachmentError("ATTACHMENT_UNSAFE_CONTENT", "附件包含宏、嵌入文件或异常压缩内容。", status_code=422)


def _text_chunks(ext: str, content: bytes) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    if b"\x00" in content:
        raise AttachmentError("ATTACHMENT_PARSE_FAILED", "文本附件包含二进制内容。", status_code=424)
    text = content.decode("utf-8-sig", errors="strict")
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(text)))
        if len(rows) > 10_000:
            raise AttachmentError("ATTACHMENT_LIMIT_EXCEEDED", "CSV 行数超过限制。", status_code=422)
        safe_rows = [[f"[formula-text]{cell}" if cell.lstrip().startswith(("=", "+", "-", "@")) else cell for cell in row] for row in rows]
        text = "\n".join(",".join(row) for row in safe_rows)
        location = {"row_range": f"1-{len(rows)}"}
        parser = "csv"
    else:
        location, parser = {"section": "document"}, ext.removeprefix(".")
    text = text[:MAX_TEXT_CHARS]
    return parser, [{"chunk_id": f"achunk_{uuid4().hex}", "text": text, "location": location}], {}


def _parse(ext: str, content: bytes) -> tuple[str, list[dict[str, Any]], dict[str, Any]]:
    if ext in {".txt", ".md", ".csv"}:
        return _text_chunks(ext, content)
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image
        try:
            with Image.open(io.BytesIO(content)) as image:
                image.verify()
            with Image.open(io.BytesIO(content)) as image:
                width, height, detected = image.width, image.height, str(image.format or "").lower()
        except Exception as exc:
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "图片内容无效。", status_code=424) from exc
        expected = "jpeg" if ext in {".jpg", ".jpeg"} else ext.removeprefix(".")
        if detected != expected or width * height > 40_000_000:
            raise AttachmentError("ATTACHMENT_UNSAFE_CONTENT", "图片格式或像素数量不符合限制。", status_code=422)
        return "pillow", [], {"width": width, "height": height, "vision": True}
    if ext == ".pdf":
        if not content.startswith(b"%PDF") or any(token in content for token in (b"/JavaScript", b"/Launch", b"/EmbeddedFile")):
            raise AttachmentError("ATTACHMENT_UNSAFE_CONTENT", "PDF 包含不支持的活动或嵌入内容。", status_code=422)
        import fitz
        try:
            document = fitz.open(stream=content, filetype="pdf")
            if document.page_count > 200 or document.embfile_count() > 0:
                raise AttachmentError("ATTACHMENT_LIMIT_EXCEEDED", "PDF 页数或嵌入文件超过限制。", status_code=422)
            chunks = [{"chunk_id": f"achunk_{uuid4().hex}", "text": page.get_text("text")[:8000], "location": {"page": index + 1}} for index, page in enumerate(document)]
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "PDF 解析失败。", status_code=424) from exc
        if not any(chunk["text"].strip() for chunk in chunks):
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "PDF 没有可解析文本。", status_code=424)
        return "pymupdf", chunks, {"pages": len(chunks)}
    _check_zip(content)
    if ext == ".docx":
        from docx import Document
        try:
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())[:MAX_TEXT_CHARS]
        except Exception as exc:
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "DOCX 解析失败。", status_code=424) from exc
        if not text:
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "DOCX 没有可解析文本。", status_code=424)
        return "python-docx", [{"chunk_id": f"achunk_{uuid4().hex}", "text": text, "location": {"section": "document"}}], {}
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_vba=False)
        chunks: list[dict[str, Any]] = []
        for sheet in workbook.worksheets[:50]:
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_index > 10_000:
                    raise AttachmentError("ATTACHMENT_LIMIT_EXCEEDED", "XLSX 行数超过限制。", status_code=422)
                values = [f"[formula-text]{value}" if isinstance(value, str) and value.startswith("=") else str(value or "") for value in row]
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                chunks.append({"chunk_id": f"achunk_{uuid4().hex}", "text": "\n".join(rows)[:MAX_TEXT_CHARS], "location": {"sheet": sheet.title, "row_range": f"1-{len(rows)}"}})
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError("ATTACHMENT_PARSE_FAILED", "XLSX 解析失败。", status_code=424) from exc
    if not chunks:
        raise AttachmentError("ATTACHMENT_PARSE_FAILED", "XLSX 没有可解析单元格。", status_code=424)
    return "openpyxl", chunks, {"sheets": len(chunks)}


def upload_attachment(identity: IdentityContext, *, session_id: str, file_name: str, media_type: str, content: bytes) -> dict[str, Any]:
    if not SESSION_RE.fullmatch(session_id or ""):
        raise AttachmentError("VALIDATION_FAILED", "session_id 无效。", status_code=422)
    safe_name, declared = _safe_name(file_name), (media_type or "").split(";", 1)[0].strip().lower()
    ext = Path(safe_name).suffix.lower()
    if ext not in MEDIA_TYPES or declared not in MEDIA_TYPES[ext]:
        raise AttachmentError("ATTACHMENT_TYPE_UNSUPPORTED", "附件类型不在白名单中。", status_code=415)
    if not content or len(content) > MAX_BYTES:
        raise AttachmentError("ATTACHMENT_TOO_LARGE", "附件为空或超过大小限制。", status_code=413)
    attachment_id = f"att_{uuid4().hex}"
    data_path, meta_path = _paths(identity, attachment_id)
    now = datetime.now(timezone.utc)
    meta = {
        "attachment_id": attachment_id, "session_id": session_id, "file_name": safe_name,
        "media_type": declared, "size_bytes": len(content), "sha256": hashlib.sha256(content).hexdigest(),
        "status": "uploading", "parser": None, "created_at": now.isoformat(),
        "expires_at": (now + timedelta(hours=int(os.environ.get("AI_ATTACHMENT_TTL_HOURS", "24")))).isoformat(),
        "tenant_id": identity.tenant_id, "user_id": identity.user_id, "error": None, "chunks": [], "properties": {},
    }
    data_path.parent.mkdir(parents=True, exist_ok=True)
    data_path.write_bytes(content)
    _write_json(meta_path, meta)
    meta["status"] = "parsing"
    _write_json(meta_path, meta)
    try:
        parser, chunks, properties = _parse(ext, content)
        combined = "\n".join(str(chunk.get("text") or "") for chunk in chunks)
        meta.update(parser=parser, chunks=chunks, properties=properties, status="ready", prompt_injection_detected=bool(PROMPT_INJECTION.search(combined)))
        _write_json(meta_path, meta)
        return _public(meta)
    except AttachmentError as exc:
        _discard_attachment_payload(
            meta,
            data_path,
            status="failed",
            error={"code": exc.code, "message": exc.message},
        )
        _write_json(meta_path, meta)
        exc.attachment_id = attachment_id
        raise


def get_attachment(identity: IdentityContext, attachment_id: str, *, session_id: str | None = None) -> dict[str, Any]:
    data_path, meta_path = _paths(identity, attachment_id)
    if not meta_path.is_file():
        raise AttachmentError("ATTACHMENT_NOT_FOUND", "附件不存在。", status_code=404)
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if meta.get("tenant_id") != identity.tenant_id or meta.get("user_id") != identity.user_id:
        raise AttachmentError("PERMISSION_DENIED", "无权访问该附件。", status_code=403)
    if session_id and meta.get("session_id") != session_id:
        raise AttachmentError("PERMISSION_DENIED", "附件不属于当前会话。", status_code=403)
    if datetime.fromisoformat(meta["expires_at"]) <= datetime.now(timezone.utc) and meta.get("status") not in TERMINAL_STATUSES:
        _discard_attachment_payload(
            meta,
            data_path,
            status="deleted",
            error={"code": "ATTACHMENT_EXPIRED", "message": "附件已过期。"},
        )
        _write_json(meta_path, meta)
    return meta


def delete_attachment(identity: IdentityContext, attachment_id: str) -> dict[str, Any]:
    meta = get_attachment(identity, attachment_id)
    data_path, meta_path = _paths(identity, attachment_id)
    if meta.get("status") not in {"cancelled", "deleted"}:
        terminal_status = "cancelled" if meta.get("status") in {"uploading", "parsing"} else "deleted"
        _discard_attachment_payload(meta, data_path, status=terminal_status)
        _write_json(meta_path, meta)
    return _public(meta)


def attachment_context(identity: IdentityContext, attachment_ids: list[str], *, session_id: str) -> dict[str, Any]:
    if len(attachment_ids) > 8:
        raise AttachmentError("ATTACHMENT_LIMIT_EXCEEDED", "单次附件数量超过限制。", status_code=422)
    records, citations, images, texts = [], [], [], []
    for attachment_id in attachment_ids:
        meta = get_attachment(identity, attachment_id, session_id=session_id)
        status = str(meta.get("status") or "")
        if status in {"deleted", "cancelled"}:
            raise AttachmentError("ATTACHMENT_NOT_AVAILABLE", "附件已删除或取消。", status_code=410)
        if status == "failed":
            raise AttachmentError("ATTACHMENT_PARSE_FAILED", "附件解析失败。", status_code=424)
        if status != "ready":
            raise AttachmentError("ATTACHMENT_NOT_READY", "附件尚未完成解析。", status_code=409)
        records.append(_public(meta))
        for chunk in meta.get("chunks") or []:
            text = str(chunk.get("text") or "")[:12_000]
            texts.append(f"<untrusted_attachment name={json.dumps(meta['file_name'], ensure_ascii=False)}>\n{text}\n</untrusted_attachment>")
            citations.append({
                "citation_id": f"acit_{uuid4().hex}", "attachment_id": attachment_id,
                "file_name": meta["file_name"], "location": chunk.get("location") or {},
                "quote": text[:240], "chunk_id": chunk.get("chunk_id"),
            })
        if (meta.get("properties") or {}).get("vision"):
            data_path, _ = _paths(identity, attachment_id)
            encoded = base64.b64encode(data_path.read_bytes()).decode("ascii")
            images.append({"attachment_id": attachment_id, "data_url": f"data:{meta['media_type']};base64,{encoded}"})
    return {"records": records, "citations": citations, "images": images, "untrusted_text": "\n".join(texts)[:60_000]}
