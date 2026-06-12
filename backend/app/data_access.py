from __future__ import annotations

import json
import logging
import math
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from sqlalchemy import inspect, text

from .config import APP_VERSION, PLATFORM_NAME, PROJECT_ROOT, project_config, project_paths
from .core.config import get_settings
from .observability import log_suppressed_exception


logger = logging.getLogger(__name__)


_DATABASE_STATUS: dict[str, Any] = {
    "primary": "unknown",
    "active": "none",
    "available": False,
    "message": "",
    "fallback_allowed": False,
}


FORECAST_FILE = "18_未来24小时预测结果_正式版.xlsx"
BUSINESS_SUMMARY_FILE = "19_业务统计摘要.xlsx"

DATA_SOURCE_CATALOG = {
    "日前价格": {
        "source": "PJM Data Miner 2 / Day-Ahead LMP",
        "source_url": "https://dataminer2.pjm.com/feed/da_hrl_lmps",
    },
    "实时价格": {
        "source": "PJM Data Miner 2 / Real-Time LMP",
        "source_url": "https://dataminer2.pjm.com/feed/rt_hrl_lmps",
    },
    "实际负荷": {
        "source": "PJM Data Miner 2 / Actual Load",
        "source_url": "https://dataminer2.pjm.com/feed/hrl_load_metered",
    },
    "预测负荷": {
        "source": "PJM Data Miner 2 / Load Forecast",
        "source_url": "https://dataminer2.pjm.com/feed/load_frcstd_7_day",
    },
    "天气数据": {
        "source": "NOAA / National Weather Service",
        "source_url": "https://www.weather.gov/documentation/services-web-api",
    },
    "建模主表": {
        "source": "项目特征工程主表（PJM + NOAA 清洗汇总）",
        "source_url": "https://dataminer2.pjm.com/",
    },
}


def jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat(sep=" ")
    if hasattr(value, "item"):
        return jsonable(value.item())
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, dict):
        return {str(k): jsonable(v) for k, v in value.items()}
    if isinstance(value, list):
        return [jsonable(v) for v in value]
    return value


def records(df: pd.DataFrame) -> list[dict[str, Any]]:
    if df is None or df.empty:
        return []
    clean = df.replace({pd.NA: None})
    return [jsonable(row) for row in clean.to_dict(orient="records")]


def read_excel_safe(path: Path) -> pd.DataFrame:
    settings = get_settings()
    if settings.has_database_url and not settings.database_allow_legacy_fallback:
        logger.warning("Legacy Excel fallback disabled by DATABASE_ALLOW_LEGACY_FALLBACK=0; source_type=excel")
        return pd.DataFrame()
    if not path.exists() or path.stat().st_size <= 0:
        return pd.DataFrame()
    try:
        return pd.read_excel(path, engine="openpyxl")
    except Exception as exc:
        log_suppressed_exception("data_access.read_excel_safe", exc, path=str(path))
        return pd.DataFrame()


def read_json_safe(path: Path) -> dict[str, Any]:
    settings = get_settings()
    if settings.has_database_url and not settings.database_allow_legacy_fallback:
        logger.warning("Legacy JSON fallback disabled by DATABASE_ALLOW_LEGACY_FALLBACK=0; source_type=json")
        return {}
    if not path.exists() or path.stat().st_size <= 0:
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        log_suppressed_exception("data_access.read_json_safe", exc, path=str(path))
        return {}


def price_column(df: pd.DataFrame) -> str | None:
    preferred = [
        "predicted_price",
        "corrected_predicted_price",
        "预测的未来24小时日前电价",
        "forecast_price",
        "da_price_pred",
    ]
    for column in preferred:
        if column in df.columns:
            return column
    for column in df.columns:
        name = str(column).lower()
        if "预测" in str(column) or "pred" in name:
            return str(column)
    return None


def risk_probability_column(df: pd.DataFrame) -> str | None:
    for column in ["spike_risk_prob", "尖峰风险概率", "spike_prob"]:
        if column in df.columns:
            return column
    return None


def load_latest_forecast() -> dict[str, Any]:
    try:
        from .repositories.forecast_repository import load_latest_forecast_from_postgres

        postgres_payload = load_latest_forecast_from_postgres()
        if postgres_payload and postgres_payload.get("available"):
            return postgres_payload
    except Exception as exc:
        log_suppressed_exception("data_access.load_latest_forecast.postgres", exc)

    paths = project_paths()
    path = paths.result_table_dir / FORECAST_FILE
    df = read_excel_safe(path)
    if df.empty:
        return {
            "run_id": "latest",
            "source": str(path),
            "available": False,
            "message": "未找到可用的未来 24 小时正式预测结果。",
            "summary": {},
            "records": [],
        }

    if "datetime" in df.columns:
        df["datetime"] = pd.to_datetime(df["datetime"], errors="coerce")
    pcol = price_column(df)
    if pcol:
        df["__price"] = pd.to_numeric(df[pcol], errors="coerce")
    else:
        df["__price"] = pd.NA
    prob_col = risk_probability_column(df)
    if prob_col:
        df["__risk_prob"] = pd.to_numeric(df[prob_col], errors="coerce").fillna(0)
    else:
        df["__risk_prob"] = 0.0
    if "risk_level" not in df.columns:
        df["risk_level"] = df["__risk_prob"].map(lambda v: "high" if v >= 0.5 else ("medium" if v >= 0.2 else "low"))

    summary = forecast_summary(df, pcol)
    summary_path = paths.current_dir / "ai_input_summary.json"
    ai_summary = read_json_safe(summary_path)
    run_id = str(ai_summary.get("run_id") or df.get("run_id", pd.Series(["latest"])).iloc[0] if not df.empty else "latest")
    return {
        "run_id": run_id,
        "source": str(path),
        "available": True,
        "generated_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(sep=" "),
        "price_column": pcol,
        "risk_probability_column": prob_col,
        "summary": summary,
        "records": records(df.drop(columns=[c for c in ["__price", "__risk_prob"] if c in df.columns])),
    }


def forecast_summary(df: pd.DataFrame, pcol: str | None = None) -> dict[str, Any]:
    if df.empty:
        return {}
    pcol = pcol or price_column(df)
    if not pcol:
        return {"rows": len(df)}
    prices = pd.to_numeric(df[pcol], errors="coerce")
    high_risk = df[df.get("risk_level", pd.Series(index=df.index, dtype=str)).astype(str).str.lower().isin(["high", "高", "高风险"])]
    max_idx = prices.idxmax() if prices.notna().any() else None
    min_idx = prices.idxmin() if prices.notna().any() else None
    start = pd.to_datetime(df["datetime"], errors="coerce").min() if "datetime" in df.columns else None
    end = pd.to_datetime(df["datetime"], errors="coerce").max() if "datetime" in df.columns else None
    max_row = df.loc[max_idx] if max_idx is not None else pd.Series(dtype=object)
    min_row = df.loc[min_idx] if min_idx is not None else pd.Series(dtype=object)
    return jsonable(
        {
            "rows": int(len(df)),
            "forecast_start": start,
            "forecast_end": end,
            "max_price": float(prices.max()) if prices.notna().any() else None,
            "min_price": float(prices.min()) if prices.notna().any() else None,
            "avg_price": float(prices.mean()) if prices.notna().any() else None,
            "peak_valley_spread": float(prices.max() - prices.min()) if prices.notna().any() else None,
            "max_hour": max_row.get("datetime") if not max_row.empty else None,
            "min_hour": min_row.get("datetime") if not min_row.empty else None,
            "high_risk_hours": int(len(high_risk)),
            "focus_hours": [
                str(v)
                for v in pd.to_datetime(high_risk.get("datetime", pd.Series(dtype=object)), errors="coerce")
                .dt.strftime("%Y-%m-%d %H:%M")
                .dropna()
                .head(6)
                .tolist()
            ],
        }
    )


def latest_business_summary() -> list[dict[str, Any]]:
    path = project_paths().result_table_dir / BUSINESS_SUMMARY_FILE
    return records(read_excel_safe(path))


def data_status() -> dict[str, Any]:
    paths = project_paths()
    sources = [
        ("日前价格", paths.data_dir / "da_price_raw.xlsx", "raw_da_price", "datetime", "da_price"),
        ("实时价格", paths.data_dir / "rt_price_raw.xlsx", "raw_rt_price", "datetime", "rt_price"),
        ("实际负荷", paths.data_dir / "actual_load_raw.xlsx", "raw_actual_load", "datetime", "actual_load"),
        ("预测负荷", paths.data_dir / "forecast_load_selected.xlsx", "raw_forecast_load_selected", "datetime", "forecast_load"),
        ("天气数据", paths.data_dir / "weather_raw.xlsx", "raw_weather", "datetime", "temperature"),
        ("建模主表", paths.data_dir / "master_table.xlsx", "model_master_table", "datetime", "da_price"),
    ]
    rows: list[dict[str, Any]] = []
    for name, path, table_name, dt_col, value_col in sources:
        status = database_table_status(table_name, dt_col, value_col)
        if not status.get("available"):
            status = excel_status_quick(path, dt_col, value_col)
        if not status.get("available"):
            source_info = DATA_SOURCE_CATALOG.get(name, {})
            rows.append(
                {
                    "name": name,
                    "status": "异常",
                    "source": source_info.get("source", "外部采集数据源"),
                    "source_url": source_info.get("source_url"),
                    "message": "文件不存在或无可读数据",
                }
            )
            continue
        source_info = DATA_SOURCE_CATALOG.get(name, {})
        rows.append(
            jsonable(
                {
                    "name": name,
                    "status": "正常" if status.get("latest_time") else "需检查",
                    "source": source_info.get("source", "外部采集数据源"),
                    "source_url": source_info.get("source_url"),
                    "internal_file": path.name,
                    "latest_time": status.get("latest_time"),
                    "start_time": status.get("start_time"),
                    "rows": status.get("rows"),
                    "missing_values": status.get("missing_values"),
                    "value_column": value_col if status.get("has_value_column") else "",
                    "updated_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(sep=" ") if path.exists() else None,
                }
            )
        )
    return {"sources": rows}


def _set_database_status(**values: Any) -> None:
    _DATABASE_STATUS.update(values)


def database_runtime_status() -> dict[str, Any]:
    if _DATABASE_STATUS.get("active") == "none":
        try:
            database_engine()
        except Exception as exc:
            log_suppressed_exception("data_access.database_runtime_status.probe", exc)
    return dict(_DATABASE_STATUS)


def database_table_status(table_name: str, datetime_col: str, value_col: str) -> dict[str, Any]:
    engine = database_engine()
    if engine is None:
        return {"available": False}
    dialect = engine.dialect.name
    try:
        if not inspect(engine).has_table(table_name):
            return {"available": False, "message": f"数据库表不存在：{table_name}"}
    except Exception as exc:
        log_suppressed_exception("data_access.database_table_status.inspect", exc, table_name=table_name)
        return {"available": False}
    safe_table = _quote_identifier(table_name, dialect)
    safe_datetime = _quote_identifier(datetime_col, dialect)
    safe_value = _quote_identifier(value_col, dialect)
    sql = f"""
        SELECT
            COUNT(*) AS rows_count,
            MIN({safe_datetime}) AS start_time,
            MAX({safe_datetime}) AS latest_time,
            SUM(CASE WHEN {safe_datetime} IS NULL THEN 1 ELSE 0 END)
              + SUM(CASE WHEN {safe_value} IS NULL THEN 1 ELSE 0 END) AS missing_values
        FROM {safe_table}
    """
    df = query_dataframe(sql)
    if df.empty:
        return {"available": False}
    row = df.iloc[0]
    return {
        "available": True,
        "rows": int(row.get("rows_count") or 0),
        "start_time": row.get("start_time"),
        "latest_time": row.get("latest_time"),
        "missing_values": int(row.get("missing_values") or 0),
        "has_value_column": True,
    }


def excel_status_quick(path: Path, datetime_col: str, value_col: str) -> dict[str, Any]:
    if not path.exists() or path.stat().st_size <= 0:
        return {"available": False}
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = workbook[workbook.sheetnames[0]]
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            header_map = {str(value): idx + 1 for idx, value in enumerate(headers) if value is not None}
            dt_idx = header_map.get(datetime_col)
            value_idx = header_map.get(value_col)
            return {
                "available": True,
                "rows": max(int(ws.max_row or 0) - 1, 0),
                "start_time": None,
                "latest_time": None,
                "missing_values": None,
                "has_value_column": bool(value_idx),
            }
        finally:
            workbook.close()
    except Exception as exc:
        log_suppressed_exception("data_access.excel_status_quick", exc, path=str(path))
        return {"available": False}


def database_engine():
    from .core.config import get_settings

    settings = get_settings()
    _set_database_status(
        primary=settings.database_primary,
        active="none",
        available=False,
        message="",
        fallback_allowed=settings.database_allow_legacy_fallback,
    )
    postgres_error = ""
    if settings.wants_postgres_primary:
        if not settings.has_database_url:
            postgres_error = "DATABASE_URL 未配置，PostgreSQL 主库不可用。"
            _set_database_status(message=postgres_error)
            if not settings.database_allow_legacy_fallback:
                return None
        else:
            try:
                from .db.session import get_engine

                engine = get_engine()
                if engine.dialect.name != "postgresql":
                    postgres_error = f"DATABASE_URL 当前方言不是 PostgreSQL：{engine.dialect.name}"
                    _set_database_status(message=postgres_error)
                    if not settings.database_allow_legacy_fallback:
                        return None
                else:
                    with engine.connect() as conn:
                        conn.execute(text("SELECT 1"))
                    _set_database_status(active="postgresql", available=True, message="PostgreSQL 主库连接正常。")
                    return engine
            except Exception as exc:
                postgres_error = log_suppressed_exception("data_access.database_engine.postgres", exc)
                _set_database_status(message=postgres_error)
                if not settings.database_allow_legacy_fallback:
                    return None

    if settings.wants_postgres_primary and not settings.database_allow_legacy_fallback:
        return None

    try:
        from .db.session import get_engine

        if settings.has_database_url:
            engine = get_engine()
            with engine.connect() as conn:
                conn.execute(text("SELECT 1"))
            _set_database_status(active=engine.dialect.name, available=True, message=f"{engine.dialect.name} 连接正常。")
            return engine
    except Exception as exc:
        log_suppressed_exception("data_access.database_engine.generic_url", exc)

    from database_utils import create_database_engine, get_database_config

    cfg = project_config()
    if not get_database_config(cfg).enabled:
        _set_database_status(message=postgres_error or "旧版数据库配置未启用。")
        return None
    try:
        engine = create_database_engine(cfg)
        _set_database_status(active=engine.dialect.name, available=True, message="已使用旧版数据库兼容连接。")
        return engine
    except Exception as exc:
        message = log_suppressed_exception("data_access.database_engine.legacy", exc)
        _set_database_status(message=postgres_error or message)
        return None


def query_dataframe(sql: str, params: dict[str, Any] | None = None) -> pd.DataFrame:
    engine = database_engine()
    if engine is None:
        return pd.DataFrame()
    try:
        with engine.connect() as conn:
            return pd.read_sql(text(sql), conn, params=params or {})
    except Exception as exc:
        log_suppressed_exception("data_access.query_dataframe", exc, sql=sql[:500])
        return pd.DataFrame()


def _quote_identifier(name: str, dialect: str = "mysql") -> str:
    if dialect == "postgresql":
        return '"' + str(name).replace('"', '""') + '"'
    return "`" + str(name).replace("`", "``") + "`"


def _text_search_expression(column_name: str, dialect: str) -> str:
    quoted = _quote_identifier(column_name, dialect)
    if dialect == "postgresql":
        return f"CAST({quoted} AS TEXT) ILIKE :keyword"
    return f"CAST({quoted} AS CHAR) LIKE :keyword"


def database_tables(search: str | None = None) -> dict[str, Any]:
    engine = database_engine()
    if engine is None:
        return {"available": False, "tables": [], "message": "数据库未启用或连接不可用。"}
    keyword = str(search or "").strip()
    schema_df = pd.DataFrame()
    if engine.dialect.name != "postgresql":
        schema_df = query_dataframe("SELECT DATABASE() AS schema_name")
    schema_name = str(schema_df.iloc[0].get("schema_name") or "") if not schema_df.empty else ""
    if schema_name:
        params = {"schema": schema_name, "keyword": f"%{keyword}%", "has_keyword": 1 if keyword else 0}
        info_df = query_dataframe(
            """
            SELECT
                t.TABLE_NAME AS table_name,
                COALESCE(t.TABLE_ROWS, 0) AS rows_count,
                COUNT(c.COLUMN_NAME) AS columns_count,
                GROUP_CONCAT(c.COLUMN_NAME ORDER BY c.ORDINAL_POSITION SEPARATOR ', ') AS column_names
            FROM information_schema.TABLES t
            LEFT JOIN information_schema.COLUMNS c
                ON c.TABLE_SCHEMA = t.TABLE_SCHEMA AND c.TABLE_NAME = t.TABLE_NAME
            WHERE t.TABLE_SCHEMA = :schema
              AND (:has_keyword = 0 OR t.TABLE_NAME LIKE :keyword)
            GROUP BY t.TABLE_NAME, t.TABLE_ROWS
            ORDER BY t.TABLE_NAME
            """,
            params,
        )
        if not info_df.empty:
            rows = []
            for _, row in info_df.iterrows():
                columns_text = str(row.get("column_names") or "")
                sample_columns = ", ".join([item for item in columns_text.split(", ") if item][:6])
                rows.append(
                    {
                        "table_name": row.get("table_name"),
                        "rows": int(row.get("rows_count") or 0),
                        "columns": int(row.get("columns_count") or 0),
                        "primary_key": "",
                        "sample_columns": sample_columns,
                    }
                )
            return {"available": True, "tables": jsonable(rows)}

    try:
        inspector = inspect(engine)
        table_names = sorted(inspector.get_table_names())
    except Exception as exc:
        log_suppressed_exception("data_access.database_tables.inspect", exc)
        return {"available": False, "tables": [], "message": f"读取数据库表失败：{exc}"}

    lower_keyword = keyword.lower()
    if lower_keyword:
        table_names = [name for name in table_names if lower_keyword in name.lower()]

    rows: list[dict[str, Any]] = []
    for table_name in table_names:
        try:
            columns = inspector.get_columns(table_name)
            pk = inspector.get_pk_constraint(table_name).get("constrained_columns") or []
            rows.append(
                {
                    "table_name": table_name,
                    "rows": None,
                    "columns": len(columns),
                    "primary_key": ", ".join(pk),
                    "sample_columns": ", ".join(str(col.get("name")) for col in columns[:6]),
                }
            )
        except Exception as exc:
            log_suppressed_exception("data_access.database_tables.table_metadata", exc, table_name=table_name)
            rows.append({"table_name": table_name, "rows": None, "columns": None, "primary_key": "", "sample_columns": ""})
    return {"available": True, "tables": jsonable(rows)}


def database_table_rows(table_name: str, search: str | None = None, limit: int = 100, offset: int = 0) -> dict[str, Any]:
    engine = database_engine()
    if engine is None:
        return {"available": False, "table_name": table_name, "columns": [], "records": [], "message": "数据库未启用或连接不可用。"}
    try:
        inspector = inspect(engine)
        table_names = set(inspector.get_table_names())
        if table_name not in table_names:
            return {"available": False, "table_name": table_name, "columns": [], "records": [], "message": "未找到该数据库表。"}
        columns = inspector.get_columns(table_name)
    except Exception as exc:
        log_suppressed_exception("data_access.database_table_rows.schema", exc, table_name=table_name)
        return {"available": False, "table_name": table_name, "columns": [], "records": [], "message": f"读取表结构失败：{exc}"}

    limit = max(1, min(int(limit or 100), 200))
    offset = max(0, int(offset or 0))
    dialect = engine.dialect.name
    safe_table = _quote_identifier(table_name, dialect)
    column_names = [str(col.get("name")) for col in columns if col.get("name")]
    searchable_columns = column_names[:12]
    params: dict[str, Any] = {}
    where = ""
    keyword = str(search or "").strip()
    if keyword and searchable_columns:
        params["keyword"] = f"%{keyword}%"
        predicates = [_text_search_expression(col, dialect) for col in searchable_columns]
        where = " WHERE " + " OR ".join(predicates)
    total_df = query_dataframe(f"SELECT COUNT(*) AS rows_count FROM {safe_table}{where}", params)
    data_df = query_dataframe(f"SELECT * FROM {safe_table}{where} LIMIT {limit} OFFSET {offset}", params)
    return {
        "available": True,
        "table_name": table_name,
        "columns": [{"name": name, "type": str(col.get("type") or "")} for name, col in zip(column_names, columns)],
        "records": records(data_df),
        "total": int(total_df.iloc[0].get("rows_count") or 0) if not total_df.empty else len(data_df),
        "limit": limit,
        "offset": offset,
        "search": keyword,
    }


def model_status() -> dict[str, Any]:
    try:
        from .repositories.model_repository import model_status_from_postgres

        postgres_payload = model_status_from_postgres()
        if postgres_payload:
            return postgres_payload
    except Exception as exc:
        log_suppressed_exception("data_access.model_status.postgres", exc)

    active = query_dataframe(
        """
        SELECT *
        FROM model_registry
        WHERE is_active = 1
        ORDER BY activated_at DESC, created_at DESC
        LIMIT 1
        """
    )
    registry = query_dataframe(
        """
        SELECT model_version, status, is_active, test_mae, test_rmse, peak_rmse, spike_rmse, feature_version, created_at, activated_at, artifact_path
        FROM model_registry
        ORDER BY created_at DESC
        LIMIT 20
        """
    )
    errors = query_dataframe(
        """
        SELECT model_version, COUNT(*) AS sample_count, AVG(abs_error) AS mae, MAX(created_at) AS latest_record
        FROM prediction_tracking
        WHERE actual_price IS NOT NULL
        GROUP BY model_version
        ORDER BY latest_record DESC
        LIMIT 20
        """
    )
    active_record = active.iloc[0].to_dict() if not active.empty else {}
    return {
        "active": jsonable(active_record),
        "versions": records(registry),
        "errors": records(errors),
        "source": "database" if not registry.empty or active_record else "local",
    }


def report_status(report_id: str = "latest") -> dict[str, Any]:
    try:
        from .repositories.report_repository import report_status_from_postgres

        postgres_payload = report_status_from_postgres(report_id)
        if postgres_payload:
            return postgres_payload
    except Exception as exc:
        log_suppressed_exception("data_access.report_status.postgres", exc, report_id=report_id)

    paths = project_paths()
    report_path = paths.current_dir / "电价智能分析综合报告.docx"
    structured = read_json_safe(paths.current_dir / "ai_report_structured.json")
    ai_summary = read_json_safe(paths.current_dir / "ai_input_summary.json")
    rid = str(ai_summary.get("run_id") or report_id or "latest")
    return {
        "report_id": rid if report_id == "latest" else report_id,
        "run_id": rid,
        "available": report_path.exists(),
        "report_path": str(report_path),
        "generated_at": datetime.fromtimestamp(report_path.stat().st_mtime).isoformat(sep=" ") if report_path.exists() else None,
        "fallback_used": bool(structured.get("fallback_used", False)),
        "summary": structured,
        "source": str(paths.current_dir),
    }


def dashboard_summary() -> dict[str, Any]:
    forecast = load_latest_forecast()
    model = model_status()
    report = report_status()
    data = data_status()
    return {
        "platform": PLATFORM_NAME,
        "version": APP_VERSION,
        "generated_at": datetime.now().isoformat(sep=" ", timespec="seconds"),
        "market": project_config().get("market", {}).get("pjm_node_name", "DOM"),
        "forecast": forecast.get("summary", {}),
        "forecast_available": forecast.get("available", False),
        "forecast_source": "正式预测结果（基于 PJM、NOAA 等采集数据清洗生成）",
        "model": model.get("active", {}),
        "report": {
            "available": report.get("available"),
            "report_id": report.get("report_id"),
            "generated_at": report.get("generated_at"),
        },
        "data_sources": data.get("sources", []),
    }


def artifact_inventory() -> list[dict[str, Any]]:
    root = PROJECT_ROOT / "model_artifacts"
    if not root.exists():
        return []
    rows = []
    for child in sorted([p for p in root.iterdir() if p.is_dir()], key=lambda p: p.stat().st_mtime, reverse=True):
        manifest = read_json_safe(child / "manifest.json")
        metrics = read_json_safe(child / "metrics.json")
        rows.append(
            {
                "model_version": manifest.get("model_version") or child.name,
                "artifact_path": str(child),
                "created_at": manifest.get("created_at") or datetime.fromtimestamp(child.stat().st_mtime).isoformat(sep=" "),
                "base_model_name": metrics.get("base_model_name"),
                "peak_model_name": metrics.get("peak_model_name"),
                "classifier_name": metrics.get("classifier_name"),
                "has_p90_model": metrics.get("has_p90_model"),
            }
        )
    return rows
