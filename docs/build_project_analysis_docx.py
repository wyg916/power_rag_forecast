from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "《智能运营分析项目》项目全面分析说明文档_v20260605.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(9)
    run.font.name = "Microsoft YaHei"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
            if row_idx == 0:
                set_cell_shading(cell, "E8EEF5")


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[object]]) -> None:
    data = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True)
    for row in data:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    style_table(table)
    doc.add_paragraph()


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_note(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{title}：")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run2 = p.add_run(text)
    run2.font.name = "Microsoft YaHei"
    run2.font.size = Pt(10)
    run2._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("《智能运营分析项目》\n项目全面分析说明文档")
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string("0B2545")
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_before = Pt(20)
    sub.add_run("面向开发、架构评审、项目交接与技术答辩").font.size = Pt(12)

    add_table(
        doc,
        ["项目项", "说明"],
        [
            ["项目路径", str(PROJECT_ROOT)],
            ["文档版本", "v20260605"],
            ["分析范围", "Python 自动化预测、FastAPI 后端、React/Vite 前端、数据层、模型层、AI 助手、数据库迁移、部署与运维脚本"],
            ["重要声明", "本文基于本地代码、配置、README、迁移脚本、测试和样例数据字段分析；涉及推断处均标注“根据代码推断”。"],
        ],
    )
    doc.add_page_break()


def build_doc() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("1. 项目概述", level=1)
    doc.add_paragraph(
        "智能运营分析项目是一个面向售电交易与电力市场运营分析的本地化 AI 决策平台。"
        "项目以 PJM/DOM 日前电价预测为主线，围绕数据采集、特征工程、模型训练、未来 24 小时正式前瞻预测、"
        "AI 日报生成、模型运维、Web 工作台、桌面 GUI 和知识库问答形成完整闭环。"
    )
    add_table(
        doc,
        ["项目基本项", "结论"],
        [
            ["项目名称", "智能运营分析项目 / AI 售电交易决策平台"],
            ["项目类型", "Python 数据分析与机器学习项目 + FastAPI 后端 + React/Vite 前端 + PySide6 桌面控制台"],
            ["业务定位", "面向售电公司或电力交易运营人员，提供日前电价预测、风险识别、交易辅助建议、AI 问答和报告生成"],
            ["核心问题", "把电价、负荷、天气、历史误差、模型运行状态转化为可解释的预测结果和运营决策参考"],
            ["完成度判断", "已具备本地端到端可运行能力；Web、GUI、模型运维和 AI 助手已落地。生产级登录鉴权、权限持久化、向量召回、线上监控仍需完善。"],
            ["运行方式概览", "本地可通过 run_project.bat、run_web_platform.bat、python main_daily_run.py、python -m uvicorn backend.app.main:app、npm run dev 等入口运行"],
        ],
    )

    doc.add_heading("2. 项目业务目标与应用场景", level=1)
    add_bullets(
        doc,
        [
            "日前价格预测：生成未来 24 小时正式前瞻电价，并输出预测明细、图表和业务统计摘要。",
            "交易风险识别：识别高价、低价、高峰、尖峰概率、峰谷价差和重点关注时段。",
            "运营报告生成：将预测、异常、模型效果、特征重要性等结构化为 AI 日报和 Word 报告。",
            "模型生命周期管理：登记模型 artifact、追踪预测误差、判断是否重训、支持 Active 模型快速推理。",
            "Web 工作台：提供数据中心、预测中心、策略中心、AI 助手、报告中心、模型中心、知识库、任务中心和系统设置页面。",
            "知识库问答：围绕电价、负荷、天气、售电交易规则、光伏电价政策等内容提供证据驱动问答。",
        ],
    )

    doc.add_heading("3. 项目技术栈", level=1)
    add_table(
        doc,
        ["层级", "技术/依赖", "实际用途"],
        [
            ["语言与运行时", "Python 3.11（根据 README 和启动脚本推断）、TypeScript、Node/Vite", "预测、后端 API、桌面 GUI、前端工作台"],
            ["数据处理", "pandas、numpy、openpyxl", "读取 Excel/CSV、构建主表、特征工程、结果统计"],
            ["机器学习", "scikit-learn、可选 xgboost、lightgbm、catboost（代码中 catboost 为可选）", "基础回归、高峰专项回归、尖峰分类、模型融合"],
            ["后端", "FastAPI、Pydantic、SQLAlchemy、Alembic", "REST API、参数校验、PostgreSQL/兼容数据库访问、迁移管理"],
            ["任务调度", "Celery、Redis、Windows 任务计划、subprocess", "Web 任务异步执行、计划任务、脚本流水线"],
            ["前端", "React 18、Vite、TypeScript、Ant Design、ECharts", "Web 工作台、图表、表格、导航与操作界面"],
            ["桌面端", "PySide6、PySide6-Fluent-Widgets", "现代化桌面控制台和旧版 GUI 兼容入口"],
            ["AI/LLM", "Ollama qwen3:4b、OpenAI-compatible 网关、可选 Dify、本地规则工具链", "AI 助手、报告生成、多 Agent 分析、结构化回答"],
            ["部署", "Docker、docker-compose、Nginx 前端静态服务", "本地轻量部署和企业版 PostgreSQL + Redis + Worker 组合"],
        ],
    )

    doc.add_heading("4. 项目整体架构", level=1)
    doc.add_paragraph(
        "整体链路可以概括为：用户操作或计划任务 → 前端/GUI/批处理入口 → FastAPI 或 main_daily_run 调度 → 数据刷新与数据库同步 → "
        "预测引擎训练或 Active 模型快速推理 → 模型登记与预测追踪 → AI 摘要与 LLM 报告 → 归档派发 → 前端/GUI 展示。"
    )
    add_table(
        doc,
        ["架构层", "核心文件/目录", "职责"],
        [
            ["前端层", "frontend/src/app、frontend/src/pages、frontend/src/api.ts", "单页应用、页面导航、统一 API 封装、图表和表格展示"],
            ["接口层", "backend/app/main.py、backend/app/api/v1/endpoints/*.py", "FastAPI 应用、CORS、路由注册、接口鉴权、审计记录"],
            ["服务层", "services/*、backend/app/services/*、backend/app/platform_services.py", "流水线调度、Web 数据聚合、策略/异常/报告服务、核心数据同步"],
            ["数据层", "database_utils.py、backend/app/data_access.py、backend/app/repositories/*、migrations/*", "数据库连接、迁移、表浏览、结果同步、PostgreSQL 仓储"],
            ["模型层", "prediction_engine/*、高峰尖刺增强版...v4_fix1.py、model_ops/*", "特征工程、训练、预测、artifact 管理、误差追踪、模型上线和偏差校正"],
            ["AI 层", "backend/app/ai、backend/app/ai_assistant、backend/langgraph_agent、backend/model_gateway、knowledge_base", "意图识别、工具调用、Prompt、RAG 检索、本地 LLM 总结、Trace 和 Answer Guard"],
            ["任务层", "main_daily_run.py、backend/app/workers/*、create_windows_task.ps1、run_*.bat", "端到端自动化、Celery 任务、Windows 定时任务、一键启动"],
            ["输出层", "结果-3、自动化输出、全流程预测结果数据存放、model_artifacts、output", "预测表、报告、归档、日志、模型版本和数据快照"],
        ],
    )

    doc.add_heading("5. 项目目录结构说明", level=1)
    add_table(
        doc,
        ["路径", "类型", "作用", "是否核心", "说明"],
        [
            ["backend/app", "后端应用", "FastAPI 主应用、接口、服务、仓储、安全、AI 助手", "是", "Web 平台 API 的主要实现位置"],
            ["backend/model_gateway", "模型网关", "OpenAI-compatible 本地模型转发接口", "是", "封装 LOCAL_LLM_BASE_URL、LOCAL_LLM_MODEL、/api/model-gateway/chat/completions"],
            ["backend/langgraph_agent", "多 Agent 工作流", "意图、数据、预测、风险、报告、回答整合节点", "是", "代码实现为顺序节点图，命名为 LangGraph Agent 工作流"],
            ["frontend/src", "Web 前端", "React 页面、布局、图表、API 封装、mock 降级数据", "是", "Vite SPA，默认端口 5173"],
            ["prediction_engine", "预测引擎包", "特征工程、模型候选、训练封装、正式前瞻、快速推理", "是", "大量逻辑委托到 v4_fix1 兼容引擎"],
            ["model_ops", "模型运维", "模型注册、Active 加载、误差记忆、偏差校正、重训判断、上线管理", "是", "支撑模型生命周期闭环"],
            ["services", "自动化服务层", "主流程调度、数据刷新、预测、报告、归档、模型优化", "是", "main_daily_run 的服务化入口"],
            ["migrations", "数据库迁移", "MySQL 兼容 SQL 和 Alembic PostgreSQL 迁移", "是", "包含 forecast、model、task、ai、knowledge、rbac、audit 等表"],
            ["knowledge_base", "业务知识库", "电力市场、价格机制、风险等级、天气负荷关系等 Markdown 文档", "是", "供轻量 RAG/关键词检索使用"],
            ["output", "数据快照", "master_table、weather、load、price 等数据文件", "是", "预测输入的本地文件模式来源"],
            ["model_artifacts", "模型产物", "base_model、peak_model、spike_classifier、metrics、schema、model_card", "是", "Active 快速预测和模型回滚依赖"],
            ["自动化输出", "运行产物", "current、dispatch、logs、review_status 等", "是", "AI 摘要、报告、归档和日志输出"],
            ["全流程预测结果数据存放", "历史归档", "按运行批次保存预测、AI 报告和日志", "否", "用于历史追溯，不是主要源码"],
            ["docker", "部署配置", "后端/前端 Dockerfile、Nginx、entrypoint", "是", "docker-compose 使用"],
            ["tests", "测试", "Web、AI、模型、迁移、安全、自检测试", "是", "覆盖关键接口和核心规则"],
        ],
    )

    doc.add_heading("6. 核心功能模块分析", level=1)
    add_table(
        doc,
        ["模块", "作用", "主要文件", "输入", "输出", "与其他模块关系"],
        [
            ["数据采集与刷新", "从 PJM 和 Open-Meteo 拉取电价、负荷、天气并构建主表", "fetch_power_market_data.py、services/data_service.py", "PJM Data Miner 2、Open-Meteo、output 历史文件", "output/*.xlsx、master_table、数据库核心表", "为预测引擎和 Web 数据中心提供输入"],
            ["特征工程", "构造时间、滞后、滚动、EWM、同小时历史统计、安全业务特征并过滤泄漏字段", "prediction_engine/features.py、v4_fix1", "master_table", "建模特征表、feature_cols", "模型训练和快速预测共享特征定义"],
            ["模型训练", "训练基础模型、高峰模型、尖峰分类器并搜索融合参数", "prediction_engine/model_trainer.py、model_selection.py、v4_fix1", "训练/验证/测试集、feature_cols", "模型评估、预测结果、model_artifacts", "输出登记到 model_registry 并供 Active 推理"],
            ["正式前瞻预测", "从最后历史时点向后逐小时递推未来 24 小时", "prediction_engine/formal_forecast.py、fast_forecast.py", "历史主表、未来负荷、天气、Active 模型", "未来 24 小时预测表、输入特征表、预测图", "被 Web 预测中心、AI 摘要、报告和策略模块消费"],
            ["模型运维", "真实值回填、误差统计、退化判断、候选模型上线、偏差校正", "model_ops/*.py、05-12 脚本", "prediction_tracking、raw_da_price、model_registry", "performance_daily、retrain_jobs、error_memory", "为模型中心和快速预测提供闭环"],
            ["AI 摘要与报告", "把预测、模型、异常、特征重要性汇总为结构化摘要并调用 LLM 生成 Word 报告", "02_build_ai_summary.py、03_llm_generate_report.py、04_dispatch_report.py", "结果表、ai_input_summary.json、Ollama", "ai_report_structured.json、Word 报告、归档", "报告中心、AI 助手和派发流程读取结果"],
            ["Web API", "把预测、数据、模型、任务、报告、策略、知识库、AI 问答封装为 REST 接口", "backend/app/api/v1/endpoints/*.py", "HTTP 请求、数据库、文件结果", "JSON、FileResponse、任务记录", "前端和自动化入口的服务边界"],
            ["Web 前端", "业务工作台，提供数据、预测、策略、AI、报告、模型、知识库和任务页面", "frontend/src/pages、frontend/src/api.ts", "后端 API、mock 降级数据", "浏览器 UI、图表、表格、下载", "面向用户操作和展示"],
            ["AI 助手", "意图识别、工具调用、上下文记忆、LLM 总结、Answer Guard、Trace", "backend/app/ai_assistant、backend/app/ai、backend/langgraph_agent", "用户问题、会话状态、工具结果、知识库", "结构化回答、证据、工具调用、Trace", "前端 AI 助手与报告/预测/数据/知识库联动"],
            ["任务调度", "手动或定时触发刷新、预测、报告、自检和模型优化", "main_daily_run.py、backend/app/workers、schedule_service.py、run_*.bat", "用户操作、计划任务、Celery 队列", "任务状态、日志、结果文件", "连接前端、服务层和脚本流水线"],
        ],
    )

    doc.add_heading("7. 核心代码文件说明", level=1)
    add_table(
        doc,
        ["文件", "关键函数/类", "作用", "输入", "输出", "调用关系"],
        [
            ["backend/app/main.py", "create_app", "创建 FastAPI 应用、配置 CORS、注册模型网关和业务路由", "无", "FastAPI app", "uvicorn 启动入口"],
            ["backend/app/api/v1/router.py", "api_router.include_router", "集中注册 system/data/forecast/strategy/assistant/knowledge/report/model/task 路由", "路由模块", "APIRouter", "被 main.py include"],
            ["main_daily_run.py", "main、build_sequence、run_script", "每日自动化总控，构建 run_id/run_mode 并按模式执行流水线", "命令行参数、配置", "日志、事件、结果文件", "调用 services.pipeline_steps.run_pipeline_step"],
            ["services/prediction_service.py", "run_prediction、_run_fast_forecast", "预测服务入口，支持完整训练和 Active 快速预测", "配置、运行上下文、env", "PredictionRunResult、数据库同步", "调用 prediction_engine 和 model_ops"],
            ["fetch_power_market_data.py", "fetch_pjm_feed、fetch_weather、build_master_table", "外部数据拉取和主表构建", "PJM/Open-Meteo/API Key", "output 数据集", "被 refresh_market_data 调用"],
            ["prediction_engine/model_trainer.py", "save_model_artifacts、load_model_artifacts", "保存/加载模型 artifact", "训练结果、特征列", "joblib + JSON + model_card", "model_ops Active 加载依赖"],
            ["prediction_engine/fast_forecast.py", "forecast_with_saved_model", "加载 Active 模型快速预测未来 24 小时", "ActiveModelBundle、master_table", "预测表、特征快照、业务摘要、图", "被 prediction_service 快速模式调用"],
            ["model_ops/active_model_loader.py", "get_active_model_record、load_active_model_artifacts、predict_with_active_model", "查找 Active 模型并加载 artifact", "数据库或配置路径", "ActiveModelBundle、预测 DataFrame", "被 fast_forecast 调用"],
            ["backend/app/ai_assistant/service.py", "answer_chat_accurate", "新版准确问答主链路：标准化、上下文、意图、工具、答案、Guard、Trace", "用户问题、session、run_id", "回答、证据、工具、Trace", "被 /api/ai/agent/analyze 和 /api/ai/chat 调用"],
            ["backend/app/ai/tool_registry.py", "TOOLS、INTENT_TOOLS、call_tools", "旧版/兼容工具注册中心", "意图和实体", "工具结果、调用记录", "AI 助手和 Agent 节点复用"],
            ["backend/langgraph_agent/nodes.py", "intent_agent、data_agent、prediction_agent、risk_agent、report_agent、answer_agent", "多 Agent 顺序工作流节点", "AgentState", "增强后的 AgentState", "PowerAgentGraph.run 逐节点调用"],
            ["backend/app/core/security.py", "get_current_user、require_permission", "基于 Header/dev fallback 的轻量权限校验", "X-User/X-Role/AUTH_REQUIRED", "CurrentUser 或 401/403", "接口依赖注入"],
            ["database_utils.py", "apply_database_migrations、sync_result_tables_to_database、save_ai_report", "兼容数据库初始化、结果同步、报告/摘要入库", "配置、DataFrame/文件", "数据库表记录", "自动化和后端服务共用"],
            ["frontend/src/api.ts", "request、api.*", "前端统一 HTTP 封装和接口路径集中管理", "path、payload", "Promise JSON/File URL", "所有页面通过该文件访问后端"],
        ],
    )

    doc.add_heading("8. 项目启动与运行流程", level=1)
    add_numbered(
        doc,
        [
            "准备环境：复制 .env.example 为 .env，按需配置 DATABASE_URL、PJM_SUBSCRIPTION_KEY、LLM_BASE_URL、LLM_MODEL、Redis 和端口。",
            "安装依赖：Python 侧执行 python -m pip install -r requirements.txt；前端进入 frontend 执行 npm install。",
            "端到端流水线：执行 python main_daily_run.py，可附加 --refresh-data、--skip-prediction、--fast-forecast、--retrain-model。",
            "Web 后端：执行 python -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000。",
            "Web 前端：进入 frontend 执行 npm run dev，默认访问 http://127.0.0.1:5173。",
            "桌面 GUI：执行 python -m ui.app 或双击 启动智能运营分析GUI.bat。",
            "Docker 企业版：执行 docker-compose -f docker-compose.enterprise.yml up，启动 PostgreSQL、Redis、backend、worker、frontend。",
        ],
    )
    add_note(
        doc,
        "启动链路说明",
        "run_project.bat 是统一菜单入口；run_web_platform.bat 用于启动 Web；run_daily_pipeline*.bat 用于不同流水线模式；create_windows_task.ps1 用于创建 Windows 计划任务。",
    )

    doc.add_heading("9. 数据来源与数据处理流程", level=1)
    add_table(
        doc,
        ["数据集", "字段/含义", "来源", "处理方式", "用途"],
        [
            ["da_price_raw", "datetime、da_price、拥塞、损耗、系统能量价格", "PJM Data Miner 2 da_hrl_lmps，DOM 节点 34964545", "按小时标准化、增量合并去重", "预测目标变量和价格分析"],
            ["rt_price_raw", "rt_price", "PJM Data Miner 2 rt_hrl_lmps", "可选增强数据，按小时合并", "实时价格参考和扩展特征"],
            ["actual_load_raw", "actual_load", "PJM hrl_load_metered，DOM 负荷区", "小时粒度标准化", "负荷影响分析和主表特征"],
            ["forecast_load_raw / selected", "forecast_load、forecast_evaluated_at", "PJM load_frcstd_hist", "按日前预测版本筛选，缺失时可从 raw 修复", "未来前瞻预测的负荷输入"],
            ["weather_raw", "temperature、wind_speed、precipitation", "Open-Meteo Historical API / ERA5，Richmond, Virginia", "缺口用 best_match 回填并记录 weather_model/source", "天气特征和问答依据"],
            ["master_table", "datetime、region_id、da_price、actual_load、forecast_load、weather、is_holiday", "由价格、负荷、天气合并生成", "按 datetime 合并、排序、去重、补充节假日和来源字段", "建模主表和快速预测历史输入"],
            ["prediction_tracking", "run_id、model_version、forecast_datetime、predicted_price、actual_price、error", "预测结果 + 真实值回填", "写入数据库/CSV 样例，用于误差统计", "模型监控、误差记忆、偏差校正"],
        ],
    )
    doc.add_paragraph(
        "数据清洗逻辑包括：时间字段标准化、空时间行删除、按时间排序、重复记录删除、PJM 分页数据合并、历史缓存增量补抓、"
        "负荷预测按 forecast_evaluated_at 选择日提前版本、天气缺失回填、主表按 datetime 合并。"
    )

    doc.add_heading("10. 数据分析、特征工程与建模流程", level=1)
    add_numbered(
        doc,
        [
            "读取 output/master_table.xlsx 或数据库导出的预测输入数据。",
            "standardize_datetime_frame 统一时间字段、去重、排序。",
            "create_time_features 构造年、月、日、小时、星期、周末、节假日等时间特征。",
            "create_lag_features 为 da_price 等目标/关键字段生成 1、2、3、6、12、18、24、25、48、72、96、168、169 小时滞后特征。",
            "create_rolling_features / create_ewm_features 生成 6、12、24、48、168 小时滚动和指数滑动统计。",
            "create_history_group_features 生成同小时、星期、月份、周末等历史分组统计。",
            "add_safe_business_features 增加安全业务特征；leakage_check_report 与 build_feature_list 过滤 LEAKAGE_BANNED_RAW_COLS，降低数据泄漏风险。",
            "按时间切分训练、验证、测试集，训练基础模型、高峰专项模型和尖峰分类器。",
            "通过验证集搜索 best_alpha、best_peak_floor、best_spike_threshold 等融合参数。",
            "输出测试集评估、特征重要性、异常明细、滚动回测、未来 24 小时正式预测、业务统计摘要和模型 artifact。",
        ],
    )

    doc.add_heading("11. 模型与算法逻辑说明", level=1)
    add_table(
        doc,
        ["模型/算法", "所在文件", "输入", "输出", "作用", "评估方式", "备注"],
        [
            ["基础回归模型", "v4_fix1、prediction_engine/model_selection.py", "feature_cols 训练特征、da_price", "base_prediction", "拟合常规日前电价水平", "MAE、RMSE、MAPE、R2", "候选包含 XGBoost、LightGBM、HistGradientBoosting、ExtraTrees、RandomForest、Ridge、LinearRegression 等"],
            ["高峰专项回归模型", "v4_fix1、model_selection.py", "高峰时段样本和相同特征", "peak_prediction", "增强早晚高峰、尖峰价格拟合能力", "整体 RMSE、高峰 RMSE、尖峰 RMSE", "包含 Peak_XGBoost、Peak_LightGBM、Peak_ExtraTrees、Peak_RandomForest 等候选"],
            ["尖峰风险分类器", "v4_fix1、model_selection.py", "建模特征、尖峰标签", "spike_risk_prob", "识别价格尖峰或高风险概率", "分类验证、阈值搜索", "候选包含 LightGBM/XGBoost/CatBoost/RF/HGB/Logistic 等，缺依赖时跳过"],
            ["融合预测", "model_selection.blend_predictions、active_model_loader.predict_with_active_model", "base_pred、peak_pred、risk_prob、peak_hour_flag、负荷/误差高风险标志", "blended_pred/predicted_price、blend_weight", "在尖峰概率或高峰条件下动态提升高峰模型权重", "验证集搜索 best_alpha/best_floor", "是最终预测主逻辑"],
            ["递推正式前瞻", "prediction_engine/formal_forecast.py、v4_fix1", "历史主表、feature_cols、模型 artifact、未来负荷/天气代理", "未来 24 小时预测和输入特征", "从最后可用历史时点逐小时预测并回填滞后特征", "与回测/测试集结合验证", "README 明确已替代演示版直接取历史末 24 行"],
            ["残差区间", "formal_forecast.add_residual_prediction_intervals", "测试集残差、未来点预测", "P10/P90 下上界、区间宽度", "表达价格波动风险范围", "按小时残差分位数，缺小时则用全局分位数", "不改变点预测，只补充风险解释字段"],
            ["偏差校正", "model_ops/bias_corrector.py", "预测行、model_error_memory", "corrected_predicted_price、bias_adjustment", "按历史误差记忆校正预测偏差", "比较 correction effect，按样本数门槛启用", "最大调整幅度默认 10，避免过度修正"],
            ["模型退化判断", "model_ops/model_monitor.py、auto_retrain_policy.py", "prediction_tracking 近期/基线误差", "should_retrain、retrain_reason", "发现 RMSE 上升或超阈值时登记重训建议", "近 7 天 vs 30 天、RMSE>1.2 倍或 >80", "依赖真实值回填"],
        ],
    )

    doc.add_heading("12. AI 大模型接入与调用逻辑", level=1)
    doc.add_paragraph(
        "当前项目已发现多处 AI 大模型接入代码：报告生成使用根目录 llm_client.py 与 03_llm_generate_report.py；"
        "Web AI 助手使用 backend/app/ai_assistant、backend/app/ai、backend/langgraph_agent；模型网关使用 backend/model_gateway。"
    )
    add_table(
        doc,
        ["AI 能力", "调用方式", "入口文件", "Prompt/上下文", "输出处理", "可靠性设计"],
        [
            ["AI 日报生成", "Ollama native API，默认 qwen3:4b", "03_llm_generate_report.py、llm_client.py", "build_prompts 基于 ai_input_summary 构造事实输入，要求严格 JSON", "JSON schema 校验，缺字段重试，失败回退模板，写 Word 报告", "JSON 片段提取、schema 校验、指数退避重试、领域关键词检查、fallback report"],
            ["本地模型网关", "OpenAI-compatible /chat/completions", "backend/model_gateway/router.py", "messages 由调用方传入", "透传 QwenClient 返回", "空 messages 返回 400，模型调用失败返回 502"],
            ["新版准确问答", "工具优先 + 可选本地 LLM 总结", "backend/app/ai_assistant/service.py", "normalize_question、route_intent、tools_for_intent、evidence、draft_answer", "answer、intent、tools、evidence、trace、data_used", "Answer Guard、Trace、上下文记忆、LLM 思考泄漏过滤、工具模板兜底"],
            ["旧版数据增强助手", "Dify 可选 + 本地工具链", "backend/app/ai/assistant_service.py", "classify、call_tools、build_answer", "结构化回答、证据和相关操作", "Dify 失败自动回退；当前入口直接委托 answer_chat_accurate，旧逻辑位于不可达分支，属于兼容遗留代码"],
            ["多 Agent 分析", "顺序 Agent 节点 + 本地模型兜底", "backend/langgraph_agent", "intent_agent、data_agent、prediction_agent、risk_agent、report_agent、answer_agent 构造 _brief_payload", "专业回答、workflow、agent_trace、focus_periods", "本地模型失败 60 秒熔断，fallback_answer 基于工具证据生成"],
            ["知识库/RAG", "PostgreSQL kb_documents/kb_chunks + 关键词/轻量检索", "backend/app/services/rag_service.py、repositories/knowledge_repository.py", "本地 knowledge_base 和 tariff policy 文件切片入库", "items、evidence、stats", "当前 embedding_json 字段存在，但未发现真正向量 embedding 生成；根据代码推断仍以关键词/JSONB 检索为主"],
        ],
    )
    doc.add_paragraph(
        "AI 助手完整调用链路：用户输入 → /api/ai/chat 或 /api/ai/agent/analyze → 输入标准化 → 多轮上下文解析 → 意图识别 → "
        "按意图选择工具 → 查询预测/天气/负荷/模型/报告/知识库/电价规则 → 生成模板草稿 → 可选 Ollama qwen3 总结 → "
        "Answer Guard 校验/改写 → 保存 Trace、会话状态和反馈入口 → 前端展示回答、证据、工具和 Trace。"
    )

    doc.add_heading("13. 接口与 API 调用说明", level=1)
    add_table(
        doc,
        ["接口路径/方法", "请求方式", "所在文件", "入参", "出参", "作用", "调用方"],
        [
            ["/api/health", "GET", "system.py", "无", "ok、platform、version", "健康检查", "前端/部署 healthcheck"],
            ["/api/dashboard/summary", "GET", "forecast.py", "无", "forecast、data_sources 等", "首页汇总", "DashboardPage"],
            ["/api/data/status", "GET", "data.py", "无", "sources、状态", "数据源状态", "DataCenterPage、AI 工具"],
            ["/api/data/tables", "GET", "data.py", "search", "tables", "数据库表列表", "数据中心"],
            ["/api/data/tables/{table}/rows", "GET", "data.py", "search、limit、offset", "columns、records", "表数据预览", "数据中心"],
            ["/api/data/tables/{table}/export", "GET", "data.py", "search", "CSV 文件", "导出表", "数据中心"],
            ["/api/data/refresh", "POST", "data.py", "无", "task_id/run_id", "刷新外部数据", "数据中心/任务"],
            ["/api/data/sync-core", "POST", "data.py", "无", "task_id/run_id", "核心数据入库", "数据中心"],
            ["/api/forecast/run", "POST", "forecast.py", "mode", "任务信息", "触发快速预测/刷新预测/重训", "预测中心"],
            ["/api/forecast/latest", "GET", "forecast.py", "无", "records、run_id、available", "最新预测明细", "预测中心、AI 工具"],
            ["/api/prediction/latest", "GET", "forecast.py", "market、date", "预测概览", "阶段一标准预测接口", "AI 工具/外部调用"],
            ["/api/weather/forecast", "GET", "forecast.py", "city、date", "天气摘要/记录", "标准天气接口", "AI 工具"],
            ["/api/load/forecast", "GET", "forecast.py", "market、date", "负荷摘要/记录", "标准负荷接口", "AI 工具"],
            ["/api/model/explain", "GET", "forecast.py", "market、date", "特征/置信度", "模型解释", "AI 工具"],
            ["/api/risk/level", "GET", "forecast.py", "market、date", "risk_level、focus_periods", "风险等级", "AI 工具"],
            ["/api/strategy/latest", "GET", "strategy.py", "无/run_id", "items", "策略建议", "策略中心"],
            ["/api/anomaly/latest", "GET", "strategy.py", "无/run_id", "items", "异常解释", "策略中心/AI"],
            ["/api/ai/chat", "POST", "assistant.py", "ChatRequest", "answer、intent、tools、trace", "AI 聊天", "AI 助手"],
            ["/api/ai/agent/analyze", "POST", "assistant.py", "AgentAnalyzeRequest", "workflow、agent_trace、answer", "复杂分析", "AI 助手"],
            ["/api/ai/chat/feedback", "POST", "assistant.py", "session_id、trace_id、rating、comment", "ok", "问答反馈", "AI 助手"],
            ["/api/ai/traces", "GET", "assistant.py", "limit、session_id", "traces", "Trace 列表", "AI 助手 Trace 页"],
            ["/api/knowledge/search", "GET", "knowledge.py", "q、top_k", "items、evidence、stats", "知识库检索", "知识库页/AI"],
            ["/api/knowledge/index-local", "POST", "knowledge.py", "无", "indexed_documents、failed、stats", "本地知识库索引", "知识库页"],
            ["/api/reports/generate", "POST", "report.py", "run_id", "任务信息", "生成报告", "报告中心/首页"],
            ["/api/reports/latest", "GET", "report.py", "无", "report_status", "最新报告", "报告中心"],
            ["/api/reports/{id}/download", "GET", "report.py", "report_id", "docx 文件", "下载报告", "报告中心"],
            ["/api/models/metrics", "GET", "model.py", "无", "active、metrics、artifacts", "模型指标", "模型中心"],
            ["/api/models/retrain-suggestion", "GET", "model.py", "无", "should_retrain、reason", "重训建议", "模型中心"],
            ["/api/tasks/run", "POST", "task.py", "kind", "任务信息", "启动后台任务", "任务中心/首页/模型中心"],
            ["/api/tasks/{id}/logs", "GET", "task.py", "task_id", "日志文本", "查看任务日志", "任务中心/首页"],
            ["/api/scheduled-tasks", "GET/POST", "task.py", "创建时 name/mode/run_time/highest", "任务列表或创建结果", "计划任务管理", "任务中心"],
            ["/api/model-gateway/chat/completions", "POST", "model_gateway/router.py", "messages、model、temperature", "OpenAI-compatible 响应", "Agent/外部本地模型调用"],
        ],
    )

    doc.add_heading("14. 数据库与存储设计", level=1)
    add_table(
        doc,
        ["存储对象", "位置/表", "用途", "关键字段"],
        [
            ["预测运行", "forecast_runs / forecast_results", "保存预测批次和小时级预测结果", "run_id、forecast_datetime、predicted_price、risk_level、raw_json"],
            ["模型版本", "model_registry、model_versions、model_metrics", "保存候选/Active 模型、artifact 路径和指标", "model_version、artifact_path、is_active、mae、rmse、peak_rmse"],
            ["预测追踪", "prediction_tracking", "未来预测和真实值回填后的误差闭环", "forecast_datetime、predicted_price、actual_price、abs_error、pct_error"],
            ["模型运维", "model_performance_daily、model_retrain_jobs、model_comparison_runs", "日级误差、重训任务和候选对比", "metric_date、mae、rmse、status、reason"],
            ["误差/策略记忆", "model_error_memory、model_strategy_memory", "偏差校正和策略推荐记忆", "hour、risk_level、bias_mean、sample_count、strategy"],
            ["任务运行", "task_runs、task_logs、analysis_runs", "Web/Celery/脚本任务状态和日志", "task_id、run_id、status、command_json、log_text"],
            ["AI 轨迹", "ai_traces、ai_tool_call_logs、ai_chat_feedback、ai_answer_feedback、ai_conversation_state", "问答 trace、工具调用、反馈和会话记忆", "trace_id、session_id、question、intent、tools_json、evidence_json"],
            ["报告", "report_runs、report_reviews、ai_report_review_runs", "报告生成、审核、发布记录", "report_id、status、file_path、reviewer、comment"],
            ["知识库", "kb_documents、kb_chunks", "本地 Markdown/政策文件切片索引", "doc_id、chunk_id、content、keywords_json、embedding_json"],
            ["电价政策", "pv_tariff_rules、pv_policy_files、pv_station_tariff_check、market_power_price_rules、southern_grid_tax_rules", "光伏/市电/南网税率等政策问答", "province、city、price、doc_number、summary、raw_json"],
            ["本地文件存储", "output、结果-3、自动化输出、model_artifacts", "无数据库时的降级数据源和产物目录", "Excel、JSON、DOCX、PNG、joblib"],
        ],
    )
    add_note(
        doc,
        "数据库兼容性风险",
        "项目同时存在 MySQL 风格 SQL（ON DUPLICATE KEY、DATE_SUB）和 PostgreSQL Alembic 迁移。enterprise docker 已引入 PostgreSQL，但部分 database_utils/model_ops SQL 仍偏 MySQL 方言，后续需要统一数据库方言。",
    )

    doc.add_heading("15. 配置文件与环境变量说明", level=1)
    add_table(
        doc,
        ["配置项", "位置", "作用", "默认/说明"],
        [
            ["DATABASE_URL", ".env.example、backend/app/core/config.py", "PostgreSQL 主库连接", "为空时主库不可用；DATABASE_PRIMARY 默认 postgresql"],
            ["DB_HOST/DB_PORT/DB_NAME/DB_USER/DB_PASSWORD", ".env.example、llm_service_config.yaml、database_utils.py", "旧版 MySQL 兼容连接", "未配置 DB_PASSWORD 时降级本地文件模式"],
            ["REDIS_URL", ".env.example、workers/celery_app.py", "Celery broker/backend", "redis://localhost:6379/0"],
            ["PJM_SUBSCRIPTION_KEY", ".env.example、fetch_power_market_data.py", "PJM Data Miner 2 API Key", "可为空；自动尝试从 Data Miner 2 settings.json 发现公开 key"],
            ["PJM_NODE_ID/PJM_REGION", ".env.example", "PJM 节点和区域", "默认 DOM / 34964545"],
            ["WEATHER_LAT/WEATHER_LON", ".env.example", "Open-Meteo 天气点", "Richmond, Virginia"],
            ["LLM_PROVIDER/LLM_BASE_URL/LLM_MODEL", ".env.example、llm_service_config.yaml", "报告生成模型配置", "默认 ollama、http://localhost:11434、qwen3:4b"],
            ["LOCAL_LLM_BASE_URL/LOCAL_LLM_MODEL", ".env.example、backend/model_gateway/config.py", "OpenAI-compatible 本地模型网关", "默认 http://localhost:11434/v1、qwen3:4b"],
            ["DIFY_ENABLED/DIFY_API_KEY", ".env.example", "可选 Dify 接入", "默认关闭；当前 /api/ai/chat 已委托准确问答链路"],
            ["AUTH_REQUIRED", ".env.example、security.py", "是否强制登录", "默认 0，未传 Header 时使用 dev_admin/admin"],
            ["ACTIVE_MODEL_ARTIFACT_PATH", ".env.example、active_model_loader.py", "无数据库 Active 记录时指定本地 Active artifact", "为空时需先设置数据库 Active 模型"],
            ["BIAS_CORRECTION_*", ".env.example、bias_corrector.py", "偏差校正开关、样本门槛、最大校正幅度", "默认启用，最小样本 10，最大调整 10"],
        ],
    )

    doc.add_heading("16. 项目部署与运行方式", level=1)
    add_table(
        doc,
        ["方式", "命令/入口", "适用场景", "说明"],
        [
            ["本地完整流水线", "python main_daily_run.py --refresh-data", "每日刷新、训练/预测、报告、归档", "可能耗时较长，依赖外部 API 和本地模型"],
            ["快速预测", "python main_daily_run.py --fast-forecast", "已有 Active 模型时快速生成预测", "不重训，需 model_registry Active 或 ACTIVE_MODEL_ARTIFACT_PATH"],
            ["只生成报告", "python main_daily_run.py --skip-prediction", "复用已有预测结果生成 AI 日报", "前置校验正式预测表和 AI 依赖文件"],
            ["Web 后端", "python -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000", "本地 API 服务", "接口文档 /docs"],
            ["Web 前端", "cd frontend；npm run dev", "本地浏览器工作台", "默认 http://127.0.0.1:5173"],
            ["统一菜单", "run_project.bat", "Windows 日常使用", "聚合 Web、每日刷新、健康检查、模型优化、GUI 等入口"],
            ["桌面 GUI", "python -m ui.app", "本机操作员控制台", "PySide6 新版；旧版 gui_launcher.py 保留"],
            ["Docker 轻量", "docker-compose up", "backend + frontend", "基础 compose 文件存在路径中文乱码风险，需部署前复核挂载路径"],
            ["Docker 企业版", "docker-compose -f docker-compose.enterprise.yml up", "PostgreSQL + Redis + backend + worker + frontend", "推荐作为服务端部署基础，但需统一数据库方言"],
        ],
    )

    doc.add_heading("17. 当前项目完成度评估", level=1)
    add_table(
        doc,
        ["能力", "完成度", "证据", "说明"],
        [
            ["端到端预测流水线", "较高", "main_daily_run.py、services、prediction_engine、README_AI自动化说明.md", "支持数据刷新、预测、摘要、报告、归档"],
            ["模型工程化", "中高", "model_artifacts、model_ops、tests/test_model_artifacts.py", "artifact、registry、tracking、monitor 已具备；需统一 DB 和上线治理"],
            ["Web 产品化", "中高", "backend/app、frontend/src、tests/test_web_platform.py", "页面和接口较完整，部分页面仍保留 mock/演示降级"],
            ["AI 助手", "中高", "ai_assistant、langgraph_agent、model_gateway、tests/test_ai_assistant_accuracy.py", "工具优先、Trace、Guard 已有；RAG 向量化和真实权限仍需加强"],
            ["数据治理", "中", "data_dictionary_fields.csv、database_utils.py、migrations", "字段字典和同步机制存在；数据质量规则、血缘和方言统一不足"],
            ["权限与审计", "中", "security.py、audit_repository、0002_rag_rbac_audit.py", "Header/dev fallback 可用；生产登录、密码/JWT 未完整实现"],
            ["部署运维", "中", "docker-compose.enterprise.yml、health_check、smoke_test", "有 Docker、健康检查、自检；需修复中文路径编码、密钥和监控"],
            ["文档与测试", "中高", "README、多份方案文档、tests", "说明较多；终端显示存在编码问题，部分中文文件名/字符串需规范"],
        ],
    )

    doc.add_heading("18. 存在问题与风险点", level=1)
    add_table(
        doc,
        ["风险点", "影响", "证据/位置", "建议"],
        [
            ["中文编码显示异常", "README、Docker 挂载、前端文案、Python 字符串在终端中出现乱码，影响维护和部署", "README_WEB平台说明.md、docker-compose.yml、多个源码输出", "统一 UTF-8 编码和编辑器配置，逐步修复乱码字符串"],
            ["数据库方言混用", "PostgreSQL 企业版部署可能遇到 MySQL SQL 语法不兼容", "database_utils.py、model_ops SQL、migrations/versions PostgreSQL", "抽象方言层或全面迁移到 SQLAlchemy Core/PostgreSQL 语法"],
            ["旧版 AI 逻辑存在不可达分支", "维护者可能误判 Dify/旧工具链仍在主路径", "backend/app/ai/assistant_service.py answer_chat 开头直接 return answer_chat_accurate", "删除或明确标注兼容代码，统一问答入口"],
            ["认证仍偏演示模式", "AUTH_REQUIRED=0 时默认 dev_admin/admin，生产安全风险较高", "backend/app/core/security.py、.env.example", "引入 JWT/会话、密码哈希、用户表、前端登录页和权限 UI"],
            ["RAG 未发现真实向量生成", "kb_chunks 有 embedding_json，但检索更像关键词/JSONB 召回，复杂语义召回不足", "rag_service.py、knowledge_repository.py、0002_rag_rbac_audit.py", "接入 embedding 模型和 pgvector 索引，保存向量并支持混合召回"],
            ["模型泄漏风险仍需持续审计", "电价时序特征复杂，未来信息泄漏会显著抬高离线指标", "LEAKAGE_BANNED_RAW_COLS、leakage_check_report、tests/test_leakage_filter.py", "把泄漏检查固化到 CI，增加时序切分和特征生成单测"],
            ["LLM 调用超时/成本控制有限", "本地模型不可用时体验降级；远程模型成本策略未完善", "llm_client、ai_assistant/core/llm_client.py", "增加 token 估算、缓存、重试预算、模型可用性监控和异步流式输出"],
            ["文件路径中文和空格较多", "Docker/Linux/CI/跨环境部署容易失败", "项目目录、结果目录、compose 挂载", "核心运行目录提供英文别名或配置化路径映射"],
            ["部分前端页面仍有 mock 降级", "真实数据不可用时容易误把演示数据当生产数据", "frontend/src/mock、DataStateBanner", "强化 mockFallback 标识，生产环境禁用 mock 或显著提示"],
            ["模型上线流程缺少人工审批闭环", "候选模型可能因单一指标改善而上线，缺少业务验收", "promotion_manager、compare_and_promote", "增加模型审批状态、回滚演练、冠军/挑战者看板和线上监控"],
        ],
    )

    doc.add_heading("19. 优化建议与后续开发方案", level=1)
    doc.add_heading("19.1 短期可立即优化", level=2)
    add_bullets(
        doc,
        [
            "统一 UTF-8 编码，优先修复 README、docker-compose、前端菜单、关键错误信息中的乱码。",
            "在 .env.example 中增加生产安全提示，并在 AUTH_REQUIRED=0 时让前端和后端健康页显示“演示鉴权模式”。",
            "补充一份真实运行 Runbook：首次安装、数据刷新、快速预测、重训、报告生成、常见错误处理。",
            "把 /api/model-gateway/health 扩展为实际模型可用性检查，而不只是返回配置。",
            "对所有 FileResponse 下载和表导出增加权限依赖与审计记录。",
            "将旧版 answer_chat 不可达代码移入 deprecated 模块或删除，降低维护歧义。",
        ],
    )
    doc.add_heading("19.2 中期需要重构优化", level=2)
    add_bullets(
        doc,
        [
            "统一数据库方言：优先以 PostgreSQL + SQLAlchemy/Alembic 为主线，替换 DATE_SUB、ON DUPLICATE KEY 等 MySQL 写法。",
            "把 v4_fix1 中仍集中的训练、特征和报表逻辑继续拆分为可测试模块，减少 legacy_engine 动态加载依赖。",
            "将特征工程做成可版本化 Feature Pipeline，保存特征 schema、生成时间和数据窗口，支撑可复现实验。",
            "完善模型注册审批：candidate → validated → approved → active，全链路保留审批人、指标、数据窗口和回滚策略。",
            "把任务运行统一到 Celery/Redis，减少 Web 请求直接启动本地脚本和内存任务管理的双轨复杂度。",
            "将知识库检索升级为关键词 + embedding + rerank 混合召回，并将引用片段带入 LLM Prompt。",
        ],
    )
    doc.add_heading("19.3 长期产品化/企业级落地优化", level=2)
    add_bullets(
        doc,
        [
            "引入正式用户体系：登录、JWT、密码哈希、角色管理、接口级权限、审计查询和操作留痕。",
            "建设 MLOps：模型实验追踪、数据漂移检测、线上/离线指标对齐、自动回滚和模型卡审批。",
            "建设数据治理：数据血缘、字段字典版本、数据质量规则、异常告警、外部 API SLA 监控。",
            "建设 AI 治理：Prompt 版本库、回答质量评测集、工具调用回放、敏感信息过滤、成本预算和多模型路由。",
            "建设部署体系：CI/CD、容器镜像、数据库迁移流水线、备份恢复、可观测性指标和集中日志。",
            "扩展业务域：更多市场节点、更多天气/新能源数据源、合同/库存/储能约束、交易策略仿真和收益回测。",
        ],
    )

    doc.add_heading("20. 总结", level=1)
    doc.add_paragraph(
        "该项目已经从单体预测脚本演进为包含数据采集、机器学习预测、模型运维、AI 报告、知识库问答、Web 工作台和桌面控制台的综合平台。"
        "从代码结构看，当前最有价值的主线是“PJM/DOM 数据 → 特征工程 → 高峰尖刺增强模型 → 正式未来 24 小时预测 → AI 解读与报告 → Web/GUI 展示”。"
        "项目具备演示、内部试用和继续工程化的基础，但若要企业级落地，需要优先解决编码规范、数据库方言统一、生产鉴权、RAG 向量化、模型审批和部署运维治理问题。"
    )
    add_note(
        doc,
        "未能完全确认的内容",
        "未执行完整训练和外部 API 拉取，因此未确认当前 .env 私有配置、数据库实际连通状态、Ollama 本机模型实际可用状态、最新线上数据质量和完整模型效果；这些均需要在目标运行环境中通过健康检查、smoke test 和一次完整流水线验证。",
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
