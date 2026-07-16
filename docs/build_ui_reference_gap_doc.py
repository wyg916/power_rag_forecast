from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
SCREENSHOT_ROOT = next(p for p in PROJECT.iterdir() if p.is_dir() and p.name.startswith("20260621-"))
REFERENCE_DIR = SCREENSHOT_ROOT / "20260621-UI界面优化参考样式图"
CURRENT_CONTACT_DIR = PROJECT / "ui_gap_assets"
OUTPUT = PROJECT / "docs" / "智能运营分析项目_UI优化参考图与当前功能差异分析说明文档_v20260621.docx"


reference_images = [
    "全局框架层.png",
    "首页-总览驾驶舱.png",
    "数据中心 - 数据总览主页面.png",
    "数据中心 - 数据质量-数据目录组合页.png",
    "预测中心 - 24小时预测主页面.png",
    "预测中心 - 历史对比页面.png",
    "预测中心 - 峰谷分析-模型评估组合页.png",
    "策略中心 - 总览主页面.png",
    "策略中心 - 低价窗口-储能策略页面.png",
    "策略中心 - 人工复核页面.png",
    "AI 助手 - 智能问答主页面.png",
    "报告中心 - 报告列表-预览主页面.png",
    "报告中心 - 审核发布页面.png",
    "模型中心 - Active-Candidate -误差趋势一体化页面.png",
    "知识库 - 文档-索引-检索一体化页面.png",
    "任务中心 + 系统设置（后台运营组合页）.png",
    "系统设置主页面.png",
]


analysis_rows = [
    {
        "板块": "全局框架层",
        "图片": "全局框架层.png",
        "参考图功能": "统一左侧导航、顶部项目/时间/地区/模型/系统状态栏、全局搜索、通知、用户菜单、页面筛选条、KPI 卡片、图表区、任务/告警/快捷入口。",
        "当前已实现": "已有 BasicLayout、HeaderBar、Sidebar、PageContainer；已有 10 个一级业务页面；已有用户登录态、模型版本/数据时间等头部概念；已有 MetricCard、SectionCard、TableCard、DetailDrawer、状态/空态组件。",
        "主要差距": "导航分组、顶部筛选、全局搜索、通知中心、收藏/新建任务、右侧运营状态区未统一；当前截图中页面缩放和卡片排列较松散，浏览器通知会遮挡内容。",
        "数据/API支撑": "settingsHealth、tasksHealth、dbHealth、knowledgeHealth、authMe、securityMe 可支撑状态；全局搜索与通知需要新增轻量聚合或前端索引。",
        "落地判断": "可直接落地。主要是前端 AppShell/布局层改造，少量新增全局搜索/通知接口。",
        "优先级": "P0",
    },
    {
        "板块": "首页/总览驾驶舱",
        "图片": "首页-总览驾驶舱.png",
        "参考图功能": "首页集中展示供需风险、预测可信度、策略预计收益、报告待审、任务提醒、数据健康；主图展示 24 小时供需风险；右侧展示 AI 建议、策略执行、任务提醒、快捷入口。",
        "当前已实现": "DashboardPage 已有指标卡、风险列表、24 小时电价趋势、储能建议、最新报告、任务日志、快捷操作；dashboardApi 汇总 forecastLatest、tasks、dataFreshness、tasksHealth、knowledgeHealth、dbHealth、modelBacktestSummary。",
        "主要差距": "当前首页仍有分散 Tab，参考图要求首屏完成核心运营驾驶舱；供需风险指数、预测可信度、策略执行摘要、报告审核摘要、数据健康环图需要统一成主页面组件。",
        "数据/API支撑": "现有接口可支撑大部分指标；供需风险指数、策略执行摘要、报告审核计数可由现有预测/策略/报告/任务数据聚合，必要时新增 dashboard summary 聚合接口。",
        "落地判断": "可落地。建议先做前端聚合，后续再把聚合下沉为 /api/dashboard/summary。",
        "优先级": "P0",
    },
    {
        "板块": "数据中心/数据总览",
        "图片": "数据中心 - 数据总览主页面.png",
        "参考图功能": "数据域、时间范围、源状态和同步状态筛选；数据源数量、同步任务、异常表、完整率、昨日更新量；采集-清洗-校验-入库-特征/服务流程；同步记录、健康摘要、快捷操作和异常提醒。",
        "当前已实现": "DataCenterPage 已有数据接入、数据质量、数据目录、数据源表、导入导出页面；dataApi 接入 dataStatus、dataQuality、databaseTables、importExportRecords、dataCatalog、dataFreshness；支持核心数据入库、全部同步、表预览、导出、只读 SQL。",
        "主要差距": "参考图把总览、流程、同步记录、健康摘要集中到一个页面；当前页面分散在多个 Tab，流程图、昨日更新量、异常提醒聚类、快捷操作面板需要补齐。",
        "数据/API支撑": "P1 数据目录、freshness、quality、tables 已有；昨日更新量和流程成功率需要 dataQuality/importExport 进一步聚合。",
        "落地判断": "可落地。主要是页面重排和聚合字段补齐。",
        "优先级": "P1",
    },
    {
        "板块": "数据中心/数据质量与目录",
        "图片": "数据中心 - 数据质量-数据目录组合页.png",
        "参考图功能": "质量 KPI、近 7 天质量趋势、异常明细表、批量处理、数据目录卡片、选中表详情、字段预览、血缘关系、可追溯标签、数据源状态条。",
        "当前已实现": "已实现质量概览、质量明细、目录列表、字段抽屉、表预览、只读 SQL、导入导出记录；corePredictionTables 已覆盖 raw_market、raw_load、raw_weather、forecast_results、raw_renewable。",
        "主要差距": "重复率、平均延迟、批量处理、字段血缘、责任人、用途说明、目录卡片详情尚不完整；raw_renewable 需要明确标注可用性或缺口。",
        "数据/API支撑": "dataQuality、dataCatalog、dataFields、dataFreshness、databaseTables 支撑基础展示；责任人、血缘和处理状态需要扩展 catalog 元数据。",
        "落地判断": "可落地但需补元数据。UI 可先落地，责任人/血缘/处理闭环作为接口增强。",
        "优先级": "P1",
    },
    {
        "板块": "预测中心/24小时预测",
        "图片": "预测中心 - 24小时预测主页面.png",
        "参考图功能": "预测日期/区域/模型/数据源筛选；最高/最低/均价/峰谷价差/可信度；24 小时曲线、置信区间、低价窗口、高风险时段、峰值点；明细表；右侧策略洞察、风险与可信度说明、模型状态、数据健康。",
        "当前已实现": "ForecastCenterPage 已有 24h、历史、明细、峰谷、模型评估 Tab；支持 runForecast；已有 KPI、PriceCurveChart、预测明细、单小时解释、策略洞察。",
        "主要差距": "置信区间、数据健康侧栏、风险/可信度说明、低价/高风险标注需要更清晰；导出结果按钮、生成策略动作与策略中心联动需要补强。",
        "数据/API支撑": "forecastLatest、predictionLatest、forecastRun、modelBacktestSummary、featureSchema、leakageCheck 已有；置信区间可先由预测误差规则估算，长期应由模型输出。",
        "落地判断": "可直接落地，置信区间属于增强项。",
        "优先级": "P1",
    },
    {
        "板块": "预测中心/历史对比",
        "图片": "预测中心 - 历史对比页面.png",
        "参考图功能": "最新预测、上一批次、近 30 天历史均值对比；关键变化洞察、变化原因说明、AI 建议摘要、数据健康状态。",
        "当前已实现": "当前有历史对比 Tab 和曲线；后端存在 /api/market/history；forecastApi 已生成 history 数据，但部分 actual/history 仍可能由当前预测衍生。",
        "主要差距": "真实上一批次预测、历史均值、变化率、变化归因和差值表未形成完整闭环；需要保存 forecast_run 历史并暴露对比接口。",
        "数据/API支撑": "市场历史接口和预测结果表已有基础；需要预测运行批次、上一批次查询、历史均值聚合。",
        "落地判断": "可落地但需接口聚合。优先做读取真实历史，再做 AI 归因。",
        "优先级": "P2",
    },
    {
        "板块": "预测中心/峰谷分析与模型评估",
        "图片": "预测中心 - 峰谷分析-模型评估组合页.png",
        "参考图功能": "峰谷价差指数、峰谷时段、业务解释；模型评估摘要展示 Baseline、高峰、尖峰、极端天气、Feature Schema、Leakage Gate、Backtest、训练与回测状态。",
        "当前已实现": "P2 已产出 train/validation/test 数据集、feature_schema、leakage_check_result、metrics、backtest_report；ForecastCenterPage 已有 peak 与 model Tab，能展示 baseline、schema、leakage、对比表。",
        "主要差距": "参考图要求峰谷分析和 P2 工程可信度组合展示；当前分散在两个 Tab；训练任务 ID、完整回测曲线、评估结论、模型准入状态需要统一。",
        "数据/API支撑": "modelBacktestSummary、modelFeatureSchema、modelLeakageCheck、retrainSuggestion、tasks 可支撑；P2 test persistence_24h MAE≈34.294、RMSE≈93.671、MAPE≈35.302、R2≈0.3287，可作为展示基线。",
        "落地判断": "可直接落地。注意不要把基线指标包装成真实精度提升结论。",
        "优先级": "P1",
    },
    {
        "板块": "策略中心/总览",
        "图片": "策略中心 - 总览主页面.png",
        "参考图功能": "策略生成结论、高价风险时段、低价采购窗口、预计收益、人工复核数；策略时间轴、建议说明、执行优先级、风险提示、关键操作建议、风险来源分布、收益对比。",
        "当前已实现": "StrategyCenterPage 有高价风险、低价窗口、储能策略、人工复核；支持 generateStrategy、strategyLatest、anomalyLatest、strategyConfig、saveReview；已有时间轴和策略建议。",
        "主要差距": "总览态信息架构、收益对比、风险来源分布、执行状态和优先级评分不完整；当前低价/储能/复核较像分页面示例。",
        "数据/API支撑": "strategyLatest、anomalyLatest 可支撑基础；收益基准、执行状态、风险来源分类需要扩展策略结果结构。",
        "落地判断": "可落地但需策略结果字段增强。",
        "优先级": "P2",
    },
    {
        "板块": "策略中心/低价窗口与储能策略",
        "图片": "策略中心 - 低价窗口-储能策略页面.png",
        "参考图功能": "低价采购窗口列表、储能充放电计划、SOC 趋势、小时建议表、选中时段详情、收益测算、风险说明、执行条件、人工复核建议、加入执行清单/标记已执行。",
        "当前已实现": "当前低价窗口和储能策略页面共用策略时间表，能展示时段、动作、功率、收益并打开详情；有策略配置入口。",
        "主要差距": "SOC 曲线、充放电约束、执行清单、已执行状态、选中详情面板、人工复核闭环未完整实现。",
        "数据/API支撑": "可由 forecast/strategy 数据生成基础建议；SOC 和执行清单需要引入储能参数、状态持久化和执行记录表。",
        "落地判断": "中等可行。先做 UI 和静态/派生数据，真实执行闭环需要后端补表。",
        "优先级": "P3",
    },
    {
        "板块": "策略中心/人工复核",
        "图片": "策略中心 - 人工复核页面.png",
        "参考图功能": "待复核/已通过/已驳回/紧急风险/平均处理时长统计；筛选、批量通过/驳回/标记复核；复核列表；右侧详情、风险摘要、证据来源、相关数据、备注、处理意见和审计时间线。",
        "当前已实现": "当前有人工复核表和标记复核动作；后端 /api/strategy/reviews 可保存复核记录。",
        "主要差距": "状态流转、批量操作、驳回/通过/复核三态、证据链、责任人、审核日志和 SLA 统计未完整实现。",
        "数据/API支撑": "现有 reviews 接口只够最小保存；需扩展 review schema、audit logs 和风险证据绑定。",
        "落地判断": "需要后端业务模型补齐。UI 可先做，但不能假装已具备真实审批闭环。",
        "优先级": "P3",
    },
    {
        "板块": "AI助手/智能问答",
        "图片": "AI 助手 - 智能问答主页面.png",
        "参考图功能": "左侧会话与常见问题；中间结构化专业回答，包含结论、数据依据、原因解释、业务建议、风险提示；顶部回答类型切换；底部输入支持附件/引用数据/截图图表；右侧数据依据、关键指标、知识引用和建议动作。",
        "当前已实现": "AssistantPage 已有会话历史、常用问题、专业/通俗等回答风格、模型 provider、Trace 开发者模式、RAG 证据、最近知识库/模型运行；后端 answer_chat_accurate 已包含意图、工具、RAG、LLM、Guard、Trace。",
        "主要差距": "普通用户视图需要隐藏 Trace/工具日志，改为右侧依据栏；回答结构需更接近业务卡片；附件和截图输入尚非核心能力；离线 mock fallback 必须明确标记。",
        "数据/API支撑": "ai/chat、chat/sessions、ai/traces、knowledgeSearch、tools 和 answer_guard 已有；上传附件需新增。",
        "落地判断": "可直接落地。重点是前端信息架构和回答渲染。",
        "优先级": "P1",
    },
    {
        "板块": "报告中心/报告列表与预览",
        "图片": "报告中心 - 报告列表-预览主页面.png",
        "参考图功能": "报告类型/状态/日期筛选、搜索、生成报告、导出；生成/待审/发布/驳回指标；报告列表、预览区、数据概览、趋势图、风险提醒、摘要、快捷操作、发布记录。",
        "当前已实现": "ReportCenterPage 有日报/周报/审核/发布 Tab、报告列表、报告预览、下载、详情、重新生成、审核通过/驳回/发布；reportApi 接入 latest、detail、download、generate、approve、reject、publish、reviews。",
        "主要差距": "真实多报告列表、分页、复制链接、归档/删除、Excel 导出、发布记录统计不完整；当前更偏 latest report 预览。",
        "数据/API支撑": "当前报告接口能支撑单报告和审核动作；完整列表/归档/删除/链接需要扩展 report repository/API。",
        "落地判断": "中高可行。列表管理需要后端补齐。",
        "优先级": "P2",
    },
    {
        "板块": "报告中心/审核发布",
        "图片": "报告中心 - 审核发布页面.png",
        "参考图功能": "报告版本列表、审核预览区、审核意见、通过/驳回/重新生成/发布归档、发布归档流程、时间线、审核记录、版本信息、操作日志。",
        "当前已实现": "已有审核 Tab、审核操作、发布操作、时间线和报告详情；后端已有 approve/reject/publish/reviews。",
        "主要差距": "版本列表、版本差异、审核人协作状态、操作日志面板、发布归档流程节点不完整；审核意见未形成强约束。",
        "数据/API支撑": "已有 reviews 基础；版本和日志需要扩展 report metadata 和 audit logs。",
        "落地判断": "可落地但需版本模型增强。",
        "优先级": "P2",
    },
    {
        "板块": "模型中心/Active-Candidate-误差趋势",
        "图片": "模型中心 - Active-Candidate -误差趋势一体化页面.png",
        "参考图功能": "模型生命周期总览、Active/Candidate、MAE/RMSE/高峰误差、最近训练；效果对比图、误差趋势、版本表、候选准入规则、训练状态、评估摘要和回滚操作。",
        "当前已实现": "ModelCenterPage 有 Active、Candidate、误差、回滚 Tab；支持启动重训任务；展示模型对比、误差趋势、详情、训练状态；backend 提供 metrics/errors/retrain-suggestion/backtest/feature-schema/leakage-check。",
        "主要差距": "设为 Active、回滚等按钮当前更像前端确认动作，缺少明确后端 mutation；Candidate 注册、准入规则自动判断、回滚审计和训练日志详情需补齐。",
        "数据/API支撑": "P2 artifacts、model metrics、task runtime 可支撑展示；模型生命周期写操作需新增 activate/rollback/register endpoints。",
        "落地判断": "展示可直接落地，生命周期闭环需后端增强。",
        "优先级": "P2",
    },
    {
        "板块": "知识库/文档-索引-检索",
        "图片": "知识库 - 文档-索引-检索一体化页面.png",
        "参考图功能": "文档管理、索引管理、知识检索、QA 测试；上传、重建索引、刷新 Embedding、批量校验、导出；文档表、索引向量化状态、RAG 健康、检索测试、Top5 结果、AI 整理答案、索引流程。",
        "当前已实现": "KnowledgeBasePage 有文档/索引/RAG/QA Tab、stats、health、search、rebuildIndex、embeddingRefresh、RAG 降级提示；P3 已验证 BGE embedding/reranker 和 health。",
        "主要差距": "上传文档、批量校验、真实 QA 通过率、文档列表分页、索引流程视觉化和导出结果尚需完善；参考图中的降级原因需要更醒目。",
        "数据/API支撑": "knowledge/stats、health、search、index-local、embedding-refresh 支撑核心；上传和 QA 验证需新增接口。",
        "落地判断": "可落地。健康与检索已具备，文档管理操作需增强。",
        "优先级": "P1",
    },
    {
        "板块": "任务中心/后台运营组合页",
        "图片": "任务中心 + 系统设置（后台运营组合页）.png",
        "参考图功能": "任务调度、运行日志、运行健康、失败重试；任务筛选、总数/成功/运行排队/失败超时/队列数 KPI、任务表、健康卡、运行日志、失败重试、趋势图、队列概览。",
        "当前已实现": "TaskCenterPage 已有任务调度、日志、运行态健康、失败重试；支持运行任务、取消、重试、分页日志、新建定时任务；P4 已完成任务生命周期、幂等、超时、重试、取消、队列健康。",
        "主要差距": "当前截图为空记录时可读性弱；参考图要求总览态、趋势图、队列概览和失败聚类；运行日志/失败重试应做右侧运营面板。",
        "数据/API支撑": "tasks、tasksHealth、taskLogs、scheduledTasks、run/cancel/retry 已有；趋势和队列分布需要从 task_runs 统计。",
        "落地判断": "可直接落地。建议补聚合统计。",
        "优先级": "P1",
    },
    {
        "板块": "系统设置",
        "图片": "系统设置主页面.png",
        "参考图功能": "系统状态、用户管理、角色配置、参数配置、接口配置；在线用户、角色数、接口数、系统健康度；运行状态摘要、健康明细、用户表、权限矩阵、参数、数据库/模型/Web 服务配置。",
        "当前已实现": "SettingsPage 有系统状态、用户权限、角色配置、参数配置、接口配置；settingsApi 接入 settingsConfig、settingsHealth、permissions、securityMe、dbHealth、tasksHealth、knowledgeHealth、localModelStatus；UserManagementPage 存在。",
        "主要差距": "参考图中的在线用户数、接口数、连接测试、健康明细大表和多个中间件（Kafka/Airflow/MinIO/Prometheus/ELK）不应照抄；应映射到项目真实服务：PostgreSQL、Redis、Celery、RAG/BGE、LLM、日志/审计。",
        "数据/API支撑": "已有健康和配置接口；在线用户、接口数量和连接测试可增补；敏感配置需脱敏。",
        "落地判断": "可落地但需按真实技术栈改造，不要伪造未部署组件。",
        "优先级": "P2",
    },
]


current_status_rows = [
    ("全局框架", "已具备", "BasicLayout、HeaderBar、Sidebar、PageContainer、10 个一级菜单已存在。", "需要统一顶部筛选、导航分组、快捷入口、全局搜索和通知样式。"),
    ("数据可信", "已具备", "P1 数据目录、字段映射、freshness、只读 SQL 安全、data quality、catalog 接口已存在。", "目录元数据、责任人、血缘、异常处理闭环仍需补齐。"),
    ("预测与 P2 工程", "已具备基础", "P2 产出 62 特征、schema、leakage、backtest，Forecast 页面已接 forecast/model 接口。", "模型精度提升、预测置信区间、真实历史批次对比和 Candidate 晋升闭环需继续。"),
    ("策略中心", "部分具备", "strategyLatest、anomalyLatest、config、review 保存、前端策略页面已存在。", "执行清单、SOC、复核状态流、证据链和收益对比仍需增强。"),
    ("AI 助手", "已具备", "answer_chat_accurate、工具路由、RAG、LLM Router、Guard、Trace、前端开发者模式已存在。", "普通用户态要改成业务结构化答案，隐藏默认 Trace，明确 fallback 状态。"),
    ("报告中心", "部分具备", "生成、下载、审核、驳回、发布、reviews 接口和前端页面已存在。", "多报告列表、版本管理、归档、发布日志和 Excel/链接操作需增强。"),
    ("模型中心", "部分具备", "metrics/errors/backtest/feature/leakage 接口、重训任务入口、Active/Candidate 展示已存在。", "设为 Active、回滚、注册 Candidate 等真实写操作需后端闭环。"),
    ("知识库/RAG", "已具备", "stats、health、search、index-local、embedding-refresh；P3 验证 BGE 和 reranker。", "上传、批量 QA、文档分页和流程可视化需补齐。"),
    ("任务中心", "已具备", "P4 完成 task lifecycle、幂等、超时、重试、取消、日志分页、队列健康。", "前端总览趋势、失败聚类、队列分布和空态体验需优化。"),
    ("系统设置", "部分具备", "settings config/health、permissions、securityMe、用户管理、健康接口已存在。", "在线用户、接口配置测试、真实服务健康明细和脱敏展示需完善。"),
]


implementation_phases = [
    ("UI-0：统一页面骨架", "改造 AppShell、HeaderBar、Sidebar、PageContainer、PageTabs、MetricCard、SectionCard、TableCard、DetailDrawer、状态色和空态。目标是先解决错位、溢出、大片空白和缩放异常。"),
    ("UI-1：高价值首屏", "优先落地全局框架、首页驾驶舱、数据中心总览、预测 24 小时主页面、任务中心总览。它们已有真实接口，风险最小，收益最大。"),
    ("UI-2：可信与闭环", "重排数据质量/目录、预测峰谷/模型评估、AI 助手、知识库。重点展示数据源、schema、leakage、RAG health、Trace/fallback，而不是只做视觉美化。"),
    ("UI-3：审批与生命周期", "增强策略人工复核、报告审核发布、模型 Candidate/Active/回滚。该阶段涉及业务状态机，必须同步扩展后端 schema/API。"),
    ("UI-4：业务执行增强", "补储能 SOC、执行清单、策略收益回填、真实上一批次预测对比、任务 SLA、线上监控告警等生产化能力。"),
]


gpt_constraints = [
    "不要推翻当前 React 18 + TypeScript + Vite + Ant Design + ECharts 技术栈，优先复用现有组件。",
    "不要为了参考图好看而伪造 Kafka、Airflow、MinIO、Prometheus、ELK 等项目未部署组件；系统设置页应展示真实服务。",
    "所有 AI 数据型回答必须基于工具、只读 SQL 或 RAG 证据；不能编造价格、日期、时段和交易结论。",
    "只读 SQL 必须继续禁止写操作、多语句、敏感表、系统表和危险函数，不能为了页面查询方便放宽 P1 安全约束。",
    "P2 模型指标只能作为工程化和 baseline 依据，不能把 persistence baseline 包装成模型精度已大幅提升。",
    "RAG 页面必须显式展示 provider、模型路径、embedding_dim、chunk_count、fallback 状态，不能隐藏降级风险。",
    "任务中心必须保持 P4 的幂等、超时、重试、取消、日志分页和队列健康约束。",
    "报告审核、策略复核、模型切换、配置修改等操作必须写审计或至少预留 audit log 字段。",
    "生产态要能关闭 mock fallback 或明确展示 fallback，不允许用户误判为真实生产数据。",
    "UI 改造应先修布局和信息架构，再补复杂业务闭环；每个阶段要有 npm build 和关键接口回归。",
]


def set_doc_defaults(doc: Document) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Inches(11.69)
    sec.page_height = Inches(8.27)
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.5)
    sec.right_margin = Inches(0.5)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].font.color.rgb = RGBColor(31, 41, 55)
    for name, size, color in [
        ("Title", 21, "0F172A"),
        ("Heading 1", 15, "0F766E"),
        ("Heading 2", 12, "1E3A8A"),
        ("Heading 3", 10.5, "0F172A"),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.2)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def set_table_width(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def add_note(doc: Document, text: str, color: str = "475569") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(item)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.8)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, head in enumerate(headers):
        shade_cell(hdr[i], "E8F3F1")
        set_cell_text(hdr[i], head, bold=True, color="0F766E")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    set_table_width(table, widths)
    doc.add_paragraph()


def add_kv_table(doc: Document, pairs: list[tuple[str, str]], key_width: float = 1.35, value_width: float = 9.0) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for key, value in pairs:
        cells = table.add_row().cells
        shade_cell(cells[0], "F1F5F9")
        set_cell_text(cells[0], key, bold=True, color="334155")
        set_cell_text(cells[1], value)
    set_table_width(table, [key_width, value_width])
    doc.add_paragraph()


def add_image(doc: Document, path: Path, caption: str, width: float = 9.5) -> None:
    if not path.exists():
        add_note(doc, f"图片缺失：{path}", "B91C1C")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string("64748B")


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("智能运营分析项目\nUI优化参考图与当前功能差异分析说明文档")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(23)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0F172A")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("版本：v20260621 | 分析范围：参考样式图、当前项目截图、前后端代码、接口与 P1-P5 验收材料")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string("475569")

    add_note(doc, "文档目的：逐张理解 UI 参考图，结合当前真实项目功能和数据/API能力，说明已实现、需补充、可落地性与后续开发重点，便于发给 GPT 助手继续制定优化落地开发文档。")
    add_note(doc, "版式：compact_reference_guide。本文嵌入 17 张参考图，并在附录保留当前项目各板块截图联系表，便于对照。")
    doc.add_page_break()


def build_doc() -> None:
    doc = Document()
    set_doc_defaults(doc)
    add_cover(doc)

    doc.add_heading("1. 分析结论摘要", level=1)
    add_bullets(
        doc,
        [
            "本次参考图不是单纯视觉美化稿，而是一套运营型产品目标态：统一框架、业务筛选、KPI、图表、明细表、右侧洞察/状态/快捷操作、审计和复核闭环。",
            "当前项目已经具备真实工程基础：FastAPI、React/TypeScript、PostgreSQL、Redis/Celery、P1 数据可信、P2 回测产物、P3 RAG/BGE、P4 任务中心、P5 前端与 AI 体验增强。",
            "多数页面可以基于现有组件和接口落地；关键风险不在技术栈，而在是否区分真实数据、派生数据和 mock/fallback，以及是否补齐审批/复核/模型生命周期等后端状态机。",
            "建议先做全局布局和高价值首屏，随后补数据可信、预测可信、AI 结构化回答、任务总览，再推进报告审核、策略复核、模型晋升/回滚等业务闭环。",
        ],
    )

    doc.add_heading("2. 当前项目真实状态", level=1)
    add_table(
        doc,
        ["模块", "当前状态", "已具备能力", "主要差距"],
        [list(row) for row in current_status_rows],
        [1.35, 1.15, 4.2, 4.2],
    )

    doc.add_heading("3. 参考图整体 UI 设计语言", level=1)
    add_bullets(
        doc,
        [
            "框架：固定左侧导航 + 顶部全局状态栏 + 页面标题/筛选条 + 内容网格。导航按业务中心、智能应用、运营管理分组。",
            "视觉：浅灰背景、白色卡片、8px 左右圆角、青绿色作为主操作与健康色，蓝色表示模型/预测，橙色表示提醒，红色表示高风险。",
            "布局：首屏强调“可行动信息”，每页都包含 KPI、主图/主表、右侧解释/状态/快捷操作，避免单页只有大表或只有图。",
            "交互：筛选、刷新、导出、生成、详情、复核、重试、取消、发布、回滚、复制链接等动作都在上下文中出现，且需要状态可追溯。",
            "数据表达：大量使用 data source、模型版本、更新时间、风险等级、可信度、fallback/health 状态，说明目标是可信运营后台，而不是静态展示页。",
        ],
    )

    doc.add_heading("4. 逐图功能差异总表", level=1)
    add_table(
        doc,
        ["参考图", "当前已实现", "需补充/重构", "落地判断", "优先级"],
        [[row["板块"], row["当前已实现"], row["主要差距"], row["落地判断"], row["优先级"]] for row in analysis_rows],
        [1.6, 3.0, 3.1, 2.1, 0.7],
    )

    doc.add_heading("5. 逐图逐板块详细分析", level=1)
    for idx, row in enumerate(analysis_rows, start=1):
        doc.add_heading(f"5.{idx} {row['板块']}", level=2)
        add_image(doc, REFERENCE_DIR / row["图片"], f"参考图：{row['图片']}", width=9.7)
        add_kv_table(
            doc,
            [
                ("参考图功能", row["参考图功能"]),
                ("当前已实现", row["当前已实现"]),
                ("需要补充", row["主要差距"]),
                ("数据/API支撑", row["数据/API支撑"]),
                ("落地判断", row["落地判断"]),
                ("优先级", row["优先级"]),
            ],
        )

    doc.add_heading("6. 能否基于真实数据与业务场景落地", level=1)
    add_table(
        doc,
        ["落地分级", "包含页面/能力", "说明"],
        [
            [
                "可直接落地",
                "全局框架、首页大部分、数据中心基础、预测 24h、预测模型评估、AI 助手、知识库 health/search、任务中心、系统设置基础健康",
                "这些已有前端页面、组件和后端接口，优先做信息架构、样式统一、聚合展示和 fallback 显式化。",
            ],
            [
                "需接口聚合",
                "首页供需风险、数据中心同步趋势、预测历史批次对比、报告多列表、任务趋势/队列统计、系统接口统计",
                "数据多已存在，但需要后端或前端聚合成参考图所需的 KPI、趋势、列表和右侧摘要。",
            ],
            [
                "需业务状态机",
                "策略人工复核、储能执行清单、报告版本/归档、模型 Candidate 晋升与回滚",
                "不能只靠 UI 假装完成，需要数据库表、状态流转、权限校验和审计记录。",
            ],
            [
                "需数据源补齐",
                "raw_renewable、真实未来天气、SOC、实际执行反馈、模型在线/离线真实误差回填",
                "若数据源暂不可用，应在 UI 与 AI 中明确显示不可用或降级，不能用 mock 冒充生产数据。",
            ],
        ],
        [1.35, 3.5, 5.4],
    )

    doc.add_heading("7. 建议优化落地路线", level=1)
    add_table(
        doc,
        ["阶段", "工作内容"],
        [[name, desc] for name, desc in implementation_phases],
        [2.2, 8.0],
    )

    doc.add_heading("8. 给后续 GPT 助手的开发约束", level=1)
    add_bullets(doc, gpt_constraints)

    doc.add_heading("9. 附录：当前项目页面截图联系表", level=1)
    add_note(doc, "以下联系表来自当前项目截图目录，用于和参考图目标态对照。当前截图显示项目已有页面和真实交互基础，但存在卡片密度、右侧辅助区、空态、滚动区域、状态闭环和整体一致性不足。")
    for contact in sorted(CURRENT_CONTACT_DIR.glob("current_*_contact.png"), key=lambda p: p.name):
        doc.add_heading(contact.stem.replace("current_", "").replace("_contact", ""), level=2)
        add_image(doc, contact, f"当前项目截图联系表：{contact.name}", width=8.9)

    doc.add_heading("10. 附录：关键代码与接口依据", level=1)
    add_table(
        doc,
        ["依据类型", "关键文件/接口", "用途"],
        [
            ["前端路由与布局", "frontend/src/app/App.tsx；frontend/src/app/router.tsx；frontend/src/layout/*", "确认 10 个一级模块、统一布局和当前导航体系。"],
            ["前端组件", "frontend/src/components/cards/*；frontend/src/components/common/*；frontend/src/components/charts/*", "确认可复用卡片、表格、图表、状态组件。"],
            ["前端页面", "frontend/src/pages/dashboard/data/forecast/strategy/assistant/report/model/knowledge/task/settings", "确认各板块当前页面与交互。"],
            ["API 客户端", "frontend/src/api.ts；frontend/src/services/*.ts", "确认页面调用真实接口和 fallback 逻辑。"],
            ["后端路由", "backend/app/api/v1/endpoints/*.py", "确认数据、预测、策略、AI、报告、模型、知识库、任务、设置、用户接口。"],
            ["数据可信", "backend/app/services/data_trust_service.py", "确认数据目录、字段、freshness、只读 SQL 安全。"],
            ["AI 助手", "backend/app/ai_assistant/service.py；core/tool_router.py；core/answer_guard.py", "确认工具证据、RAG、LLM、Guard、Trace。"],
            ["任务中心", "backend/app/services/task_runtime.py；backend/app/repositories/task_repository.py", "确认任务生命周期、幂等、超时、重试、取消和日志分页。"],
            ["P2 产物", "output/p2/feature_schema.json；leakage_check_result.json；metrics.json；backtest_report.md", "确认模型工程化、schema、leakage、baseline/backtest。"],
            ["项目分析文档", "《智能运营分析项目》项目全面分析说明文档_v20260621.docx", "确认项目整体状态、P1-P5 基础和风险。"],
        ],
        [1.5, 4.4, 4.4],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()
