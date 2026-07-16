from __future__ import annotations

import csv
import json
import math
import re
import textwrap
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from matplotlib import font_manager
import openpyxl
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
DOCS = ROOT / "docs"
PHASE2 = DOCS / "phase2"
OUT = DOCS / "reports"
ASSETS = OUT / "_report_assets_v20260714"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT / "智能运营分析项目_当前状态技术架构与优化实施规划报告_v20260714.docx"
PDF_PATH = OUT / "智能运营分析项目_当前状态技术架构与优化实施规划报告_v20260714.pdf"
META_PATH = OUT / "智能运营分析项目_当前状态技术架构与优化实施规划报告_v20260714_metadata.json"

PROJECT = "智能运营分析项目"
REPORT_TITLE = "智能运营分析项目当前状态、技术架构与后续优化实施规划报告"
BASELINE_DATE = "2026年7月14日"
VERSION = "v20260714"

COLORS = {
    "navy": "17365D",
    "blue": "2F75B5",
    "light_blue": "D9EAF7",
    "pale_blue": "EEF5FB",
    "gold": "C7901E",
    "light_gold": "FFF2CC",
    "orange": "C65911",
    "light_orange": "FCE4D6",
    "red": "A61C2A",
    "light_red": "F4CCCC",
    "green": "548235",
    "light_green": "E2F0D9",
    "gray": "6B7280",
    "light_gray": "F2F4F7",
    "dark": "1F2937",
    "white": "FFFFFF",
}

STATUS_CN = {
    "completed_verified": "已验证完成",
    "completed_unverified": "已实现但未充分验证",
    "partially_completed": "部分完成",
    "implemented_not_integrated": "已实现但未接入",
    "planned_only": "仅规划",
    "not_started": "尚未开始",
    "blocked": "被阻塞",
    "duplicate": "重复实现",
    "deprecated": "已废弃",
    "unable_to_verify": "无法验证",
    "passed": "已验证完成",
    "partially_passed": "部分完成",
    "failed": "被阻塞",
    "skipped_for_safety": "被阻塞（安全停止）",
    "not_applicable": "尚未开始（当前不适用）",
}


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def fix_utf8_mojibake(value):
    if isinstance(value, str):
        try:
            candidate = value.encode("latin1").decode("utf-8")
            if sum("\u4e00" <= c <= "\u9fff" for c in candidate) > sum("\u4e00" <= c <= "\u9fff" for c in value):
                return candidate
        except Exception:
            pass
        return value
    if isinstance(value, list):
        return [fix_utf8_mojibake(x) for x in value]
    if isinstance(value, dict):
        return {fix_utf8_mojibake(k): fix_utf8_mojibake(v) for k, v in value.items()}
    return value


def fix_gbk_mojibake(value):
    if isinstance(value, str):
        try:
            candidate = value.encode("latin1").decode("gbk")
            if sum("\u4e00" <= c <= "\u9fff" for c in candidate) > sum("\u4e00" <= c <= "\u9fff" for c in value):
                return candidate
        except Exception:
            pass
        return value
    if isinstance(value, list):
        return [fix_gbk_mojibake(x) for x in value]
    if isinstance(value, tuple):
        return tuple(fix_gbk_mojibake(x) for x in value)
    if isinstance(value, dict):
        return {fix_gbk_mojibake(k): fix_gbk_mojibake(v) for k, v in value.items()}
    return value


SOURCES = {
    "phase2_baseline": PHASE2 / "智能运营分析项目_最终项目状态基线_20260714.md",
    "phase2_e2e": PHASE2 / "智能运营分析项目_阶段二_端到端链路验证报告_20260714.md",
    "phase2_root": PHASE2 / "智能运营分析项目_阶段二_根因分析与风险清单_20260714.md",
    "phase2_matrix": PHASE2 / "智能运营分析项目_阶段二_V001-V030验证矩阵_20260714.xlsx",
    "phase2_ai": PHASE2 / "智能运营分析项目_阶段二_AI与RAG评测报告_20260714.md",
    "phase2_migration": PHASE2 / "智能运营分析项目_阶段二_M001-M038状态迁移_20260714.json",
    "phase2_tasks": PHASE2 / "智能运营分析项目_阶段二_P0-P1-P2优化任务总表_20260714.json",
    "phase2_codex": PHASE2 / "智能运营分析项目_阶段二_后续Codex任务包_20260714.md",
    "phase2_log": PHASE2 / "智能运营分析项目_阶段二_命令与测试日志_20260714.md",
    "phase2_decisions": PHASE2 / "智能运营分析项目_阶段二_未解决与需用户决策事项_20260714.md",
    "phase2_evidence": PHASE2 / "智能运营分析项目_阶段二_证据更新索引_20260714.csv",
    "phase2_safety": PHASE2 / "智能运营分析项目_阶段二_执行计划与安全边界_20260714.md",
    "phase1_baseline": DOCS / "智能运营分析项目_阶段一_全量盘点与证据基线_20260714.md",
    "phase1_modules": DOCS / "智能运营分析项目_阶段一_模块状态清单_20260714.json",
    "phase1_evidence": DOCS / "智能运营分析项目_阶段一_证据索引_20260714.csv",
    "phase1_pending": DOCS / "智能运营分析项目_阶段一_待验证问题清单_20260714.md",
}

for key, path in SOURCES.items():
    if not path.exists():
        raise FileNotFoundError(f"缺少报告源文件：{key}: {path}")

SOURCE_TEXT = {key: read_text(path) for key, path in SOURCES.items() if path.suffix.lower() == ".md"}
MIGRATION = fix_utf8_mojibake(json.loads(read_text(SOURCES["phase2_migration"])))
TASKS_OBJ = json.loads(read_text(SOURCES["phase2_tasks"]))
PHASE1_MODULES = json.loads(read_text(SOURCES["phase1_modules"]))

with SOURCES["phase2_evidence"].open(encoding="utf-8-sig", newline="") as f:
    PHASE2_EVIDENCE = list(csv.DictReader(f))
with SOURCES["phase1_evidence"].open(encoding="utf-8-sig", newline="") as f:
    PHASE1_EVIDENCE = list(csv.DictReader(f))

wb = openpyxl.load_workbook(SOURCES["phase2_matrix"], read_only=True, data_only=True)
V_ROWS = [fix_gbk_mojibake(list(row)) for row in wb["验证矩阵"].iter_rows(values_only=True)]
V_HEADER = V_ROWS[0]
V_DATA = V_ROWS[1:31]
V_SUMMARY = [fix_gbk_mojibake(list(row)) for row in wb["状态汇总"].iter_rows(values_only=True)]


def source_note(text: str) -> str:
    return f"【来源：{text}】"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def set_run_fonts(run, east="宋体", latin="Arial", size=None, bold=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str, result_text=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = result_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def add_seq_caption(doc, label: str, text: str, seq: str, n: int):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{label} ")
    set_run_fonts(r, east="微软雅黑", size=9, bold=True, color=COLORS["navy"])
    field_run = add_field(p, f" SEQ {seq} \\* ARABIC ", str(n))
    set_run_fonts(field_run, east="微软雅黑", size=9, bold=True, color=COLORS["navy"])
    r2 = p.add_run(f"  {text}")
    set_run_fonts(r2, east="微软雅黑", size=9, color=COLORS["dark"])
    p.paragraph_format.space_after = Pt(6)
    set_keep_with_next(p, False)
    return p


def add_title(doc, title: str, level=1):
    p = doc.add_heading(title, level=level)
    set_keep_with_next(p)
    return p


def add_paragraph(doc, text="", *, style=None, bold_lead=None, source=None, color=None, keep=False):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_fonts(r1, bold=True, color=color or COLORS["dark"])
        r2 = p.add_run(text[len(bold_lead):])
        set_run_fonts(r2, color=color or COLORS["dark"])
    else:
        r = p.add_run(text)
        set_run_fonts(r, color=color or COLORS["dark"])
    if source:
        r = p.add_run(" " + source_note(source))
        set_run_fonts(r, east="宋体", size=8.5, color=COLORS["gray"])
    if keep:
        set_keep_together(p)
    return p


def add_bullets(doc, items, level=0, source=None):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        r = p.add_run(str(item))
        set_run_fonts(r)
        if source:
            rs = p.add_run(" " + source_note(source))
            set_run_fonts(rs, size=8.5, color=COLORS["gray"])


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(str(item))
        set_run_fonts(r)


TABLE_COUNTER = 0
FIG_COUNTER = 0


def add_table(doc, title: str, headers, rows, widths=None, compact=False, caption_source=None):
    global TABLE_COUNTER
    TABLE_COUNTER += 1
    add_seq_caption(doc, "表", title, "Table", TABLE_COUNTER)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, COLORS["navy"])
        set_cell_margins(cell, 70, 80, 70, 80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        set_run_fonts(r, east="微软雅黑", size=8 if compact else 8.5, bold=True, color=COLORS["white"])
        if widths:
            cell.width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        if ridx % 2 == 1:
            for c in cells:
                set_cell_shading(c, "F8FAFC")
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_margins(cell, 55 if compact else 75, 70, 55 if compact else 75, 70)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = "—" if value is None or value == "" else str(value)
            r = p.add_run(text)
            set_run_fonts(r, size=7.3 if compact else 8.3)
            if widths:
                cell.width = Cm(widths[i])
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    if caption_source:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(source_note(caption_source))
        set_run_fonts(r, size=8, color=COLORS["gray"])
    return table


def add_figure(doc, image_path: Path, title: str, source: str, width_cm=16.5):
    global FIG_COUNTER
    FIG_COUNTER += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run()
    r.add_picture(str(image_path), width=Cm(width_cm))
    add_seq_caption(doc, "图", title, "Figure", FIG_COUNTER)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = p2.add_run(source_note(source))
    set_run_fonts(rs, size=8, color=COLORS["gray"])


def find_chinese_font():
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyh.ttf"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


FONT_PATH = find_chinese_font()
if FONT_PATH:
    font_manager.fontManager.addfont(FONT_PATH)
    CHINESE_FONT = font_manager.FontProperties(fname=FONT_PATH).get_name()
else:
    CHINESE_FONT = "DejaVu Sans"
plt.rcParams["font.family"] = CHINESE_FONT
plt.rcParams["axes.unicode_minus"] = False


def save_flow(name, nodes, edges, title, *, xlim=(0, 10), ylim=(0, 6), note=None):
    fig, ax = plt.subplots(figsize=(12, 6.2), dpi=180)
    ax.set_xlim(*xlim)
    ax.set_ylim(*ylim)
    ax.axis("off")
    ax.set_title(title, fontsize=16, fontweight="bold", color="#17365D", pad=16)
    for key, spec in nodes.items():
        x, y, w, h, label = spec[:5]
        fill = spec[5] if len(spec) > 5 else "#EEF5FB"
        edge = spec[6] if len(spec) > 6 else "#2F75B5"
        patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.03,rounding_size=0.08",
                               linewidth=1.4, edgecolor=edge, facecolor=fill)
        ax.add_patch(patch)
        ax.text(x+w/2, y+h/2, label, ha="center", va="center", fontsize=10,
                color="#1F2937", wrap=True, linespacing=1.25)
    for a, b, label in edges:
        sa, sb = nodes[a], nodes[b]
        x1, y1, w1, h1 = sa[:4]
        x2, y2, w2, h2 = sb[:4]
        start = (x1+w1/2, y1+h1/2)
        end = (x2+w2/2, y2+h2/2)
        dx, dy = end[0]-start[0], end[1]-start[1]
        if abs(dx) >= abs(dy):
            start = (x1+w1 if dx > 0 else x1, y1+h1/2)
            end = (x2 if dx > 0 else x2+w2, y2+h2/2)
        else:
            start = (x1+w1/2, y1+h1 if dy > 0 else y1)
            end = (x2+w2/2, y2 if dy > 0 else y2+h2)
        arrow = FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=12,
                                linewidth=1.2, color="#6B7280", connectionstyle="arc3,rad=0.0")
        ax.add_patch(arrow)
        if label:
            ax.text((start[0]+end[0])/2, (start[1]+end[1])/2+0.12, label,
                    ha="center", va="bottom", fontsize=8, color="#6B7280")
    if note:
        ax.text(xlim[0]+0.1, ylim[0]+0.05, note, fontsize=8.5, color="#6B7280", va="bottom")
    fig.tight_layout()
    path = ASSETS / f"{name}.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def build_figures():
    figs = {}
    figs["architecture"] = save_flow(
        "01_current_architecture",
        {
            "user": (0.2, 4.5, 1.25, 0.7, "用户/项目人员", "#F2F4F7", "#6B7280"),
            "web": (2.0, 4.7, 1.55, 0.75, "React Web\n当前主入口"),
            "desktop": (2.0, 3.45, 1.55, 0.75, "PySide6\n辅助入口", "#FFF2CC", "#C7901E"),
            "api": (4.25, 4.7, 1.45, 0.75, "FastAPI\n业务/API层"),
            "service": (6.25, 4.7, 1.45, 0.75, "业务服务\n仓储/任务"),
            "pg": (8.25, 4.7, 1.35, 0.75, "PostgreSQL\n58表"),
            "celery": (6.25, 3.2, 1.45, 0.75, "Redis/Celery\n部分闭环", "#FFF2CC", "#C7901E"),
            "predict": (4.25, 2.0, 1.45, 0.75, "预测流水线\nlegacy+模块化", "#FCE4D6", "#C65911"),
            "artifact": (6.25, 2.0, 1.45, 0.75, "artifact/结果\n本地文件", "#FCE4D6", "#C65911"),
            "ai": (4.25, 0.65, 1.45, 0.75, "AI Assistant\n工具/RAG/Guard"),
            "provider": (6.25, 0.65, 1.45, 0.75, "DeepSeek/Ollama\nProvider"),
            "kb": (8.25, 0.65, 1.35, 0.75, "知识库\n当前6/6", "#FCE4D6", "#C65911"),
        },
        [("user","web","主入口"),("user","desktop","辅助"),("web","api","HTTP/SSE"),("desktop","service","调用"),
         ("api","service","服务"),("service","pg","读写"),("service","celery","任务"),("service","predict","预测"),
         ("predict","artifact","产物"),("api","ai","聊天"),("ai","provider","路由"),("ai","kb","检索")],
        "当前总体技术架构（基线日）", note="橙色表示存在事实源、身份或运行闭环风险；图中仅表示已确认的结构与边界。")

    figs["data_flow"] = save_flow(
        "02_data_pipeline",
        {"src":(0.2,2.4,1.25,0.8,"PJM /\nOpen-Meteo"),"collect":(1.9,2.4,1.25,0.8,"数据采集\n分页/重试"),
         "clean":(3.6,2.4,1.25,0.8,"清洗对齐\n时区/去重"),"store":(5.3,2.4,1.25,0.8,"PostgreSQL /\n本地文件"),
         "feature":(7.0,2.4,1.25,0.8,"特征工程\n时间切分"),"model":(8.7,2.4,1.1,0.8,"训练/预测\n产物")},
        [("src","collect","24小时小窗已实调"),("collect","clean",""),("clean","store",""),("store","feature",""),("feature","model","")],
        "数据采集、处理与建模流程", ylim=(0.8,4.6), note="外部最小采集已验证；正式增量、DST、429与迟到数据矩阵仍需补齐。")

    figs["current_forecast"] = save_flow(
        "03_current_forecast_chain",
        {"input":(0.2,3.2,1.3,0.75,"历史输入\n24行冻结样本"),"loader":(2.0,3.2,1.35,0.75,"Active加载器\nmodel_registry"),
         "art":(2.0,1.6,1.35,0.75,"候选artifact\n20260620", "#FFF2CC", "#C7901E"),"schema":(4.05,3.2,1.35,0.75,"170项特征\n静态匹配"),
         "predict":(6.1,3.2,1.35,0.75,"快速预测\n未安全执行", "#F4CCCC", "#A61C2A"),"sync":(8.15,3.2,1.35,0.75,"同步/展示\n覆盖风险", "#F4CCCC", "#A61C2A")},
        [("input","loader",""),("art","loader","未晋升"),("loader","schema",""),("schema","predict","静默补0风险"),("predict","sync","无统一run_id")],
        "当前预测主链：存在结构但闭环被阻塞", ylim=(0.6,5.2), note="候选artifact存在不等于当前Active可加载；阶段二按安全停止条件未执行写链。")

    figs["ideal_forecast"] = save_flow(
        "04_target_fast_forecast",
        {"run":(0.2,3.0,1.2,0.75,"创建run_id"),"active":(1.75,3.0,1.25,0.75,"唯一Active\n校验hash"),"contract":(3.4,3.0,1.25,0.75,"严格特征契约\n不匹配即失败"),
         "forecast":(5.05,3.0,1.25,0.75,"24小时预测\n身份齐全"),"db":(6.7,3.0,1.25,0.75,"事务upsert\n旧批次共存"),"verify":(8.35,3.0,1.25,0.75,"API/Web/AI\n同源验收"),
         "rollback":(5.05,1.4,1.25,0.75,"失败回滚\n无半成品", "#FFF2CC", "#C7901E")},
        [("run","active",""),("active","contract",""),("contract","forecast",""),("forecast","db",""),("db","verify",""),("forecast","rollback","异常")],
        "理想快速预测闭环：身份、契约、事务和回滚统一", ylim=(0.6,4.8), note="目标链路对应T001—T003，并要求在新隔离库完成双跑、幂等和故障回滚验证。")

    figs["fact_split"] = save_flow(
        "05_current_fact_split",
        {"local":(0.2,4.0,1.5,0.75,"本地artifact\n候选模型", "#FCE4D6", "#C65911"),"registry":(2.2,4.0,1.5,0.75,"model_registry\n快速加载=0", "#F4CCCC", "#A61C2A"),
         "versions":(4.2,4.0,1.5,0.75,"model_versions\nWeb=5条seed", "#FCE4D6", "#C65911"),"files":(6.2,4.0,1.5,0.75,"预测文件\n跨批次current", "#FCE4D6", "#C65911"),
         "db":(8.2,4.0,1.5,0.75,"forecast_results\n0行", "#F4CCCC", "#A61C2A"),"seed":(1.2,1.8,1.5,0.75,"seed/demo\n读时可能写入", "#FFF2CC", "#C7901E"),
         "fallback":(4.2,1.8,1.5,0.75,"fallback/派生\n页面兜底", "#FFF2CC", "#C7901E"),"consumer":(7.2,1.8,1.5,0.75,"Web / AI\n无法统一追溯", "#F4CCCC", "#A61C2A")},
        [("local","registry","未注册"),("registry","versions","两套表"),("versions","consumer","读取"),("files","consumer","读取"),("db","consumer","读取"),("seed","versions","补空"),("fallback","consumer","兜底")],
        "当前事实源分裂", note="多个读写路径在模型身份、领域、来源类型和run_id上不一致，存在业务误导风险。")

    figs["fact_target"] = save_flow(
        "06_target_fact_source",
        {"candidate":(0.2,3.0,1.3,0.75,"Candidate\n只读产物"),"admit":(1.9,3.0,1.3,0.75,"准入验证\nhash/契约/指标"),"registry":(3.6,3.0,1.3,0.75,"统一模型注册\n唯一Active"),
         "run":(5.3,3.0,1.3,0.75,"统一run_id\n预测事实"),"api":(7.0,3.0,1.3,0.75,"统一API\n来源元数据"),"users":(8.7,3.0,1.1,0.75,"Web/AI/\n报告"),
         "audit":(5.3,1.3,1.3,0.75,"审计/监控\n可回滚", "#E2F0D9", "#548235")},
        [("candidate","admit",""),("admit","registry","人工晋升"),("registry","run",""),("run","api",""),("api","users",""),("run","audit","记录")],
        "目标统一事实源", ylim=(0.5,4.8), note="real、seed、demo、fallback必须显式标识；读请求不写入，旧批次不删除。")

    figs["ai_chain"] = save_flow(
        "07_ai_chain",
        {"fe":(0.1,3.0,1.0,0.75,"前端"),"api":(1.35,3.0,1.0,0.75,"API/SSE"),"intent":(2.6,3.0,1.0,0.75,"意图识别"),"tools":(3.85,3.0,1.0,0.75,"工具/RAG"),
         "provider":(5.1,3.0,1.0,0.75,"Provider"),"planner":(6.35,3.0,1.0,0.75,"Planner"),"guard":(7.6,3.0,1.0,0.75,"Answer\nGuard"),"answer":(8.85,3.0,1.0,0.75,"用户回答")},
        [("fe","api",""),("api","intent",""),("intent","tools",""),("tools","provider",""),("provider","planner",""),("planner","guard",""),("guard","answer","")],
        "AI助手调用主链", ylim=(1.1,4.8), note="链路代码已接入；上游事实源不统一时，工具返回与最终回答均不能自动视为业务事实。")

    figs["rag_flow"] = save_flow(
        "08_rag_flow",
        {"query":(0.2,3.0,1.2,0.75,"用户问题"),"rewrite":(1.8,3.0,1.2,0.75,"查询改写"),"embed":(3.4,3.0,1.2,0.75,"Embedding\n身份记录"),
         "retrieve":(5.0,3.0,1.2,0.75,"关键词+向量\n检索"),"rerank":(6.6,3.0,1.2,0.75,"reranker\n重排"),"context":(8.2,3.0,1.2,0.75,"证据上下文\n引用"),
         "kb":(5.0,1.4,1.2,0.75,"知识库\n当前6/6", "#F4CCCC", "#A61C2A"),"guard":(8.2,1.4,1.2,0.75,"回答生成\nGuard", "#EEF5FB", "#2F75B5")},
        [("query","rewrite",""),("rewrite","embed",""),("embed","retrieve",""),("kb","retrieve","数据"),("retrieve","rerank",""),("rerank","context",""),("context","guard","生成")],
        "RAG完整流程与当前阻点", ylim=(0.5,4.8), note="当前直接查询超时且2项排序回归失败；历史9/10不能替代当前评分。")

    figs["model_lifecycle"] = save_flow(
        "09_model_lifecycle",
        {"candidate":(0.3,3.0,1.2,0.75,"Candidate\n候选"),"validate":(2.0,3.0,1.2,0.75,"安全加载\n契约/指标验证"),"active":(3.7,3.0,1.2,0.75,"Active\n唯一生效"),
         "predict":(5.4,3.0,1.2,0.75,"预测\n绑定run_id"),"monitor":(7.1,3.0,1.2,0.75,"监控\n漂移/误差"),"rollback":(8.8,3.0,1.0,0.75,"回滚"),
         "gate":(2.0,1.4,1.2,0.75,"人工准入门禁", "#FFF2CC", "#C7901E")},
        [("candidate","validate",""),("validate","active","通过"),("active","predict",""),("predict","monitor",""),("monitor","rollback","异常"),("gate","active","批准")],
        "目标模型生命周期", ylim=(0.5,4.8), note="基线日仅确认Candidate；未授权安全反序列化和晋升前，不得将其称为Active。")

    # 状态迁移图
    fig, ax = plt.subplots(figsize=(11.5, 6.1), dpi=180)
    labels = ["已验证完成", "部分完成", "被阻塞", "重复实现", "尚未开始", "已废弃"]
    p1 = [12, 12, 6, 1, 1, 1]
    p2 = [14, 16, 5, 1, 1, 1]
    x = range(len(labels))
    ax.bar([i-0.18 for i in x], p1, width=0.36, label="阶段一", color="#9CA3AF")
    ax.bar([i+0.18 for i in x], p2, width=0.36, label="阶段二", color="#2F75B5")
    ax.set_xticks(list(x), labels, rotation=15, ha="right")
    ax.set_ylabel("模块数")
    ax.set_title("阶段一到阶段二模块状态变化", fontsize=16, fontweight="bold", color="#17365D")
    ax.grid(axis="y", color="#E5E7EB", linewidth=0.8)
    ax.legend(frameon=False)
    for i, v in enumerate(p1): ax.text(i-0.18, v+0.25, str(v), ha="center", fontsize=9)
    for i, v in enumerate(p2): ax.text(i+0.18, v+0.25, str(v), ha="center", fontsize=9)
    ax.spines[["top","right"]].set_visible(False)
    fig.tight_layout()
    path = ASSETS / "10_phase_change.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    figs["phase_change"] = path

    figs["roadmap"] = save_flow(
        "11_p0_p1_p2_roadmap",
        {"a":(0.2,3.5,1.0,0.7,"A\n安全/恢复"),"b":(1.4,3.5,1.0,0.7,"B\n统一事实"),"c":(2.6,3.5,1.0,0.7,"C\n安全加载"),"d":(3.8,3.5,1.0,0.7,"D\n预测事务"),
         "e":(5.0,3.5,1.0,0.7,"E\nWeb/AI同源"),"f":(6.2,3.5,1.0,0.7,"F\nRAG恢复"),"g":(7.4,3.5,1.0,0.7,"G\n报告/任务"),"h":(8.6,3.5,1.0,0.7,"H\n历史收敛"),
         "p0":(1.8,1.5,2.1,0.65,"第一批：P0 / T001—T005", "#F4CCCC", "#A61C2A"),"p1":(4.4,1.5,2.1,0.65,"第二批：P1 / T006—T010", "#FFF2CC", "#C7901E"),"p2":(7.0,1.5,2.1,0.65,"第三批：P2 / T011—T012", "#EEF5FB", "#2F75B5")},
        [("a","b",""),("b","c",""),("c","d",""),("d","e",""),("e","f",""),("f","g",""),("g","h","")],
        "P0、P1、P2实施路线", ylim=(0.5,5.2), note="安全线T004可并行启动；模型与预测主线按T001→T002→T003→T005推进。")

    figs["milestones"] = save_flow(
        "12_milestones",
        {"demo":(0.2,3.0,1.25,0.75,"当前\n开发演示", "#FCE4D6", "#C65911"),"safe":(1.85,3.0,1.25,0.75,"安全默认\n可恢复基线"),"fact":(3.5,3.0,1.25,0.75,"事实源统一\n唯一Active"),
         "forecast":(5.15,3.0,1.25,0.75,"预测闭环\nrun_id事务"),"ai":(6.8,3.0,1.25,0.75,"Web/AI/RAG\n同源可验证"),"stable":(8.45,3.0,1.25,0.75,"目标\n可稳定本地运行", "#E2F0D9", "#548235")},
        [("demo","safe","P0"),("safe","fact","P0"),("fact","forecast","P0"),("forecast","ai","P1"),("ai","stable","验收")],
        "从“开发演示”到“可稳定本地运行”的里程碑", ylim=(1.0,4.8), note="里程碑以证据门禁定义，不给出缺乏依据的精确工期。")
    return figs


FIGS = build_figures()


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(COLORS["dark"])
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in [
        ("Title", 26, COLORS["navy"], 12, 12),
        ("Heading 1", 16, COLORS["navy"], 12, 8),
        ("Heading 2", 13, COLORS["blue"], 10, 6),
        ("Heading 3", 11, COLORS["dark"], 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(COLORS["dark"])
    caption.paragraph_format.keep_with_next = True

    for list_name in ["List Bullet", "List Bullet 2", "List Number"]:
        st = styles[list_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(3)

    if "Callout" not in styles:
        st = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.font.size = Pt(10.5)
        st.font.color.rgb = RGBColor.from_string(COLORS["navy"])
        st.paragraph_format.left_indent = Cm(0.5)
        st.paragraph_format.right_indent = Cm(0.5)
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_together = True

    # 页眉与页脚
    header = sec.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(PROJECT + "｜项目状态基线与优化规划")
    set_run_fonts(r, east="微软雅黑", size=8.5, color=COLORS["gray"])
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"基线日期：{BASELINE_DATE}　｜　第 ")
    set_run_fonts(r, size=8.5, color=COLORS["gray"])
    fr = add_field(p, " PAGE ", "1")
    set_run_fonts(fr, size=8.5, color=COLORS["gray"])
    r2 = p.add_run(" 页")
    set_run_fonts(r2, size=8.5, color=COLORS["gray"])

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT)
    set_run_fonts(r, east="微软雅黑", size=20, bold=True, color=COLORS["blue"])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(35)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("当前状态、技术架构与\n后续优化实施规划报告")
    set_run_fonts(r, east="微软雅黑", size=28, bold=True, color=COLORS["navy"])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(25)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CURRENT STATE, TECHNICAL ARCHITECTURE\nAND OPTIMIZATION IMPLEMENTATION PLAN")
    set_run_fonts(r, east="微软雅黑", latin="Arial", size=10.5, color=COLORS["gray"])

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    info = [
        ("基线日期", BASELINE_DATE),
        ("文档版本", VERSION),
        ("文档性质", "项目状态基线与优化规划"),
        ("事实认定", "阶段二运行证据优先于阶段一静态结论"),
    ]
    for i, (k, v) in enumerate(info):
        table.rows[i].cells[0].width = Cm(4.0)
        table.rows[i].cells[1].width = Cm(9.5)
        set_cell_shading(table.rows[i].cells[0], COLORS["navy"])
        set_cell_shading(table.rows[i].cells[1], "F8FAFC")
        for c in table.rows[i].cells:
            set_cell_margins(c, 110, 150, 110, 150)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = table.rows[i].cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(k)
        set_run_fonts(r0, east="微软雅黑", size=10, bold=True, color=COLORS["white"])
        p1 = table.rows[i].cells[1].paragraphs[0]
        r1 = p1.add_run(v)
        set_run_fonts(r1, east="微软雅黑", size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("内部项目资料｜不含账号、密码、Token、API Key或连接串")
    set_run_fonts(r, east="微软雅黑", size=9, color=COLORS["gray"])
    doc.add_page_break()


def add_front_matter(doc):
    add_title(doc, "保密与使用说明", 1)
    add_paragraph(doc, "本报告依据2026年7月14日形成的阶段一静态审计与阶段二运行验证成果编制，用于项目事实基线、技术复盘、后续开发任务下发和阶段汇报。报告不构成生产部署证明、交易建议或真实收益承诺。未经项目负责人确认，不应将Candidate模型、seed/demo/fallback数据、历史RAG/AI评测或测试子集结果对外表述为当前生产能力。")
    add_paragraph(doc, "报告未记录任何账号、密码、Token、API Key、数据库连接串或敏感配置值；引用路径仅用于审计追溯。阶段一、阶段二原始审计文件均保持原样，报告生成仅新增交付文件和图表资产。", source="阶段二执行计划与安全边界")

    add_table(doc, "文档信息", ["项目", "内容"], [
        ["项目名称", PROJECT], ["报告名称", REPORT_TITLE], ["基线日期", BASELINE_DATE], ["文档版本", VERSION],
        ["主要读者", "项目负责人、数据分析师、Python/AI开发人员、后续Codex、技术负责人"],
        ["审计口径", "运行证据＞代码存在；测试证据＞文档声明；阶段二当前事实＞阶段一历史结论"],
        ["报告边界", "不继续审计代码、不实施业务修改、不写现有数据库、不加载未授权模型二进制"],
    ], widths=[4.0, 12.2])

    add_table(doc, "修订记录", ["版本", "日期", "修订内容", "状态"], [
        [VERSION, BASELINE_DATE, "基于阶段一和阶段二审计成果形成首版正式项目报告", "已发布基线"],
    ], widths=[2.5, 3.0, 8.7, 2.5])

    add_title(doc, "目录", 1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', "请在Word中更新目录")
    doc.add_page_break()


def add_executive_summary(doc):
    add_title(doc, "执行摘要", 1)
    add_paragraph(doc, "项目已经形成覆盖数据采集、处理、特征、预测、Web、任务调度、AI助手与RAG的模块化单体系统，并具备可恢复的最小运行栈；但当前真实阶段仍是“开发演示”，尚未达到“可稳定本地运行”。决定性差距不在功能数量，而在模型与数据事实源、身份契约、run_id事务、安全默认以及可复现验证。", bold_lead="项目已经形成", source="阶段二最终项目状态基线，第1、4、38—40节")
    add_paragraph(doc, "总体工程完成度为64.6%，口径为排除duplicate和deprecated后的36个活动模块：已验证完成14个、部分完成16个、被阻塞5个、尚未开始1个；按1/0.5/0.25/0权重得分23.25/36。该指标仅表示模块工程成熟度，不代表生产可用度、AI准确率、测试全量通过率或部署成熟度。", bold_lead="总体工程完成度为64.6%", source="阶段二最终项目状态基线，第40节；M001—M038状态迁移")

    add_table(doc, "执行摘要关键判断", ["主题", "基线结论", "管理含义"], [
        ["项目阶段", "开发演示", "不得对外描述为生产系统或稳定本地闭环"],
        ["工程完成度", "64.6%（23.25/36）", "用于模块成熟度管理，不是可用率"],
        ["候选模型", "model_20260620_063015", "仅Candidate，未证明为Active"],
        ["快速预测", "不可用/未安全执行", "需先完成T001—T003"],
        ["Web", "可启动、构建通过", "页面成功不等于真实事实闭环"],
        ["DeepSeek", "真实最小调用可用", "完整401/429/5xx矩阵待补"],
        ["Ollama", "服务可达、无模型", "本地推理不可用"],
        ["RAG", "已进主链但当前无法可信评分", "当前6/6库、超时与排序失败需修复"],
        ["测试", "安全子集98/103；无全量通过率", "21文件收集错误必须先修复"],
    ], widths=[3.1, 6.5, 6.7])

    add_title(doc, "管理层需要记住的五项P0", 2)
    add_bullets(doc, [
        "P0-1：模型事实源分裂，读请求可能写入seed并改变Active状态。",
        "P0-2：候选模型尚未完成安全加载，运行时特征缺列会静默补0。",
        "P0-3：刷新、预测、同步没有统一run_id和原子事务，且存在全表覆盖风险。",
        "P0-4：认证默认开放、默认JWT以及前后端认证口径不一致，必须改为fail-closed。",
        "P0-5：seed、demo与fallback可能在真实页面冒充业务事实，需要统一来源标记和空态。",
    ], source="阶段二根因分析与风险清单；P0-P1-P2优化任务总表")

    add_title(doc, "第一批优化任务与进入下一阶段的条件", 2)
    add_paragraph(doc, "安全线可独立启动T004；模型与预测主线必须按T001→T002→T003→T005推进。第一批完成后，项目至少应具备：唯一模型事实源、经授权的安全加载、严格特征契约、原子run_id预测链、生产认证fail-closed、真实/seed/demo/fallback清晰区分，以及在隔离库中的双跑、幂等、失败回滚和旧批次共存证据。", source="阶段二后续Codex任务包，T001—T005")
    add_paragraph(doc, "在上述条件满足前，不建议开展LoRA/SFT、LangGraph复杂化、更换向量数据库、大规模UI重构、新模型堆叠、大量新增知识资料或历史代码删除。当前瓶颈属于事实源、模型身份、契约、事务、安全和运行闭环，而不是大模型参数能力不足。", source="阶段二最终项目状态基线，第43节")
    doc.add_page_break()


def chapter(doc, n, title, paragraphs=None, bullets=None, page_break=False):
    if page_break:
        doc.add_page_break()
    add_title(doc, f"第{n}章 {title}", 1)
    for item in paragraphs or []:
        if isinstance(item, tuple):
            add_paragraph(doc, item[0], source=item[1])
        else:
            add_paragraph(doc, item)
    if bullets:
        add_bullets(doc, bullets)


def build_body(doc):
    chapter(doc, 1, "项目背景与建设目标", [
        ("项目以电力市场价格预测与智能运营分析为主线，逐步形成数据采集、数据处理、模型训练与预测、业务分析、Web平台、报告、任务调度以及AI/RAG能力。当前报告的目标不是再次检查代码，而是把阶段一、阶段二已形成的证据整合为长期可引用的事实基线和优化依据。", "阶段一全量盘点，第1、4节"),
        ("目标状态不是“页面更多”或“模型更多”，而是让数据、模型、预测、报告和AI回答共享可追溯身份，在本地环境可重复启动、失败可回滚、权限默认安全，并能够用明确验收证据证明。", "阶段二最终项目状态基线，第38—43节"),
    ])

    chapter(doc, 2, "项目发展历程与版本演进", [
        ("项目从v2/v3/v4单体预测脚本演进到v4_fix1兼容核心，再由prediction_engine、services和model_ops形成模块化外壳；用户入口由旧Tkinter迁移到PySide6，随后形成FastAPI/React Web主入口。AI侧保留旧AI链、LangGraph/Dify实验链和当前ai_assistant主链。", "阶段一全量盘点，第7节；阶段二端到端报告，历史引用图"),
        ("产品版本为v2.11.2；当前分支p5-frontend-ai-experience，HEAD为f8dc0bc，但工作区仍有86个modified和304个untracked，因此HEAD不能单独代表基线磁盘状态。阶段二已用Manifest、关键入口哈希、patch和ZIP建立恢复材料。", "阶段二最终项目状态基线，第3节；P2E002—P2E006"),
    ])

    chapter(doc, 3, "阶段一和阶段二审计说明", [
        ("阶段一采用只读静态审计、只读数据库查询和既有产物交叉核验，形成全量盘点、模块状态、证据索引和待验证队列；阶段二按Gate 0—5执行运行、迁移、外部最小调用、模型/契约、AI/RAG、测试/build等定向验证，未实施业务代码优化。", "阶段一全量盘点，第2—3节；阶段二执行计划与安全边界"),
        ("事实认定遵循：阶段二运行证据优先于阶段一静态结论；测试证据优先于文档声明；当前事实优先于历史结果；无法安全验证的内容标记为被阻塞或无法验证。阶段二对DeepSeek、FastAPI、React、PySide6和Celery获得了新证据，也下调了正式预测、Embedding、AI工具、安全和自动测试的成熟度。", "阶段二最终项目状态基线，第2节"),
    ])
    add_table(doc, "阶段一与阶段二工作范围对照", ["维度", "阶段一", "阶段二", "事实优先级"], [
        ["主要方法", "静态审计、只读SQL、既有产物核验", "隔离运行、迁移、最小外部调用、测试/build", "阶段二运行证据优先"],
        ["代码修改", "无", "无业务代码修改", "保持原始工作区"],
        ["数据库", "现有库只读", "现有库只读；新建隔离库迁移", "不将隔离数据当正式数据"],
        ["模型", "校验文件/元数据", "未授权反序列化；加载跳过", "文件存在≠可加载Active"],
        ["输出", "4份核心审计成果", "验证矩阵、状态迁移、任务与专项报告", "报告统一索引"],
    ], widths=[2.8,4.5,5.2,4.0])

    chapter(doc, 4, "当前总体状态与完成度", [
        ("当前总体阶段为“开发演示”。最小基础栈、页面构建、DeepSeek最小问答和部分数据采集有运行证据，但真实Active快速预测、统一run同步、真实模型中心、当前RAG评分、100题、报告状态机、策略闭环与Ollama推理仍不可稳定使用。", "阶段二最终项目状态基线，第38—40节"),
        ("完成度64.6%按活动模块风险加权计算：14个已验证完成、16个部分完成、5个被阻塞、1个尚未开始；得分14×1+16×0.5+5×0.25+1×0=23.25。duplicate与deprecated不进入分母。", "M001—M038状态迁移，summary"),
    ], page_break=True)
    add_table(doc, "36个活动模块状态分布", ["状态", "数量", "权重", "加权得分", "解释"], [
        ["已验证完成", 14, 1, 14, "存在当前运行/测试或强证据"],
        ["部分完成", 16, 0.5, 8, "实现或部分运行，尚缺闭环证据"],
        ["被阻塞", 5, 0.25, 1.25, "前置条件或安全风险阻断"],
        ["尚未开始", 1, 0, 0, "LoRA/SFT门槛未满足"],
        ["合计", 36, "—", 23.25, "23.25/36=64.6%"],
    ], widths=[3.0,2.0,2.0,2.5,7.0])
    add_figure(doc, FIGS["phase_change"], "阶段一到阶段二模块状态变化", "M001—M038状态迁移；阶段二最终项目状态基线")

    chapter(doc, 5, "当前真实技术架构", [
        ("当前结构是以React Web和FastAPI为主入口的模块化单体：业务服务连接PostgreSQL、Redis/Celery和预测流水线；预测侧保留v4_fix1兼容核心与fast_forecast并存；AI侧由ai_assistant调用工具、RAG、Provider、Planner和Answer Guard。", "阶段二最终项目状态基线，第4—5节"),
        ("结构已形成，但边界尚未收敛：Web、PySide6、历史脚本和多套启动/依赖仍并存；本地artifact、模型表、预测表、current目录和页面fallback没有统一为同一事实链。", "阶段二端到端链路验证报告，历史引用图与事实源"),
    ])
    add_figure(doc, FIGS["architecture"], "当前总体技术架构", "阶段一全量盘点，第4—6节；阶段二最终状态基线，第4节")
    add_table(doc, "项目核心技术栈", ["层级", "主要技术", "当前角色", "状态"], [
        ["前端", "React 18 / TypeScript / Vite / Ant Design / ECharts", "唯一主用户入口", "已验证完成（build/启动）"],
        ["后端", "FastAPI / SQLAlchemy / 原生SQL仓储", "API与业务服务", "已验证完成（最小栈）"],
        ["数据", "PostgreSQL 18.4 / 本地Excel与CSV", "结构与事实存储", "部分完成"],
        ["任务", "Redis 7 / Celery 5.6.3", "异步调度", "部分完成"],
        ["预测", "pandas / scikit-learn / v4_fix1兼容内核", "训练、正式与快速预测", "部分完成/关键链被阻塞"],
        ["AI", "DeepSeek / Ollama / LLMRouter / Answer Guard", "问答、路由与安全输出", "部分完成"],
        ["RAG", "BGE / JSON向量 / 关键词检索 / reranker", "知识检索与引用", "部分完成"],
        ["桌面", "PySide6", "辅助入口", "已验证完成（导入边界）"],
    ], widths=[2.2,5.2,4.2,4.5])

    chapter(doc, 6, "当前代码版本、工作区与可复现基线", [
        ("代码基线为v2.11.2、分支p5-frontend-ai-experience、HEAD f8dc0bc；本地分支领先远端10个提交。阶段二Gate 0记录4,855项Manifest、27项关键入口哈希、tracked patch和304项未跟踪文件ZIP。", "阶段二命令与测试日志，Gate 0"),
        ("最终按相同排除规则复核范围外4,855个文件，新增0、缺失0、内容变化0；阶段二新增材料均位于docs/phase2，证明审计未改变范围外项目文件。", "阶段二命令与测试日志，清理"),
    ])
    add_table(doc, "当前主要入口", ["入口", "命令/接口", "定位", "基线结论"], [
        ["React Web", "npm run dev / Vite", "主用户入口", "可隔离启动并构建"],
        ["FastAPI", "backend.app.main:app", "后端主入口", "健康、DB、task检查200"],
        ["日常流水线", "main_daily_run.py", "采集/训练/预测编排", "完整写链未执行"],
        ["快速预测", "forecast_with_saved_model", "24小时预测", "被Active/加载/事务阻塞"],
        ["AI聊天", "/api/ai/chat 与 /api/ai/stream", "AI主链", "代码存在、最小Provider可用"],
        ["PySide6", "python -m ui.app", "辅助入口", "29模块导入通过"],
    ], widths=[2.8,5.0,4.0,4.4])

    chapter(doc, 7, "数据采集、数据处理与数据库现状", [
        ("采集实现覆盖PJM日前/实时价格与负荷、Open-Meteo天气、分页、重试、超时、增量区间、去重合并和America/New_York时区转换。阶段二经授权完成PJM单节点单日24行与Open-Meteo同日24行最小实调，未写入业务目录。", "阶段二命令与测试日志，Gate 3；P2E025"),
        ("PostgreSQL现有schema为58表、迁移head 0012；独立空库两次迁移成功且schema diff=0。当前业务事实仍不完整：forecast_runs=1而forecast_results=0，model_registry=0而model_versions=5，报告/策略表为0，知识库为6/6。", "阶段二端到端链路验证报告，数据库"),
    ], page_break=True)
    add_figure(doc, FIGS["data_flow"], "数据采集与数据处理流程", "阶段一全量盘点，第10—13节；阶段二V021")
    add_table(doc, "数据事实源现状", ["对象", "当前事实", "风险", "状态"], [
        ["原始市场/负荷/天气", "历史表已有两年数据；小窗口采集可用", "正式增量/DST/限流矩阵待补", "已验证完成"],
        ["主表/特征表", "主表17,488行；特征表16,902行/172列（历史产物）", "需要与run_id和模型身份绑定", "已验证完成（历史批次）"],
        ["forecast_runs/results", "1/0", "ready记录无结果，不能视为有效预测", "被阻塞"],
        ["model_registry/versions", "0/5", "快速加载与Web展示使用不同表", "被阻塞"],
        ["report/strategy", "0", "无同run审核、发布、派发与收益证据", "部分完成"],
        ["kb_documents/chunks", "6/6 seed", "历史54/583尚未恢复", "部分完成"],
    ], widths=[3.2,5.1,5.1,3.0])

    chapter(doc, 8, "特征工程与数据泄漏治理", [
        ("历史特征表包含时间、周期、滞后、滚动、负荷、天气、价差、交互和峰值/尖峰特征；阶段一防泄漏证据确认按时间顺序切分、未随机打乱、目标未直接作为特征、未使用预测结果，未发现高/中风险直接泄漏。", "阶段一全量盘点，第13节"),
        ("阶段二将24行冻结输入与170项schema静态核对，缺失0、额外0、顺序/hash一致、空值0、非数值0、小时连续；但时间戳无时区，运行时_prepare_history对缺失列静默补0，三个误差记忆特征当前为0占位，因此不能宣称严格契约或完全无泄漏保证。", "阶段二端到端链路验证报告，特征契约"),
    ])

    chapter(doc, 9, "预测模型体系与当前可信模型", [
        ("模型体系包含基础Ridge、高峰专项模型、尖峰风险分类器及融合后处理。唯一最新完整候选为model_artifacts/model_20260620_063015，含11个文件，模型版本model_20260620_063015、特征版本features_140db8af25f9、目标da_price、170项特征。", "阶段二端到端链路验证报告，模型"),
        ("训练/验证/测试为15,462/720/720样本，测试窗口2026-05-19至2026-06-17。阶段二没有反序列化joblib，因此它只能定义为Candidate；文件存在、静态完整和指标可复算，不等于当前环境可安全加载或已晋升Active。", "V003；阶段二命令与测试日志，安全停止"),
    ], page_break=True)
    add_table(doc, "当前可信候选模型指标", ["指标", "融合模型", "解释"], [
        ["MAE", "16.084156", "720个测试样本平均绝对误差"],
        ["RMSE", "32.700131", "对大误差更敏感的均方根误差"],
        ["MAPE", "20.359783%", "相对误差，需注意低/负价格场景"],
        ["R²", "0.944455", "测试集解释度，不等于部署稳定性"],
        ["尖峰价格子集MAE/RMSE", "42.089 / 67.7078", "价格回归指标，不是分类precision/recall"],
    ], widths=[4.0,4.0,8.2])

    chapter(doc, 10, "模型注册、Active生命周期与特征契约", [
        ("Active模型是经准入验证后唯一生效、可被生产预测加载的模型；Candidate是待验证候选。当前model_registry为空，Web模型中心读model_versions，且读请求可能通过seed逻辑改变Active，因此没有可信唯一Active。", "阶段二最终项目状态基线，第9—12节；R001"),
        ("目标生命周期必须包含Candidate、验证、人工晋升Active、绑定run_id预测、监控和回滚；artifact、feature_version、schema hash与指标必须同批次。任何契约缺失应fail-closed，不得静默补0或跨表回退。", "T001—T003"),
    ])
    add_figure(doc, FIGS["model_lifecycle"], "模型生命周期目标流程", "阶段二T001—T003；V003—V005")
    add_table(doc, "Candidate与Active状态说明", ["对象", "当前身份", "已验证内容", "未验证内容"], [
        ["model_20260620_063015", "唯一最新完整Candidate", "文件清单、元数据、720样本指标复算、静态schema", "joblib安全加载、固定24行推理、晋升Active"],
        ["model_registry", "空表", "表存在", "唯一Active、生命周期、回滚"],
        ["model_versions", "5条seed/展示记录", "Web读取路径存在", "与真实电价Candidate同源、无读时写入"],
    ], widths=[4.2,3.8,4.3,4.3])

    chapter(doc, 11, "正式预测与快速预测现状", [
        ("2026-06-18存在24行正式预测文件，但文件未包含model_version或feature_version，且时间早于6月20候选artifact，无法证明来自该候选模型。阶段一“正式预测已验证完成”因此在阶段二下调为部分完成。", "阶段二最终项目状态基线，第14节；M014状态迁移"),
        ("快速预测当前不可用：无可信唯一Active、模型二进制未安全加载、运行时契约可能静默补0、同步含覆盖风险。阶段二按安全停止条件未执行注册、激活、快速预测或同步写入。", "阶段二命令与测试日志，Gate 3与安全停止"),
    ])
    add_figure(doc, FIGS["current_forecast"], "当前预测主链及阻断位置", "阶段二端到端链路验证报告，模型/契约/预测身份")
    add_figure(doc, FIGS["ideal_forecast"], "理想快速预测闭环", "T001—T003验收目标")
    add_table(doc, "快速预测阻塞点", ["阻塞点", "当前现象", "风险", "对应任务"], [
        ["Active身份", "model_registry=0；model_versions为seed", "加载错误/无模型", "T001"],
        ["安全加载", "joblib未获授权加载", "供应链/反序列化风险", "T002"],
        ["特征契约", "静态匹配但运行时静默补0", "产生看似合理的错误预测", "T002"],
        ["run_id", "预测文件缺模型/特征身份", "无法追溯报告与页面", "T003"],
        ["同步事务", "先删后插/全表覆盖语义", "失败丢历史、半成品", "T003"],
    ], widths=[3.0,5.1,4.8,2.5])

    chapter(doc, 12, "run_id与数据库同步链路", [
        ("run_id是一次数据刷新、预测、报告、策略和任务结果的统一批次身份。当前预测文件、数据库记录、current目录和模型身份没有共享同一run_id；sync_model_facts与sync_forecast_facts存在删除/覆盖语义，不能支持旧批次共存和失败回滚。", "R002—R003"),
        ("目标链应以新run_id创建批次，预测24行原子upsert，model/artifact/feature身份随记录保存；同run重复执行hash一致，不同run并存，异常无半成品。旧批次和既有artifact不得清空。", "T003验收标准"),
    ])

    chapter(doc, 13, "业务分析、策略建议与报告系统", [
        ("业务分析、储能/交易建议、风险时段、峰谷价差与报告生成均有代码或本地产物，但report/strategy事实表为空，current目录跨批次。没有同run_id的预测→报告→审核→发布→派发，也没有规则版本、约束与理论收益闭环。", "阶段二端到端链路验证报告，报告与策略"),
        ("因此报告与策略只能描述为“部分完成”：可展示历史归档或代码能力，但不得作为已发布运营报告、真实交易建议或实际收益证明。后续T008应在上游T003完成后建立状态机、权限、幂等和审计证据。", "R008；T008"),
    ], page_break=True)

    chapter(doc, 14, "FastAPI后端状态", [
        ("FastAPI在127.0.0.1:8001隔离启动，health连续3次200，DB/task health 200，未认证trace 401，CORS允许200/拒绝400。阶段一“已实现但未验证”升级为“已验证完成（最小运行栈）”。", "阶段二命令与测试日志，Gate 2；M030"),
        ("这一结论不覆盖全部业务路由、认证角色矩阵或端到端写链。后端仍受默认认证、依赖收集错误、事实表空和运行闭环阻塞影响。", "V025—V027"),
    ])

    chapter(doc, 15, "React Web前端状态", [
        ("React Web在Vite 5174隔离启动，首页与/api/health代理均返回200；TypeScript检查与隔离Vite build通过，共3,675模块。阶段一“已实现但未验证”升级为“已验证完成（启动/build）”。", "阶段二命令与测试日志，Gate 2与Gate 5；M031"),
        ("但两项前端契约测试失败，模型/设置/来源契约仍不一致；浏览器E2E因工具沙箱初始化失败未执行。页面能启动和构建不等于预测、模型、报告与AI页面展示真实统一数据。", "阶段二最终项目状态基线，第34节"),
    ])

    chapter(doc, 16, "Redis、Celery与任务调度", [
        ("隔离Redis ping通过，Celery worker注册9个任务，无副作用health任务0.811秒成功；因此调度由阶段一blocked升级为部分完成。", "阶段二命令与测试日志，Gate 2；M034"),
        ("尚未验证queued/running/success/failed/cancel/retry完整状态机、worker重启、重复任务去重和业务任务副作用。Celery任务需要幂等，即同一任务重复执行不产生重复结果或不可逆覆盖。", "V022；T009"),
    ])

    chapter(doc, 17, "AI助手总体架构与调用链", [
        ("当前主链为前端→API/SSE→意图识别→工具/RAG→Provider→Planner→Answer Guard→用户回答。工具、RAG和Guard已经进入answer_chat_accurate主链；旧AI链和LangGraph/Dify被归类为重复/实验实现。", "阶段一全量盘点，第22节；阶段二最终状态基线，第17、37节"),
        ("Answer Guard是回答输出前的安全与质量门禁，负责隐藏内部Trace、工具细节和敏感字段。当前单测与未认证负向有证据，但普通用户、管理员、调试用户三角色的API/SSE/UI/导出矩阵尚未完成。", "阶段二AI与RAG评测报告，普通用户边界"),
    ], page_break=True)
    add_figure(doc, FIGS["ai_chain"], "AI助手调用链", "阶段一AI助手主链；阶段二V010—V012")

    chapter(doc, 18, "DeepSeek、Ollama及多模型路由", [
        ("Provider是向AI主链提供推理能力的外部或本地模型服务。DeepSeek经授权完成真实health和固定16-token最小问答，health约945ms、问答约1,599ms，结论为真实可用；但401、429、5xx和网络中断矩阵未完整注入。", "阶段二AI与RAG评测报告，DeepSeek"),
        ("Ollama 0.31.2服务可达但模型列表为空，qwen3:4b未安装，因此只能说服务可达，不能说本地推理可用。LLMRouter的4项确定性测试通过，但真实auto多Provider故障回退尚未闭环。", "阶段二AI与RAG评测报告，Ollama与路由"),
    ])
    add_table(doc, "AI Provider状态", ["Provider/组件", "当前状态", "已验证", "缺口"], [
        ["DeepSeek", "已验证完成（最小调用）", "health与短问答成功", "401/429/5xx/断网矩阵"],
        ["Ollama", "被阻塞", "服务版本与11434可达", "无可推理模型；未做性能/结构化测试"],
        ["LLMRouter", "部分完成", "4项确定性路由测试通过", "真实多Provider auto故障闭环"],
        ["Answer Guard", "部分完成", "脱敏/边界单测与401负向", "三角色API/SSE/UI/E2E"],
    ], widths=[3.5,3.8,4.5,4.5])

    chapter(doc, 19, "AI工具与业务事实一致性", [
        ("AI工具实现覆盖预测、负荷、天气、模型、报告和策略等业务对象，但11类工具尚不能证明共享同一run_id与事实源。若工具读到空表、旧文件、seed或fallback，LLM即使语言质量正常，也可能给出错误业务事实。", "阶段二最终项目状态基线，第21节；R001—R003"),
        ("目标API应为每个数值返回source_type、domain、run_id、generated_at和可追溯证据；无真实数据时明确返回不可用或可解释空态。禁止把推测、派生值或演示数据包装为实时生产事实。", "T005验收标准"),
    ])
    add_figure(doc, FIGS["fact_split"], "当前事实源分裂", "阶段二端到端链路验证报告，事实源；R001—R003")
    add_figure(doc, FIGS["fact_target"], "目标统一事实源", "T001、T003、T005目标状态")

    chapter(doc, 20, "知识库、Embedding与RAG", [
        ("RAG（检索增强生成）通过查询改写、Embedding向量表示、关键词/向量检索、reranker重排和引用上下文，为AI回答提供项目知识证据。代码已进入聊天主链。", "阶段二最终项目状态基线，第27节"),
        ("当前磁盘原始资料83份、69个唯一内容hash、14份内容重复；历史成功导入54文档/583 chunk，当前数据库仅6文档/6 chunk。当前RAG直接查询两次超时，专项10测8过2败，失败为政策类排序，因此当前smoke只能标记无法验证/存在失败证据。历史9/10不能代表当前。", "阶段二AI与RAG评测报告"),
    ])
    add_figure(doc, FIGS["rag_flow"], "RAG完整流程与当前阻点", "阶段二AI与RAG评测报告；V013—V017")
    add_table(doc, "RAG当前状态", ["环节", "当前事实", "结论/风险"], [
        ["资料Manifest", "83份原始、69唯一hash、14重复、12历史解析失败", "资料可盘点，失败件需用户决定"],
        ["当前数据库", "6文档/6 chunk", "仅seed/sample，覆盖不足"],
        ["历史恢复目标", "54文档/583 chunk", "历史证据，不代表当前已恢复"],
        ["Embedding", "历史BGE-large-zh-v1.5/1024；当前6条身份不完整", "不能仅凭维度判断Provider"],
        ["运行查询", "120秒与45秒超时", "当前链不稳定"],
        ["专项测试", "8/10，2项排序失败", "不给出虚假当前通过率"],
    ], widths=[3.2,6.2,6.2])

    chapter(doc, 21, "AI评测和回答质量现状", [
        ("当前AI 100题没有重新执行，因为预测run_id、知识库、RAG运行和Ollama模型身份均未冻结。历史42/100只能作为历史基线，不能写成当前通过率。", "阶段二AI与RAG评测报告，AI 100题"),
        ("首要预期失败类型是事实数据缺失/版本错配、seed/fallback误用、RAG恢复与排序、Provider可用性和工具事实契约。应在T003、T006、T007后冻结题集、评分规则、Provider、数据库与run_id，再保存逐题证据。", "阶段二AI与RAG评测报告，AI 100题"),
    ], page_break=True)

    chapter(doc, 22, "PySide6和历史界面边界", [
        ("PySide6 6.11.1下29个UI Python文件AST与offscreen import全部通过，定位为辅助入口；未实例化GUI或执行写操作。Tkinter保持已废弃。", "阶段二命令与测试日志，Gate 5；V028"),
        ("建议Web继续作为唯一主入口，PySide6仅保留尚未在Web替代的本地辅助能力，并统一读取后端API。历史界面和入口的归档必须在核心链稳定后逐文件确认，禁止批量删除。", "阶段二端到端链路验证报告，历史引用图"),
    ])

    chapter(doc, 23, "测试、构建与运行验证结果", [
        ("后端pytest collect识别103项，但21个测试文件因缺少httpx2发生收集错误；排除TestClient阻塞后的安全子集为103项，98通过、5失败，通过率95.15%。该数值只能称为安全子集通过率，不能称为全量后端通过率。", "阶段二命令与测试日志，Gate 5"),
        ("TypeScript检查和隔离Vite build通过；两项前端契约测试失败；浏览器E2E因内置工具沙箱失败未执行，不能描述为通过或失败。测试曾写入日志和3条会话状态，阶段二已按Gate 0哈希精确恢复。", "阶段二命令与测试日志，Gate 5、浏览器、清理"),
    ])
    add_table(doc, "测试与构建结果", ["项目", "结果", "准确表述", "后续动作"], [
        ["后端收集", "21文件收集错误", "无可信全量通过率", "修复httpx2与依赖锁"],
        ["后端安全子集", "98/103，95.15%", "仅安全子集", "修复5项确定性失败"],
        ["TypeScript", "通过", "类型检查通过", "纳入发布门禁"],
        ["Vite build", "通过，3,675模块", "隔离构建通过", "继续控制大chunk"],
        ["前端契约", "2项失败", "契约未完全一致", "修复模型/设置契约"],
        ["浏览器E2E", "工具沙箱阻塞", "未执行/无法验证", "工具恢复后跑10关键页面"],
        ["PySide6", "29模块导入通过", "仅导入边界", "按辅助入口补导航验证"],
    ], widths=[3.1,3.7,5.1,4.3])

    chapter(doc, 24, "安全、认证和敏感配置风险", [
        ("当前安全不达标：后端默认AUTH_REQUIRED=false并存在默认JWT secret，前端默认认证为true；默认管理员初始化、密码依赖和角色口径也存在问题。缺少环境变量时，开发默认可能进入正式运行路径。", "R004；V025"),
        ("目标必须是fail-closed：生产环境缺secret或管理员凭据即拒绝启动；前后端认证口径一致；普通、管理员、调试用户越权负向通过；日志、SSE和导出不含敏感值。显式DEV模式可以保留，但不得默认启用。", "T004验收标准"),
    ])

    chapter(doc, 25, "阶段一到阶段二状态变化", [
        ("阶段二共记录5项升级、5项降级、28项不变。升级反映DeepSeek、FastAPI、React、PySide6和Redis/Celery获得了新的当前运行证据；降级反映正式预测、Embedding、AI工具、安全和测试在身份或运行层面未满足原静态判断。", "M001—M038状态迁移，summary"),
        ("这说明“能找到代码或文件”与“当前可运行且可追溯”之间存在显著差异。报告以后续验收证据更新状态，不以功能数量或历史截图作为完成证明。", "阶段二最终项目状态基线，第2节"),
    ], page_break=True)
    add_figure(doc, FIGS["phase_change"], "阶段一到阶段二状态变化（复用）", "M001—M038状态迁移")

    chapter(doc, 26, "当前可稳定使用或演示的功能", [
        ("可演示能力包括：最小栈健康、React页面构建、DeepSeek最小问答、外部小窗口采集，以及带来源说明的历史模型指标展示。演示时必须明确其证据边界和历史/当前身份。", "阶段二最终项目状态基线，第38节"),
    ])
    chapter(doc, 27, "当前不可稳定使用的功能", [
        ("不可稳定演示或不可用于业务决策的能力包括：真实Active快速预测、统一run同步、真实模型中心、当前RAG、AI 100题、报告审核发布状态机、策略闭环和Ollama本地推理。", "阶段二最终项目状态基线，第39节"),
    ])
    add_table(doc, "可演示与不可演示功能对照", ["能力", "当前可否", "使用条件/限制"], [
        ["最小服务栈健康", "可以演示", "使用隔离配置；不代表业务闭环"],
        ["React Web构建与页面", "可以演示", "需强调部分页面可能为空/seed/fallback"],
        ["DeepSeek最小问答", "可以演示", "不展示Key；完整故障矩阵待补"],
        ["PJM/Open-Meteo小窗口采集", "可以演示", "仅最小窗口，不写正式目录"],
        ["历史模型指标", "可以展示", "标注候选artifact和历史测试窗口"],
        ["真实Active快速预测", "不可", "T001—T003未完成"],
        ["模型中心作为真实事实", "不可", "存在seed与两套表"],
        ["当前RAG/AI评分", "不可", "当前库6/6、超时、排序失败、100题未跑"],
        ["策略/收益/发布闭环", "不可", "无同run状态机和真实收益证据"],
        ["Ollama推理", "不可", "服务可达但无模型"],
    ], widths=[4.5,3.0,8.0])

    chapter(doc, 28, "根因分析", [
        ("五项P0并非孤立缺陷，而是同一类工程问题的不同表现：系统没有把“谁产生、何时产生、属于哪个领域、使用哪个模型/特征、是否真实、如何回滚”固化为单一可执行契约。", "阶段二根因分析与风险清单"),
        ("直接后果是模型中心与快速预测看见不同事实、预测和报告无法绑定同一run、读请求可能改变数据库、页面兜底可能冒充业务数据，以及认证缺失时仍进入运行态。", "R001—R004；T005"),
    ], page_break=True)
    p0_details = [
        ("P0-1 模型事实源分裂及读时seed写入", "model_registry为空而model_versions含seed，读请求可能改变Active；影响模型展示、快速预测、Web/AI事实。直接原因是两套模型表与补空设计，深层原因是缺少唯一领域模型与只读边界。涉及M015/M031，ui_platform_service.py、model_repository.py，model_registry/model_versions/model_metrics。应以T001/T005修复：唯一表/视图、读无写、领域与来源强标记。验收为唯一Active、跨端一致、重复读无数据库变化；回滚保留旧读兼容，不删除旧数据。", "R001；P2E023—P2E024"),
        ("P0-2 模型加载与特征契约不安全", "候选artifact未获授权反序列化，_prepare_history对缺失列静默补0。影响预测正确性、审计和模型准入。直接原因是加载与契约校验未形成fail-closed，深层原因是artifact、feature_version与schema身份未强制绑定。涉及M006/M007/M013/M015、fast_forecast.py和model_registry。应以T002在一次性无网络只读沙箱验证hash/签名并加载；任何缺失、额外、顺序、dtype或时区不一致即失败。验收含固定24行推理、异常负向和无业务写入；未获授权时保持安全跳过。", "V003—V005；P2E019—P2E024"),
        ("P0-3 刷新、预测、同步缺少统一run_id与事务链", "预测文件缺模型/特征身份，同步先删后插，不同批次不能安全共存。影响历史保留、报告/策略追溯和失败恢复。直接原因是产物规范未强制身份字段，深层原因是唯一键、事务、幂等和审计契约缺失。涉及M013/M014/M017/M030/M031、main_daily_run.py、fast_forecast.py、core_data_sync.py及forecast_runs/results。T003要求新run原子写24行、重复执行幂等、旧run共存、异常回滚。", "R002—R003；P2E024"),
        ("P0-4 认证默认开放及默认JWT风险", "缺少env时后端免认证并使用默认JWT，前端却默认认证；管理员初始化与密码依赖存在失败。影响越权、部署误配置和审计可信度。涉及M030/M031/M036/M037、backend/app/config.py、auth、frontend、users/roles/audit_logs。T004应建立生产fail-closed、强制secret、角色负向、日志脱敏；安全线可与模型主线并行。回滚只能保留显式DEV模式，不能恢复不安全默认。", "R004；P2E011/P2E015/P2E035"),
        ("P0-5 seed、demo与fallback冒充真实数据", "空表时自动seed，负荷模型数据可能出现在电价模型页面，前端/AI兜底缺少统一来源标签。影响用户决策、演示诚信和后续评测。涉及M015/M020/M024/M031、ui_platform_service.py、frontend/src/pages、model_versions与kb_documents。T005要求空数据返回可解释空态，所有API携带source_type、domain、run_id、generated_at；不得将派生或历史值标为real。验收覆盖API、UI、AI回答和导出；回滚保留显式演示模式但不进入真实页面。", "T005；P2E024/P2E031"),
    ]
    for title, detail, src in p0_details:
        add_title(doc, title, 2)
        add_paragraph(doc, detail, source=src)

    chapter(doc, 29, "P0、P1、P2问题和风险", [
        ("任务总表包含P0 5项、P1 5项、P2 2项。P0直接阻断安全与预测主链；P1用于恢复知识库、测试、报告/策略、Celery和环境复现；P2补齐Provider/E2E和PySide6/历史边界。", "阶段二P0-P1-P2优化任务总表"),
    ])
    task_rows = []
    for t in TASKS_OBJ["tasks"]:
        task_rows.append([t["task_id"], t["priority"], t["title"], t["estimated_complexity"], "是" if t["user_decision_required"] else "否", ",".join(t["dependencies"]) or "无"])
    add_table(doc, "P0、P1、P2优化任务总览", ["编号", "优先级", "任务", "复杂度", "需用户决策", "依赖"], task_rows, widths=[1.5,1.5,7.1,1.7,2.2,2.5], compact=True)

    chapter(doc, 30, "后续优化总体原则", [
        ("优化应遵循安全优先、事实源优先、身份先于功能、事务先于自动化、隔离验证、最小变更、可回滚和证据驱动。任何写入型验证必须使用隔离数据库/输出目录；不得覆盖current/latest、既有artifact或正式结果。", "阶段二后续Codex任务包，通用执行门禁"),
        ("实施时先补失败回归并冻结输入/身份，再做最小文件级变更；单元、集成、端到端与回归测试必须分别定义。旧接口若需要兼容，应显式兼容，禁止隐式跨事实表或跨来源回退。", "T001—T012通用任务字段"),
    ], page_break=True)

    chapter(doc, 31, "第一批T001—T005实施计划", [
        ("推荐顺序分为两条线：安全线T004可立即并行；模型与预测主线按T001→T002→T003→T005。T002涉及模型反序列化和T004涉及凭据/认证，均需要用户明确授权或决策。", "阶段二后续Codex任务包"),
    ])
    for t in TASKS_OBJ["tasks"][:5]:
        add_title(doc, f"{t['task_id']} {t['title']}", 2)
        detail_rows = [
            ["当前问题", t["problem_statement"]], ["任务目标", t["target_state"]], ["业务价值", t["business_value"]],
            ["涉及模块", "、".join(t["affected_modules"])], ["涉及文件", "；".join(t["affected_files"])], ["前置依赖", "、".join(t["dependencies"]) or "无"],
            ["实施步骤", "；".join(t["implementation_steps"])], ["明确不修改", "；".join(t["do_not_change"])],
            ["数据库/迁移", t["migration_impact"]], ["兼容性", t["backward_compatibility"]], ["安全", t["security_considerations"]],
            ["单元测试", "；".join(t["unit_tests"])], ["集成测试", "；".join(t["integration_tests"])], ["端到端测试", "；".join(t["e2e_tests"])],
            ["回归测试", "；".join(t["regression_tests"])], ["验收标准", t["acceptance_criteria"]], ["回滚方案", t["rollback_plan"]],
            ["Codex独立执行", "适合" if t["can_codex_execute_independently"] else "不适合，需用户授权/决策"],
            ["推荐模型/推理", f"{t['recommended_codex_model']} / {t['recommended_reasoning_level']}"], ["预期成熟度提升", "关闭对应P0，形成可复现、可回滚的局部事实闭环"],
        ]
        add_table(doc, f"{t['task_id']}实施卡", ["字段", "内容"], detail_rows, widths=[3.3, 13.0], compact=True,
                  caption_source=f"任务总表；验证{','.join(t['source_validation_ids'])}；证据{','.join(t['source_evidence_ids'])}")

    chapter(doc, 32, "中长期实施路线图", [
        ("路线划分为A安全和可恢复基线、B统一模型事实与Active、C安全加载和特征契约、D快速预测与run_id数据库闭环、E Web和AI统一事实源、F知识库/RAG恢复与100题、G报告/策略/Celery稳定性、H历史收敛和非核心体验优化。", "阶段二任务总表与后续Codex任务包"),
        ("各阶段以验收门禁而非精确工期推进。A—D是进入“可稳定本地运行”的必要条件；E—G决定业务事实一致性和运营闭环；H应在核心稳定后开展。", "阶段二最终项目状态基线，第42—43节"),
    ])
    add_figure(doc, FIGS["roadmap"], "P0、P1、P2实施路线图", "阶段二P0-P1-P2任务总表")
    roadmap_rows = [
        ["A", "安全和可恢复基线", "冻结基线、认证/凭据策略", "可恢复包、安全门禁", "缺secret拒绝启动；基线可恢复", "是"],
        ["B", "统一模型事实与Active", "T001", "唯一模型表/视图", "读无写、唯一Active、来源/领域明确", "是"],
        ["C", "安全加载和特征契约", "T002", "沙箱加载证据、严格schema", "固定24行成功；契约异常失败", "是"],
        ["D", "快速预测与run_id闭环", "T003", "24行事务写入与回滚", "双跑幂等、旧批次共存", "是"],
        ["E", "Web和AI统一事实源", "T005及工具契约", "来源元数据与空态", "API/UI/AI同源", "是"],
        ["F", "知识库/RAG/100题", "T006/T007", "54/583恢复、当前评测", "导入幂等、评分可复现", "是"],
        ["G", "报告/策略/Celery", "T008/T009/T010", "状态机与环境收敛", "审核/派发/任务可审计", "否"],
        ["H", "历史收敛与体验", "T011/T012", "E2E与逐文件归档", "无误删、主入口唯一", "否"],
    ]
    add_table(doc, "阶段A—H路线与验收", ["阶段", "目标", "主要任务", "交付物", "验收标准", "阻塞下一阶段"], roadmap_rows,
              widths=[1.2,3.1,3.0,3.2,4.5,1.5], compact=True)

    chapter(doc, 33, "用户待决策事项", [
        ("以下事项涉及模型晋升、反序列化授权、凭据、隔离资产、资源分配和历史边界，Codex不应自行扩大授权。建议在对应任务启动前或进入验收前做出书面决定。", "阶段二未解决与需用户决策事项"),
    ], page_break=True)
    decisions = [
        ["是否将model_20260620_063015设为唯一Candidate", "建议：是，但仅设Candidate", "消除候选歧义；不等于晋升Active", "是", "T001启动前"],
        ["是否授权隔离沙箱加载joblib", "建议：在一次性无网络只读容器中授权", "否则T002保持安全跳过", "是", "T002验收前"],
        ["统一采用哪份PostgreSQL配置", "建议：明确开发唯一配置并修复漂移", "禁止猜测或回显密码", "是", "知识恢复/写测试前"],
        ["是否轮换凭据和JWT", "建议：轮换并移除默认fallback", "关闭已知安全风险", "是", "T004实施时"],
        ["是否保留隔离测试数据库", "建议：短期保留至P0复核后由用户处置", "保留迁移复核证据", "否", "P0复核完成后"],
        ["是否部署Ollama模型", "建议：按资源决定；未部署则auto不验收", "影响本地推理与真实回退", "否", "T011前"],
        ["是否继续保留PySide6", "建议：保留最小辅助只读能力", "Web保持唯一主入口", "否", "历史收敛阶段"],
        ["如何处置隔离Redis容器/镜像", "建议：验收后由用户逐项手动处置", "遵守禁止批量删除规则", "否", "阶段资产复核后"],
        ["如何处理12份知识解析失败件", "建议：人工转换或定向补解析器", "不得批量覆盖原件", "否", "T006前"],
        ["历史代码何时归档", "建议：核心闭环稳定后逐文件确认", "避免误删v4_fix1兼容核心", "否", "阶段H"],
    ]
    add_table(doc, "用户待决策事项", ["事项", "推荐选择", "理由/影响", "当前P0阻塞", "建议决策点"], decisions, widths=[4.2,4.2,4.5,1.8,2.2], compact=True)

    chapter(doc, 34, "暂不建议开展的工作", [
        ("当前暂不建议优先开展LoRA、SFT、LangGraph复杂化、更换或新增向量数据库、大规模UI重构、新模型堆叠、大量新增知识资料、删除历史实现，以及在事实源统一前继续增加业务页面。", "阶段二最终项目状态基线，第43节"),
        ("原因是现阶段主要失败来自数据与模型身份、契约、事务、安全、事实源和运行闭环，而不是大模型参数能力不足。扩展新能力会增加入口、事实源和测试矩阵，反而提高收敛成本。", "阶段二未解决与需用户决策事项，暂不启动"),
    ])

    chapter(doc, 35, "项目最终目标状态", [
        ("最终目标是“可稳定本地运行”：一套明确启动路径可恢复服务；生产认证默认关闭风险；唯一Active模型经安全准入；数据刷新、24小时预测、数据库同步、API、Web、AI、报告共享run_id和来源；失败可回滚；当前RAG和100题可复现；所有演示/降级数据显式标记。", "T001—T012验收目标"),
        ("该目标仍不等于企业生产部署、真实交易执行或收益实现。若未来进入生产化，仍需另行完成容量、SLA、灾备、密钥管理、网络安全、监控告警、数据治理和业务审批。", "报告边界；阶段二安全边界"),
    ])
    add_figure(doc, FIGS["milestones"], "从开发演示到可稳定本地运行的里程碑", "阶段二最终状态基线；T001—T012")

    chapter(doc, 36, "结论", [
        ("截至2026年7月14日，智能运营分析项目具备较完整的功能广度和可恢复的基础运行栈，但真实阶段仍是开发演示，工程完成度64.6%。可信候选模型与指标存在，DeepSeek真实最小调用可用，Web/FastAPI/Redis/Celery可在隔离配置下启动；同时快速预测、统一run同步、RAG当前评分、报告/策略状态机和安全默认仍未闭合。", "阶段二最终项目状态基线"),
        ("后续应首先完成T001—T005：统一模型事实源、安全加载与严格特征契约、run_id事务闭环、认证fail-closed和来源标签。只有这些验收通过并形成隔离双跑与回滚证据，项目才具备从“开发演示”进入“可稳定本地运行”的必要基础。", "阶段二后续Codex任务包"),
    ], page_break=True)


def add_appendices(doc):
    doc.add_page_break()
    add_title(doc, "附录A 模块状态总表", 1)
    add_paragraph(doc, "状态采用统一中文口径；duplicate与deprecated不进入36个活动模块完成度分母。")
    module_rows = []
    for m in MIGRATION["modules"]:
        module_rows.append([
            m.get("module_id"), m.get("module_name"), STATUS_CN.get(m.get("phase1_status"), m.get("phase1_status")),
            STATUS_CN.get(m.get("phase2_status"), m.get("phase2_status")), "是" if m.get("changed") else "否",
            "、".join(m.get("related_validation_ids", [])), m.get("remaining_gap", "")
        ])
    add_table(doc, "M001—M038模块状态", ["编号", "模块", "阶段一", "阶段二", "变化", "验证", "剩余缺口"], module_rows,
              widths=[1.4,4.1,2.8,2.8,1.3,2.0,4.1], compact=True,
              caption_source="阶段二M001—M038状态迁移")

    doc.add_page_break()
    add_title(doc, "附录B V001—V030验证结果摘要", 1)
    v_rows = []
    for r in V_DATA:
        v_rows.append([r[0], r[1], r[2], r[3], STATUS_CN.get(str(r[5]), r[5]), r[11], r[14], r[21]])
    add_table(doc, "V001—V030验证矩阵摘要", ["编号", "验证项", "优先级", "模块", "结论", "证据", "当前结果", "建议"], v_rows,
              widths=[1.3,4.0,1.4,2.3,2.3,2.1,4.4,3.1], compact=True,
              caption_source="阶段二V001—V030验证矩阵")

    doc.add_page_break()
    add_title(doc, "附录C M001—M038状态迁移摘要", 1)
    changed = [m for m in MIGRATION["modules"] if m.get("changed")]
    rows = [[m["module_id"],m["module_name"],STATUS_CN.get(m["phase1_status"],m["phase1_status"]),STATUS_CN.get(m["phase2_status"],m["phase2_status"]),m["change_reason"],m.get("new_runtime_evidence","")] for m in changed]
    add_table(doc, "状态发生变化的10个模块", ["编号", "模块", "阶段一", "阶段二", "修正原因", "新证据"], rows,
              widths=[1.4,3.4,2.6,2.6,6.2,3.0], compact=True,
              caption_source="阶段二M001—M038状态迁移")

    doc.add_page_break()
    add_title(doc, "附录D 测试和命令结果摘要", 1)
    add_table(doc, "Gate 0—5结果", ["Gate", "范围", "结果", "关键证据"], [
        ["0", "工作区冻结", "已验证完成", "4,855项Manifest、27入口哈希、patch、ZIP"],
        ["1", "环境/迁移/安全预检", "部分完成", "空库两次迁移至0012、58表、diff=0；安全有风险"],
        ["2", "PostgreSQL/Redis/Celery/FastAPI/Vite/Ollama", "部分完成", "最小栈成功；Ollama无模型；任务状态机未闭合"],
        ["3", "模型/契约/采集/预测", "部分完成", "指标与静态契约通过；写闭环安全停止"],
        ["4", "DeepSeek/Ollama/Router/AI/RAG/报告/策略", "部分完成", "DeepSeek可用；RAG超时/排序失败；100题未跑"],
        ["5", "测试/build/E2E/PySide6", "部分完成", "98/103安全子集；21文件收集错误；build通过；E2E阻塞"],
    ], widths=[1.3,4.6,2.8,8.0], caption_source="阶段二命令与测试日志；执行计划与安全边界")
    add_paragraph(doc, "命令级原始摘要未在本报告中重复粘贴；完整命令、耗时和安全停止原因见阶段二命令与测试日志。报告未记录任何密钥、密码、Token、连接串或完整敏感堆栈。")

    doc.add_page_break()
    add_title(doc, "附录E 证据和资料索引", 1)
    source_rows = [
        [path.name, "全文/对应章节", "相关M/V/P2E", "事实基线、运行验证或任务依据"]
        for path in SOURCES.values()
    ]
    add_table(doc, "报告使用的审计源文件", ["来源文件", "定位", "关联编号", "用途"], source_rows,
              widths=[7.2,3.0,3.0,5.0], compact=True)
    evidence_rows = []
    for e in PHASE2_EVIDENCE:
        evidence_rows.append([e["evidence_id"], e["category"], e["title"], e["related_validation_ids"], e["result"], e["confidence"]])
    add_table(doc, "阶段二证据更新索引", ["证据", "类别", "标题", "验证", "结果", "置信度"], evidence_rows,
              widths=[1.7,2.0,7.0,2.6,2.4,2.2], compact=True,
              caption_source="阶段二证据更新索引CSV")
    add_table(doc, "阶段一证据组索引", ["证据范围", "主题", "报告用途"], [
        ["E001—E009", "资产、Git、版本与入口", "项目背景、版本与可复现基线"],
        ["E010—E018", "采集、数据、数据库与质量", "数据链与事实源"],
        ["E019—E029", "特征、模型、预测与报告", "模型与预测历史证据"],
        ["E030—E039", "AI、Provider、工具和前端", "AI架构和阶段一静态状态"],
        ["E040—E048", "知识库、Embedding、RAG、评测", "历史RAG/AI基线"],
        ["E049—E058", "运行、部署、日志、文档与语法", "环境与工程边界"],
    ], widths=[3.5,6.2,7.0], caption_source="阶段一证据索引CSV")

    doc.add_page_break()
    add_title(doc, "附录F 术语表", 1)
    glossary = [
        ["Active模型", "已经通过准入验证并被唯一指定为当前生效、可用于正式预测的模型。"],
        ["Candidate模型", "等待验证或人工晋升的候选模型；存在文件和指标不等于Active。"],
        ["artifact", "一次训练或预测形成的文件化产物集合，如模型二进制、特征清单、指标和配置。"],
        ["run_id", "贯穿数据刷新、预测、报告、策略和任务的一次批次唯一标识。"],
        ["特征契约", "模型输入字段、顺序、类型、时区、缺失规则和版本的严格约定。"],
        ["事实源", "系统对某类业务事实认定的唯一可信来源及其身份规则。"],
        ["seed", "为了初始化或演示而插入的种子数据，不能冒充真实业务数据。"],
        ["demo", "明确用于演示的样例数据或能力。"],
        ["fallback", "主数据或主服务不可用时的降级结果；必须显式标识。"],
        ["Celery", "基于消息队列执行异步任务的Python任务框架。"],
        ["RAG", "先检索项目知识证据，再让模型基于证据生成回答的技术。"],
        ["Embedding", "把文本转成数值向量，用于语义相似度检索。"],
        ["reranker", "对初步检索结果进行二次排序的模型或规则。"],
        ["Provider", "向AI主链提供推理能力的模型服务，如DeepSeek或Ollama。"],
        ["Answer Guard", "回答输出前的安全、格式和事实边界门禁。"],
        ["幂等", "同一操作重复执行，结果不重复、不扩大副作用且保持一致。"],
        ["fail-closed", "缺少必要安全配置或身份时拒绝启动/拒绝执行，而不是自动放开。"],
    ]
    add_table(doc, "术语解释", ["术语", "解释"], glossary, widths=[4.0,12.3])


def build_docx():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_executive_summary(doc)
    build_body(doc)
    add_appendices(doc)
    core = doc.core_properties
    core.title = REPORT_TITLE
    core.subject = "项目状态基线、技术架构与优化实施规划"
    core.author = "Codex（基于阶段一和阶段二审计成果整理）"
    core.keywords = "智能运营分析, 项目审计, 技术架构, 优化规划, 20260714"
    core.comments = "未修改阶段一、阶段二原始审计文件。"
    doc.save(DOCX_PATH)

    metadata = {
        "project": PROJECT,
        "title": REPORT_TITLE,
        "baseline_date": BASELINE_DATE,
        "version": VERSION,
        "docx": str(DOCX_PATH),
        "pdf_expected": str(PDF_PATH),
        "figures": FIG_COUNTER,
        "tables": TABLE_COUNTER,
        "sources": [str(p.relative_to(ROOT)) for p in SOURCES.values()],
        "source_files_read": len(SOURCES),
        "phase2_evidence_rows": len(PHASE2_EVIDENCE),
        "phase1_evidence_rows": len(PHASE1_EVIDENCE),
        "v_rows": len(V_DATA),
        "module_rows": len(MIGRATION["modules"]),
        "core_conclusion": "开发演示；工程完成度64.6%；尚未达到可稳定本地运行。",
        "original_audit_files_modified": False,
    }
    META_PATH.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return metadata


if __name__ == "__main__":
    print(json.dumps(build_docx(), ensure_ascii=False, indent=2))
