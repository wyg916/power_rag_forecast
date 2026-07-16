from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "docs" / "《智能运营分析项目》项目全面分析说明文档_v20260703.docx"


def _font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    _font(r, 9, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc: Document, headers: list[str], rows: list[tuple], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, True)
        _shade(t.rows[0].cells[i], "D9EDE8")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            _cell(cells[i], value)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()


def h(doc: Document, level: int, text: str) -> None:
    doc.add_heading(text, level=level)


def p(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def note(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    _font(run, 9, None, "626F86")
    run.italic = True


def code_ref(doc: Document, paths: list[str]) -> None:
    para = doc.add_paragraph()
    run = para.add_run("代码依据：")
    _font(run, 9, True)
    for path in paths:
        r = para.add_run(f"\n- {path}")
        _font(r, 9)


TECH_ROWS = [
    ("前端", "React 18、TypeScript、Vite、Ant Design 5、ECharts / echarts-for-react", "frontend/package.json"),
    ("后端 Web", "FastAPI、Uvicorn、Pydantic、SQLAlchemy 2、psycopg", "requirements*.txt; backend/app/main.py"),
    ("数据库", "PostgreSQL，Alembic + SQL 迁移脚本，仓储层封装", "migrations/*; backend/app/repositories/*"),
    ("任务队列", "Redis + Celery，支持 celery/local_thread/auto 三种执行模式", "backend/app/workers/*; backend/app/services/task_runtime.py"),
    ("预测与建模", "pandas、numpy、scikit-learn、matplotlib、joblib，可选 XGBoost / LightGBM", "prediction_engine/*; requirements.txt"),
    ("RAG/Embedding", "本地 hash fallback、可选 sentence-transformers/BGE，Rerank 支持本地启发式/BGE", "rag_service.py; embedding_service.py; rerank_service.py"),
    ("LLM", "Ollama、本地 OpenAI compatible、DeepSeek Provider，LLMRouter 统一路由", "ai_assistant/llm_router.py; llm_providers/*; .env.example"),
    ("部署", "Docker Compose: postgres、redis、backend、celery_worker、frontend、可选 ollama profile", "docker-compose.yml; README_DEPLOY.md"),
    ("测试", "pytest + 前端构建契约测试 + AI/RAG 评估脚本", "tests/*; tests/evaluation/*"),
]

PATH_ROWS = [
    ("backend/app/main.py", "FastAPI 应用入口，创建 app、CORS、注册路由。"),
    ("backend/app/api/v1/endpoints", "后端 REST API 层，按业务模块拆分。"),
    ("backend/app/repositories", "PostgreSQL 仓储层，封装 forecast/task/model/report/settings/knowledge/user 等访问。"),
    ("backend/app/services", "业务服务层，包含 RAG、Embedding、Rerank、任务运行、系统设置、数据同步。"),
    ("backend/app/ai_assistant", "AI 助手新链路：意图、上下文、工具、LLM、Trace、模板、记忆。"),
    ("backend/app/workers", "Celery app、任务函数、投递器、任务命令映射。"),
    ("prediction_engine", "训练集构建、特征工程、训练、回测、正式预测。"),
    ("model_ops", "模型注册、对比、晋升、回滚、监控、误差记忆和偏差修正。"),
    ("migrations / migrations/versions", "幂等 SQL 与 Alembic 迁移，定义核心生产表。"),
    ("frontend/src", "React/Vite 前端，页面、服务、布局、主题和通用组件。"),
    ("tests", "pytest、评估脚本和产品化审计输出。"),
    ("docker-compose.yml", "本地/生产容器编排示例。"),
]

MODULE_ROWS = [
    ("首页/驾驶舱", "DashboardPage + dashboardApi；dashboard_home_service / forecast API", "汇总预测、策略、风险、报告、任务状态。"),
    ("数据中心", "DataCenterPage + dataApi；data.py / data_access / repositories", "数据库目录、字段、质量、SQL 查询、核心事实同步。"),
    ("预测中心", "ForecastCenterPage + forecastApi；prediction_engine；forecast.py", "24 小时电价预测、历史对比、风险/负荷/天气/新能源联动。"),
    ("策略中心", "StrategyCenterPage + strategyApi；strategy.py；AI tools/trading/tariff/storage", "高价风险、低价窗口、储能策略、人工复核。"),
    ("AI 助手", "AssistantPage + assistantApi；ai_assistant/service.py；llm_router；rag_service；tools", "业务问答、流式输出、RAG 引用、工具调用、反馈、Trace 折叠。"),
    ("报告中心", "ReportCenterPage + reportApi；report.py；report_repository", "报告列表、预览、下载、审核、发布、重新生成。"),
    ("模型中心", "ModelCenterPage + modelApi；model.py；model_ops；prediction_engine", "模型生命周期治理、Active/Candidate 对比、训练、回滚、特征与泄漏检查。"),
    ("知识库", "KnowledgeBasePage + knowledgeApi；knowledge.py；rag_service；embedding/rerank", "文档、索引、检索、QA 测试、RAG 健康。"),
    ("任务中心", "TaskCenterPage + taskApi；task.py；dispatcher；Celery tasks；task_business_handlers", "任务调度、运行日志、失败重试、队列状态、真实业务任务执行。"),
    ("系统设置", "SettingsPage/UserManagementPage + settingsApi/userApi；system.py；settings_center_service", "系统健康、接口配置、用户权限、安全策略、审计。"),
]

API_ROWS = [
    ("认证与用户", "/api/auth/*, /api/users/*", "登录、当前用户、退出、改密、用户增删改查、角色列表、重置密码", "auth.py; users.py"),
    ("系统设置", "/api/settings/*, /api/health, /api/db/health", "系统状态、运行参数、健康明细、用户权限、接口配置与测试日志", "system.py"),
    ("数据中心", "/api/data/*", "数据状态、目录、字段、质量、SQL 查询、表行数据、导入导出、核心数据同步", "data.py"),
    ("预测中心", "/api/forecast/*, /api/prediction/*, /api/market/*", "预测运行、最新预测、24h 预测、预测详情、市场历史、天气/负荷/新能源预测", "forecast.py"),
    ("策略中心", "/api/strategy/*, /api/anomaly/*", "策略生成、今日策略、配置、复核、异常解释", "strategy.py"),
    ("AI 助手", "/api/ai/*", "聊天、流式输出、附件、会话导出、反馈、会话列表、洞察、Trace", "assistant.py"),
    ("知识库", "/api/knowledge/*", "统计、文档列表、健康、搜索、QA 测试、批量校验、上传、本地索引、Embedding 刷新", "knowledge.py"),
    ("报告中心", "/api/reports/*", "报告生成、列表、概要、详情、下载、重新生成、审核、驳回、发布、审核记录", "report.py"),
    ("模型中心", "/api/models/*, /api/model/*", "Active/Candidate、指标、误差、重训建议、治理总览、训练、激活、回滚、导出、特征与泄漏检查", "model.py"),
    ("任务中心", "/api/tasks/*, /api/scheduled-tasks/*", "任务运行、概览、列表、健康、日志、重试、队列、趋势、启动、取消、计划任务", "task.py"),
]

API_DETAIL_ROWS = [
    ("POST /api/ai/chat", "非流式 AI 问答，返回 answer/evidence/trace 等。"),
    ("POST /api/ai/chat/stream", "SSE 流式问答，前端 assistantApi.askAssistantStream 消费。"),
    ("GET /api/ai/chat/sessions", "AI 会话列表，支撑左侧会话栏。"),
    ("GET /api/knowledge/stats / documents / health", "知识库指标、文档表、RAG 状态。"),
    ("POST /api/knowledge/qa-test", "RAG 检索测试与整理答案。"),
    ("GET /api/models/center/overview", "模型中心一体化总览。"),
    ("POST /api/models/center/training/start", "启动模型训练任务。"),
    ("POST /api/models/center/activate / rollback", "模型激活与回滚。"),
    ("GET /api/tasks/health", "Redis/Celery/执行模式健康状态。"),
    ("POST /api/tasks/start", "按 task kind 投递业务任务。"),
    ("GET /api/reports / summary / latest", "报告中心列表、统计、最新报告。"),
    ("POST /api/reports/{id}/approve / reject / publish", "审核与发布状态流转。"),
    ("GET /api/settings/status/overview", "系统设置状态总览。"),
    ("GET /api/settings/interfaces/configs", "接口配置卡片数据。"),
    ("POST /api/settings/interfaces/test-all", "批量测试接口配置。"),
]

DB_ROWS = [
    ("预测与模型", "forecast_runs, forecast_results, model_versions, model_metrics, model_registry, model_evaluation_runs, model_performance_daily, model_retrain_jobs, model_comparison_runs, model_governance_events, model_prediction_comparison_points"),
    ("任务运行", "task_runs, task_logs，以及 task_runs/task_logs 的多轮运行态、超时、重试、取消字段扩展"),
    ("AI 与 RAG", "ai_traces, ai_chat_sessions, ai_chat_messages, ai_tool_call_logs, ai_prompt_templates, ai_chat_feedback, ai_answer_feedback, ai_conversation_state, ai_qa_test_cases, kb_documents, kb_chunks, kb_search_results, kb_qa_tests"),
    ("报告与审核", "report_runs, report_reviews, report_approval_runs, ai_report_review_runs"),
    ("数据源与业务规则", "raw_market, raw_weather, raw_load, raw_renewable, feature_importance, pv_tariff_rules, pv_policy_files, pv_station_tariff_check, market_power_price_rules, southern_grid_tax_rules, pv_tariff_period_rules, prediction_tracking"),
    ("权限与审计", "users, roles, user_roles, audit_logs"),
    ("系统设置", "system_runtime_config, system_health_snapshots, system_api_configs, system_api_test_logs, monitoring_metrics"),
]

RISK_ROWS = [
    ("高", "配置与密钥治理", "存在较多环境变量和可选 provider，生产必须使用 AUTH_REQUIRED=1、强 JWT_SECRET_KEY、禁用开发 fallback。", "使用 .env 模板与部署检查脚本固化生产必填项；启动时对弱配置 fail-fast。"),
    ("高", "数据库来源一致性", "代码同时保留 legacy Excel/脚本资产和 PostgreSQL 主库路径，部分接口有 fallback 能力。", "明确生产环境 DATABASE_ALLOW_LEGACY_FALLBACK=0；将页面数据来源统一落库并通过 API 展示。"),
    ("中高", "RAG 质量与 Embedding 维度", "本地 hash fallback 可保障可用性，但语义效果低于真实 BGE；Docker 配置期望 1024 维。", "生产部署实际 BGE embedding/rerank 模型；建立索引维度一致性校验和批量 QA 评估。"),
    ("中高", "任务在线消费链路", "支持 Celery 与 local_thread；Celery 依赖 Redis/Worker 可用性。", "生产固定 TASK_EXECUTION_MODE=celery，/api/tasks/health 纳入监控，失败重试必须记录 error_reason。"),
    ("中", "模型训练与晋升治理", "已有模型注册、比较、回滚、准入规则，但需结合真实数据持续校准阈值。", "补齐训练样本审计、特征版本、数据漂移、候选模型灰度准入。"),
    ("中", "前端页面密度与接口状态", "近期大量 UI 重构后页面结构统一，但不同页面仍依赖各自适配器。", "沉淀统一卡片/表格/状态组件；保留接口异常态，不回退为前端假数据。"),
    ("中", "测试与验收覆盖", "测试文件多，但未在本次文档生成中全量执行。", "建立 CI：后端 pytest、前端 npm run build、迁移检查、RAG smoke、Celery smoke。"),
]

TEST_ROWS = [
    ("认证/RBAC", "test_auth_*.py, test_user_management_*.py, test_frontend_auth_contract.py", "JWT、登录、权限、用户管理、审计。"),
    ("AI/RAG", "test_ai_*.py, test_rag_hybrid.py, test_ai_rag_policy_ranking.py, evaluation/*", "问答契约、可读性、证据泄漏、LLM 路由、RAG 排序与评估。"),
    ("任务中心", "test_task_*.py, test_p4_task_*.py, productization/check_task_center_runtime.py", "任务仓储、调度、健康、日志、重试、取消、超时。"),
    ("模型/预测", "test_p2_*.py, test_model_*.py, test_formal_forecast.py", "数据集、特征、泄漏检查、回测、模型中心 API。"),
    ("知识库", "test_knowledge_*.py", "知识库业务 API 和前端布局契约。"),
    ("系统设置", "test_settings_center_business_api.py", "系统设置三 Tab 后端闭环。"),
    ("迁移/部署", "test_migrations.py, test_startup_scripts.py", "数据库迁移和启动脚本可用性。"),
]


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "微软雅黑"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        styles[name].font.color.rgb = RGBColor(15, 35, 62)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("《智能运营分析项目》项目全面分析说明文档")
    _font(run, 22, True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(subtitle.add_run("面向 AI 售电交易决策平台的代码级架构、数据、AI/RAG、模型、任务和前端分析"), 11)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(meta.add_run("生成日期：2026-07-03    生成方式：基于当前项目代码静态扫描与关键文件定向阅读"), 10)
    doc.add_paragraph()
    note(doc, "说明：本文档中的事实结论以当前工作区代码、配置、迁移、前端路由、后端 API 和测试文件为依据；无法直接由代码证明的内容均标注为“根据代码推断”。本次未连接生产数据库、未执行全量测试。")
    doc.add_page_break()

    h(doc, 1, "目录")
    chapters = [
        "项目概述", "业务目标", "技术栈", "整体架构", "目录结构", "功能模块", "核心代码", "启动运行",
        "数据来源与处理", "分析 / 特征 / 建模", "模型算法", "AI 大模型接入", "API 接口", "数据库与存储",
        "配置与环境变量", "部署与运行", "完成度评估", "问题与风险", "优化建议", "总结",
    ]
    for index, chapter in enumerate(chapters, 1):
        p(doc, f"{index}. {chapter}")
    doc.add_page_break()

    h(doc, 1, "1. 项目概述")
    p(doc, "本项目当前代码体现为“AI 售电交易决策平台 / 智能运营分析项目”：以电力交易、价格预测、策略辅助、报告管理、模型治理、RAG 知识库和任务调度为核心的企业级后台系统。")
    p(doc, "从代码结构看，项目已经从早期离线脚本与桌面端资产，演进为 FastAPI + React/Vite + PostgreSQL + Celery/Redis 的 Web 平台，并保留大量预测引擎、模型管控、RAG 评估和历史脚本。")
    p(doc, "根据代码推断，系统定位是为售电交易、虚拟电厂/新能源聚合、策略研判和运营管理提供一体化工作台。")
    code_ref(doc, ["backend/app/main.py", "frontend/src/app/router.tsx", "README_WEB平台说明.md", "README_DEPLOY.md"])
    table(doc, ["维度", "结论"], [
        ("产品名称", "AI 售电交易决策平台 / 智能运营分析项目"),
        ("核心用户", "交易运营、分析师、模型工程、管理员、审核/发布人员"),
        ("核心能力", "数据治理、预测、策略、AI 问答、报告、模型治理、知识库、任务调度、系统设置"),
        ("代码形态", "Web 平台 + 后端服务 + 预测/模型 Python 包 + Celery Worker + PostgreSQL 迁移 + 历史脚本/桌面资产"),
    ])

    h(doc, 1, "2. 业务目标")
    bullets(doc, [
        "统一展示电力交易相关数据、预测结果、策略建议、风险提示和报告状态。",
        "通过预测引擎支撑未来 24 小时价格、负荷、新能源等指标的业务分析。",
        "通过 AI 助手将查数问答、RAG 知识引用、工具调用、模型/报告/策略解释整合到业务工作台。",
        "通过模型中心实现 Active / Candidate 生命周期治理、误差趋势、训练、准入、回滚。",
        "通过任务中心打通价格预测、数据同步、日报生成等后台任务执行闭环。",
        "通过系统设置、用户权限、接口配置、审计日志提高企业级运维与安全可管理性。",
    ])
    p(doc, "根据代码推断，业务边界并不是单一“聊天助手”或“预测脚本”，而是围绕售电交易决策形成的数据、模型、知识、任务、报告和权限的闭环平台。")

    h(doc, 1, "3. 技术栈")
    table(doc, ["层级", "主要技术", "代码依据"], TECH_ROWS, [3, 8, 6])

    h(doc, 1, "4. 整体架构")
    p(doc, "项目可分为前端展示层、API 层、业务服务层、仓储层、后台任务层、模型/预测引擎、AI/RAG 层和数据库/外部服务层。")
    table(doc, ["架构层", "职责", "关键目录/文件"], [
        ("前端展示层", "页面布局、交互、图表、表格、状态展示、API 适配", "frontend/src/pages, frontend/src/services, frontend/src/api.ts"),
        ("API 层", "按业务模块暴露 REST 接口，统一进入 api_router", "backend/app/api/v1/router.py, endpoints/*"),
        ("业务服务层", "RAG、任务运行、设置中心、数据同步、UI 平台服务等", "backend/app/services/*"),
        ("仓储层", "封装 PostgreSQL 表访问与写入", "backend/app/repositories/*"),
        ("任务层", "任务投递、Celery worker、local_thread fallback、任务日志与状态", "backend/app/workers/*, backend/app/services/task_runtime.py"),
        ("预测/模型层", "训练集、特征、训练、回测、正式预测、模型注册/晋升/回滚", "prediction_engine/*, model_ops/*"),
        ("AI/RAG 层", "意图识别、上下文、工具调用、LLM 路由、RAG 检索、Trace/反馈", "backend/app/ai_assistant/*, backend/app/services/rag_service.py"),
        ("数据与外部依赖", "PostgreSQL、Redis、Ollama/DeepSeek、BGE/Milvus 或本地索引、文件知识库", "docker-compose.yml, .env.example, migrations/*"),
    ], [3, 6, 7])
    note(doc, "根据代码推断，当前系统以 PostgreSQL 为目标主数据源，但保留历史 Excel/本地文件读取逻辑用于兼容和 fallback。")

    h(doc, 1, "5. 目录结构")
    table(doc, ["路径", "作用"], PATH_ROWS, [5, 11])
    p(doc, "根目录还保留若干批处理启动脚本、历史数据处理脚本、知识库目录、模型产物目录和输出目录。它们对理解项目历史演进有价值，但生产运行主链路应以 backend、frontend、migrations、prediction_engine、model_ops 和 tests 为核心。")

    h(doc, 1, "6. 功能模块")
    table(doc, ["模块", "主要代码路径", "功能说明"], MODULE_ROWS, [3, 6, 8])

    h(doc, 1, "7. 核心代码")
    table(doc, ["文件/目录", "关键作用"], [
        ("backend/app/main.py", "create_app 创建 FastAPI，配置 CORS，注册 model gateway 和 api_router。"),
        ("backend/app/api/v1/router.py", "统一注册 auth/system/data/forecast/strategy/assistant/knowledge/report/model/task/users。"),
        ("backend/app/core/config.py", "读取 DATABASE_URL、REDIS_URL、Celery、LLM、RAG、AUTH_REQUIRED 等环境配置。"),
        ("backend/app/db/session.py", "创建 SQLAlchemy engine/session，DATABASE_URL 未配置时直接报错。"),
        ("backend/app/data_access.py", "封装 Excel/JSON/数据库状态、预测摘要、业务摘要、数据库表查询等兼容访问。"),
        ("backend/app/ai_assistant/service.py", "AI 助手主流程：问题标准化、意图路由、工具调用、RAG/LLM 决策、答案构建、Trace/会话状态保存。"),
        ("backend/app/services/rag_service.py", "RAG 检索主实现：keyword/vector/rerank/fallback file search、profile、cache、domain boost。"),
        ("backend/app/workers/dispatcher.py", "根据 TASK_EXECUTION_MODE 在 celery/local_thread/auto 之间投递任务并返回健康状态。"),
        ("backend/app/workers/tasks.py", "Celery 任务函数，包括命令任务、知识导入、embedding 刷新、报告生成、价格预测、数据同步、日报。"),
        ("backend/app/services/task_business_handlers.py", "price_predict/data_sync/report_daily 真实业务任务写库逻辑。"),
        ("prediction_engine/dataset_builder.py", "PostgreSQL/legacy 数据读取、训练集组装、时间切分、特征 schema 和泄漏检查。"),
        ("prediction_engine/feature_builder.py", "时间、lag、rolling、interaction 等特征构建与元数据输出。"),
        ("model_ops/promotion_manager.py", "候选模型晋升为 Active 与最新 Candidate 查询。"),
        ("frontend/src/api.ts", "前端统一请求层、鉴权 token、文件下载和各业务 API 封装。"),
    ], [5, 11])

    h(doc, 1, "8. 启动运行")
    numbered(doc, [
        "本地后端：根据 README 可使用 `python -X utf8 -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000`。",
        "本地前端：进入 frontend 后执行 `npm run dev`，默认 127.0.0.1:5173。",
        "平台脚本：根目录提供 run_project.bat、run_web_platform.bat、run_web_backend.bat、run_web_frontend.bat、run_celery_worker.bat 等。",
        "容器部署：README_DEPLOY 与 docker-compose.yml 支持 postgres、redis、backend、celery_worker、frontend、可选 ollama。",
        "数据库迁移：可使用 Alembic 或项目迁移脚本，migrations/README.md 中给出 apply_database_migrations 方式。",
    ])
    note(doc, "本次文档生成没有实际启动服务或连接数据库；运行方式根据 README、docker-compose 和脚本文件整理。")

    h(doc, 1, "9. 数据来源与处理")
    p(doc, "项目的数据来源包括 PostgreSQL 中的标准化业务表、历史 Excel/JSON/本地文件、知识库文档、预测/模型产物以及外部 LLM/RAG 服务。")
    bullets(doc, [
        "PostgreSQL：通过 backend/app/repositories 和 SQLAlchemy 访问，是当前 Web 平台与生产化闭环的目标主数据源。",
        "历史文件：data_access.py、prediction_engine 和根目录历史脚本保留 Excel/JSON/本地文件读取能力。",
        "核心事实同步：core_data_sync.py 与任务业务处理器可将预测、市场、负荷、规则等写入核心事实表。",
        "知识库：knowledge_pipeline、kb_documents/kb_chunks、RAG 索引与 embedding 刷新支撑文档检索。",
        "任务产物：forecast_runs/forecast_results、report_runs、task_runs/task_logs 构成任务执行结果闭环。",
    ])

    h(doc, 1, "10. 分析 / 特征 / 建模")
    p(doc, "prediction_engine 是分析与建模主目录，围绕时间序列预测、特征构建、回测与正式预测组织。")
    table(doc, ["文件", "功能"], [
        ("dataset_builder.py", "读取 PostgreSQL/legacy 数据，组装训练集，按时间切分，生成特征 schema，执行 leakage_check。"),
        ("feature_builder.py / features.py", "时间特征、滞后特征、滚动窗口特征、EWMA、业务交互特征、特征元数据。"),
        ("baseline_backtest.py / backtest_runner.py", "基线模型、历史预测对比、回归指标、峰值误差、滚动回测。"),
        ("model_trainer.py", "训练管线、特征版本、模型产物保存与加载。"),
        ("formal_forecast.py / fast_forecast.py", "正式前向预测、残差区间、图表和业务摘要输出。"),
        ("evaluation.py", "MAE/RMSE/MAPE、异常识别、滚动评估等通用评估函数。"),
    ], [5, 11])
    p(doc, "根据代码推断，项目在 P2 阶段重点治理了特征安全与数据泄漏问题，相关测试包括 test_p2_schema_guard.py、test_p2_leakage_checker.py、test_p2_feature_schema.py 等。")

    h(doc, 1, "11. 模型算法")
    p(doc, "依赖层面包含 scikit-learn，并可选 xgboost、lightgbm。代码中模型训练/注册/对比/回滚被拆分在 prediction_engine 与 model_ops 中。")
    table(doc, ["模型治理能力", "代码依据", "说明"], [
        ("模型注册", "model_ops/model_registry.py", "从产物 manifest 与指标中注册模型版本。"),
        ("Active/Candidate 晋升", "model_ops/promotion_manager.py", "支持候选模型设置 Active、查询最新候选。"),
        ("模型对比", "model_ops/model_comparator.py", "对比 Candidate 与 Active 指标。"),
        ("模型监控", "model_ops/model_monitor.py", "summarize_error_window、decide_retrain、evaluate_model_health。"),
        ("误差记忆", "model_ops/error_memory.py", "构建/更新/加载误差记忆，并为预测行提供误差 profile。"),
        ("偏差校正", "model_ops/bias_corrector.py", "基于误差 profile 计算 bias adjustment 并添加 correction columns。"),
    ], [4, 5, 7])
    note(doc, "未在本次扫描中确认生产环境具体最终模型文件和训练数据窗口；实际模型算法组合需结合 model_artifacts 与训练日志进一步核查。")

    h(doc, 1, "12. AI 大模型接入")
    p(doc, "AI 助手采用“意图识别 + 工具调用 + RAG + LLM 路由 + 答案保护 + Trace/反馈”的混合架构。")
    table(doc, ["子能力", "实现位置", "说明"], [
        ("LLM 路由", "backend/app/ai_assistant/llm_router.py", "支持 Ollama、DeepSeek、OpenAI compatible 风格的本地/远程模型路由。"),
        ("Provider", "backend/app/ai_assistant/llm_providers/*", "DeepSeekProvider、OllamaProvider。"),
        ("意图与上下文", "core/intent_router.py, context_resolver.py, input_normalizer.py", "识别预测、策略、知识、查数等意图，并处理追问上下文。"),
        ("工具调用", "ai_assistant/tools/*", "forecast、knowledge、model、report、storage、tariff、trading、weather、time、data freshness 等工具。"),
        ("RAG 检索", "services/rag_service.py", "关键词 + 向量 + rerank + fallback file search + cache + domain boost。"),
        ("Embedding/Rerank", "embedding_service.py, rerank_service.py", "local hash fallback；可选 sentence-transformers/BGE。"),
        ("答案保护", "core/answer_guard.py, service.sanitize_answer_for_display", "防止策略边界越界、证据泄漏、推理痕迹暴露。"),
        ("反馈与 Trace", "assistant.py, ai_traces, ai_chat_feedback, ai_answer_feedback", "支持会话、反馈、Trace 查询和开发者折叠。"),
    ], [4, 5, 7])
    p(doc, "环境变量中 AI_ASSISTANT_LLM_ENABLED、LLM_PROVIDER、LLM_BASE_URL、LLM_MODEL、DEEPSEEK_*、RAG_ENABLED、RAG_PROFILE、RAG_TOP_K、RAG_EMBEDDING_*、RAG_RERANK_* 控制大模型和 RAG 行为。")

    h(doc, 1, "13. API 接口")
    table(doc, ["模块", "主要路径", "能力", "代码依据"], API_ROWS, [3, 4, 7, 5])
    h(doc, 2, "重点业务接口示例")
    table(doc, ["接口", "用途"], API_DETAIL_ROWS, [6, 10])
    p(doc, "前端通过 frontend/src/api.ts 暴露 api 对象，再由 services 层进行页面级聚合，例如 assistantApi、taskApi、modelApi、knowledgeApi、reportApi、settingsApi。")

    h(doc, 1, "14. 数据库与存储")
    table(doc, ["领域", "相关表"], DB_ROWS, [3, 13])
    p(doc, "数据库迁移同时存在 migrations/*.sql 与 migrations/versions/*.py。README 说明迁移脚本应保持幂等，重复执行不能破坏已有表。")
    note(doc, "本次没有直接连接数据库确认实际表存在和行数；上述表来自迁移文件静态扫描。")

    h(doc, 1, "15. 配置与环境变量")
    p(doc, "核心配置集中在 .env.example、.env.docker.example、docker-compose.yml 与 backend/app/core/config.py。")
    table(doc, ["类别", "关键变量", "含义"], [
        ("数据库", "DATABASE_URL, POSTGRES_DB, POSTGRES_USER, POSTGRES_PASSWORD, POSTGRES_PORT", "PostgreSQL 连接与容器端口。"),
        ("Redis/Celery", "REDIS_URL, CELERY_BROKER_URL, CELERY_RESULT_BACKEND, TASK_EXECUTION_MODE", "队列、结果后端和任务执行模式。"),
        ("认证", "AUTH_REQUIRED, JWT_SECRET_KEY", "是否强制鉴权和 JWT 签名密钥。"),
        ("LLM", "LLM_PROVIDER, LLM_ROUTER_MODE, LLM_BASE_URL, LLM_MODEL, LLM_API_KEY, DEEPSEEK_*", "本地/远程大模型配置。"),
        ("RAG", "RAG_ENABLED, RAG_PROFILE, RAG_TOP_K, RAG_EMBEDDING_*, RAG_RERANK_*", "RAG 开关、检索规模、embedding/rerank provider 与模型路径。"),
        ("前端", "VITE_API_BASE_URL, VITE_API_PROXY_TARGET, VITE_AUTH_REQUIRED", "前端 API 地址、代理目标和鉴权开关。"),
    ], [3, 6, 7])
    p(doc, "生产建议在 README_DEPLOY 和 .env.example 中已有提示：APP_ENV=production、AUTH_REQUIRED=1、DATABASE_ALLOW_LEGACY_FALLBACK=0、TASK_EXECUTION_MODE=celery。")

    h(doc, 1, "16. 部署与运行")
    table(doc, ["场景", "方式", "注意事项"], [
        ("本地开发", "后端 uvicorn + 前端 npm run dev", "需配置 DATABASE_URL；前端默认 5173，后端默认 8000。"),
        ("后台任务", "run_celery_worker.bat 或 celery worker", "Redis 必须可用；TASK_EXECUTION_MODE=celery 时 Worker 不在线会失败。"),
        ("Docker Compose", "docker compose --env-file .env.docker up -d --build", "服务包括 postgres/redis/backend/celery_worker/frontend；RAG 模型路径可通过 volume 挂载。"),
        ("健康检查", "/health, /api/db/health, /api/tasks/health", "可用于确认 Web、数据库、Redis/Celery 状态。"),
        ("构建验证", "frontend: npm run build；后端: pytest/compileall", "本次未全量执行，仅在文档中基于项目说明列出。"),
    ], [3, 6, 7])

    h(doc, 1, "17. 完成度评估")
    table(doc, ["领域", "当前完成度判断", "依据"], [
        ("前端页面", "较高：主要业务页面均有 React 页面和服务层，且近期已重构为统一后台风格。", "frontend/src/pages, frontend/src/services, router.tsx"),
        ("后端 API", "较高：十大模块均有 FastAPI endpoints，覆盖列表/详情/操作/健康/导出等。", "backend/app/api/v1/endpoints/*"),
        ("数据库闭环", "中高：核心迁移已覆盖预测、任务、AI、报告、知识库、权限、设置；实际数据完整性需连接数据库验收。", "migrations/*"),
        ("任务执行链路", "中高：支持 Celery/local_thread，已实现 price_predict、data_sync、report_daily 等业务任务。", "workers/tasks.py; task_business_handlers.py"),
        ("AI/RAG", "中高：链路完整，含 RAG、LLM、工具、Trace、反馈；生产质量取决于 embedding/rerank 与知识索引。", "ai_assistant/*; rag_service.py; tests/evaluation/*"),
        ("模型治理", "中高：模型中心 API、表、前端、训练/激活/回滚已具备；需持续真实训练数据验证。", "model.py; model_ops/*; migrations/0009/0010"),
        ("测试与审计", "中高：测试文件覆盖多阶段主题；需 CI 固化和本地/容器 smoke 定期执行。", "tests/*"),
    ], [3, 8, 6])

    h(doc, 1, "18. 问题与风险")
    table(doc, ["级别", "风险点", "说明", "建议处理"], RISK_ROWS, [2, 4, 7, 6])

    h(doc, 1, "19. 优化建议")
    numbered(doc, [
        "建立“一键环境验收”脚本：检查 DATABASE_URL、Redis、Celery worker、迁移版本、后端健康、前端构建、关键 API smoke。",
        "将所有前端页面的数据来源标签和 fallback 规则统一到服务层，禁止页面直接硬编码业务假数据。",
        "为 AI 助手建立固定评测集：覆盖电价、负荷、新能源、策略、政策、报告、数据库查数和追问。",
        "对 RAG 建立索引版本管理：记录 embedding 模型、维度、chunk 策略、rerank 模型、索引时间、QA 通过率。",
        "模型中心补齐训练审计：训练样本窗口、特征版本、参数、数据质量、泄漏检查结果、准入规则和回滚依据。",
        "任务中心补齐生产监控：队列堆积、Worker 心跳、失败原因分类、重试次数、超时策略、幂等键冲突。",
        "系统设置页面将接口测试、用户操作、配置修改全部写 audit_logs，便于合规追踪。",
        "将历史脚本、桌面端资产、Web 生产链路分层归档，减少新开发误用 legacy 数据源的风险。",
        "补充 API OpenAPI 文档导出和前后端契约测试，避免字段变更造成页面空态或错映射。",
        "引入 CI：pytest 分组、frontend npm run build、迁移幂等检查、RAG smoke、Celery smoke。",
    ])

    h(doc, 1, "20. 总结")
    p(doc, "当前项目已经具备企业级 AI 售电交易决策平台的主体架构：前端页面、FastAPI 后端、PostgreSQL 表结构、Celery/Redis 任务、预测与模型治理、RAG/LLM 问答、报告审核发布、系统设置和用户权限均有实际代码支撑。")
    p(doc, "从架构演进看，项目正在从“脚本 + 文件 + 局部页面”过渡到“数据库主线 + API + 前端工作台 + 后台任务 + AI/RAG/模型治理”的平台化阶段。")
    p(doc, "后续最重要的不是继续扩展静态界面，而是持续打通真实业务数据闭环、生产配置、任务在线消费、模型训练审计和 AI/RAG 质量评测。")
    note(doc, "本文件为代码级分析文档，不替代生产安全审计、数据库现场验收和模型效果评测报告。")

    h(doc, 1, "附录 A：前端结构概览")
    table(doc, ["类别", "路径", "说明"], [
        ("路由/菜单", "frontend/src/app/router.tsx", "定义首页、数据中心、预测中心、策略中心、AI 助手、报告中心、模型中心、知识库、任务中心、系统设置等一级入口。"),
        ("全局骨架", "BasicLayout/HeaderBar/Sidebar/PageContainer/styles/theme", "顶部全局栏、左侧导航、页面容器、主题 token。"),
        ("统一 API", "frontend/src/api.ts", "封装 request、requestForm、requestBlob、token、auth、各业务 API。"),
        ("页面服务层", "frontend/src/services/*.ts", "每页聚合后端 API，做状态、字段、fallback 边界适配。"),
        ("页面组件", "frontend/src/pages/*", "各业务页面主体，近期已大量改造成企业级 SaaS 后台工作台。"),
    ], [3, 5, 8])

    h(doc, 1, "附录 B：测试覆盖概览")
    table(doc, ["测试领域", "代表文件", "覆盖重点"], TEST_ROWS, [3, 6, 7])

    h(doc, 1, "附录 C：本次读取/扫描依据")
    bullets(doc, [
        "根目录结构、rg --files 文件索引、requirements.txt、requirements.web.txt、frontend/package.json。",
        "backend/app/main.py、backend/app/api/v1/router.py、backend/app/api/v1/endpoints/*。",
        "backend/app/core/config.py、backend/app/db/session.py、backend/app/data_access.py。",
        "backend/app/ai_assistant/service.py、llm_router.py、tools、backend/app/services/rag_service.py、embedding_service.py、rerank_service.py。",
        "backend/app/workers/tasks.py、dispatcher.py、celery_app.py、services/task_runtime.py、task_business_handlers.py。",
        "prediction_engine/*、model_ops/*、migrations/*、migrations/versions/*。",
        "frontend/src/api.ts、frontend/src/app/router.tsx、frontend/src/pages/*、frontend/src/services/*。",
        "tests/*、tests/evaluation/*、tests/productization/output/* 的文件索引。",
    ])
    p(doc, f"文档生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}；输出文件：{OUTPUT}")

    for para in doc.paragraphs:
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.15
        for run in para.runs:
            _font(run)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
