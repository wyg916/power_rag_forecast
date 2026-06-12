from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\智能运营分析项目")
OUT = ROOT / "《智能运营分析项目》项目全面分析说明文档_v20260610.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def add_table(doc: Document, headers: list[str], rows: list[tuple[object, ...]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True)
        set_cell_shading(hdr[i], "D9EAF7")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_kv_table(doc: Document, rows: list[tuple[object, object]]) -> None:
    add_table(doc, ["项目", "内容"], rows, widths=[1.8, 5.8])


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def style_doc(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.2
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.color.rgb = RGBColor(18, 52, 86)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11)


def build_doc() -> Path:
    doc = Document()
    style_doc(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智能运营分析项目\n项目全面分析说明文档")
    r.bold = True
    r.font.size = Pt(24)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("电力市场电价预测与智能分析平台\n")
    r.font.size = Pt(14)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r = p.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n项目路径：{ROOT}")
    r.font.size = Pt(10)

    add_kv_table(
        doc,
        [
            ("文档用途", "面向开发、项目负责人、面试官和团队成员的项目全景说明，覆盖业务、架构、代码、数据、模型、AI、部署、风险和后续规划。"),
            ("分析依据", "基于当前项目本地代码、配置示例、数据库迁移、前端页面、知识库管道、测试报告和部署报告进行静态分析；未读取 .env、密钥、模型权重和大体积二进制文件。"),
            ("重要说明", "涉及未实际运行或无法从代码完全确认的内容，文档中以“根据代码推断”标注；未发现的功能明确写明“当前项目中未发现”。"),
        ],
    )
    doc.add_page_break()

    headings = [
        "项目概述",
        "项目业务目标与应用场景",
        "项目技术栈",
        "项目整体架构",
        "项目目录结构说明",
        "核心功能模块分析",
        "核心代码文件说明",
        "项目启动与运行流程",
        "数据来源与数据处理流程",
        "数据分析、特征工程与建模流程",
        "模型与算法逻辑说明",
        "AI 大模型接入与调用逻辑",
        "接口与 API 调用说明",
        "数据库与存储设计",
        "配置文件与环境变量说明",
        "项目部署与运行方式",
        "当前项目完成度评估",
        "存在问题与风险点",
        "优化建议与后续开发方案",
        "总结",
    ]
    doc.add_heading("目录", level=1)
    add_numbered(doc, headings)
    doc.add_paragraph("说明：此目录为章节清单，Word 中可按需插入自动目录。")
    doc.add_page_break()

    doc.add_heading("1. 项目概述", level=1)
    add_kv_table(
        doc,
        [
            ("项目名称", "智能运营分析项目 / AI 售电交易决策平台"),
            ("项目类型", "Python 数据工程与机器学习项目 + FastAPI 后端 + React/Vite 前端 + PostgreSQL/Redis/Celery 产品化平台 + AI/RAG 智能助手。"),
            ("项目定位", "面向电力市场、售电交易和电价预测场景的智能运营分析平台，重点服务日前/实时电价分析、负荷天气影响解释、尖峰风险识别、模型运维、报告生成和 AI 问答。"),
            ("项目解决的问题", "将 PJM/DOM/LMP 等电力市场数据、负荷、天气和预测模型结果统一沉淀到数据库与工作台，并通过 AI 助手把预测数据、知识库和业务规则转化为可读的运营建议。"),
            ("目标用户", "售电公司交易员、运营分析师、模型运维人员、管理层、系统开发/运维人员。"),
            ("当前完成度判断", "已达到“准产品化工程阶段”：核心预测、RAG AI 助手、PostgreSQL 数据骨架、任务中心、JWT/RBAC、审计脱敏、Docker Compose 部署均已具备；仍需继续完善生产密钥治理、长期任务稳定性、外部数据全链路 SLA、模型治理和企业级运维。"),
            ("主要运行方式", "本地脚本运行、FastAPI + Vite 本地开发、Docker Compose 一键启动、Celery/Redis 异步任务、PostgreSQL 主事实源。"),
        ],
    )

    doc.add_heading("2. 项目业务目标与应用场景", level=1)
    doc.add_paragraph("项目以电力市场运营分析为核心，业务目标不是单纯展示电价曲线，而是形成“数据获取 -> 特征工程 -> 预测建模 -> 风险解释 -> 策略建议 -> 报告输出 -> AI 助手问答”的闭环。")
    add_table(
        doc,
        ["应用场景", "用户问题", "系统能力", "主要产出"],
        [
            ("日前/实时电价分析", "明天浙江省电力供需如何？晚高峰是否有缺口风险？", "读取预测结果、负荷、天气、市场规则和历史指标，结合 AI 助手解释风险。", "风险判断、重点时段、建议关注指标。"),
            ("售电交易策略", "峰谷价差变大时应关注什么？", "结合价格预测、尖峰概率、储能约束和知识库规则生成建议。", "购电/储能/风险边界建议。"),
            ("模型运维", "模型误差变大是否需要重训？", "查看模型指标、预测误差、特征重要性和数据新鲜度。", "误差归因、重训建议、检查项。"),
            ("知识库问答", "LMP、绿证、代理购电、现货出清是什么意思？", "BGE embedding + reranker 的 RAG 检索，引用知识片段解释。", "专业术语解释和业务影响。"),
            ("报告生成", "生成日报/策略报告。", "报告接口和异步任务中心生成报告，并保留审核流。", "报告草稿、审核记录、发布状态。"),
        ],
    )

    doc.add_heading("3. 项目技术栈", level=1)
    add_table(
        doc,
        ["层级", "技术/依赖", "说明"],
        [
            ("后端框架", "FastAPI、Uvicorn、Pydantic", "backend/app/main.py 创建应用，api_router 注册各业务路由。"),
            ("数据库", "PostgreSQL、SQLAlchemy、Alembic、psycopg", "PostgreSQL 是主事实源；Alembic 迁移覆盖预测、模型、任务、知识库、RBAC、审计、Stage1 原始数据等表。"),
            ("异步任务", "Celery、Redis、本地后台线程 fallback", "任务中心支持 auto/celery/local_thread 三种执行模式。"),
            ("前端", "React 18、Vite 5、TypeScript、Ant Design、ECharts", "工作台包括首页、数据中心、预测中心、AI 助手、报告、模型、知识库、任务中心、系统设置/用户管理。"),
            ("数据处理", "pandas、numpy、openpyxl、requests", "用于 PJM/Open-Meteo 数据获取、清洗、合并、Excel/CSV 兼容和 PostgreSQL 导入。"),
            ("机器学习", "scikit-learn、XGBoost、LightGBM、CatBoost、joblib", "根据环境可用性组合训练回归、分类、尖峰识别和预测区间模型。"),
            ("AI 大模型", "DeepSeek API、Ollama/qwen3、本地 deterministic fast-path", "LLMRouter 根据任务复杂度和配置选择在线/本地/兜底模型。"),
            ("RAG", "BGE large zh embedding、BGE reranker、关键词召回、向量召回、Rerank", "知识库 chunks 进入 kb_documents/kb_chunks，embedding_json 保存 1024 维向量。"),
            ("部署", "Docker Compose、Dockerfile.backend、Dockerfile.frontend、Nginx/静态前端", "compose 服务包含 postgres、redis、backend、celery_worker、frontend，可选 ollama profile。"),
        ],
    )

    doc.add_heading("4. 项目整体架构", level=1)
    doc.add_paragraph("整体架构可以理解为三条主线：数据与模型主线、AI/RAG 主线、产品化运行主线。")
    add_table(
        doc,
        ["架构层", "职责", "关键组件"],
        [
            ("前端架构", "提供业务工作台、AI 助手对话、任务中心、系统设置和用户管理。", "frontend/src/pages、frontend/src/api.ts、AuthContext、UserManagementPage、AssistantPage、TaskCenterPage。"),
            ("接口层", "对外暴露 REST API，负责鉴权、审计、请求校验和任务派发。", "backend/app/api/v1/endpoints/*.py、api_router。"),
            ("服务层", "封装业务逻辑、RAG 检索、AI 答案生成、数据状态、报告、Stage1 服务。", "backend/app/services、backend/app/ai_assistant。"),
            ("Repository 层", "统一访问 PostgreSQL，避免业务层直接拼 SQL 或读文件。", "backend/app/repositories。"),
            ("数据层", "保存预测、模型、原始市场/天气/负荷、任务、审计、知识库、用户等事实数据。", "PostgreSQL + Alembic 迁移。"),
            ("模型层", "训练、加载和执行电价预测模型，输出预测、风险和模型指标。", "prediction_engine、model_ops、model_artifacts。"),
            ("任务调度", "将知识库索引、embedding 刷新、报告生成、数据同步等耗时任务异步化。", "Celery/Redis、backend/app/workers。"),
            ("AI/RAG", "将用户问题转为意图、工具调用、RAG 证据和 LLM 回答。", "answer_chat_accurate、LLMRouter、rag_service、embedding_service、rerank_service。"),
            ("部署架构", "通过 Docker Compose 编排服务，支持本地/服务器部署。", "docker-compose.yml、docker/backend-entrypoint.sh、README_DEPLOY.md。"),
        ],
    )
    doc.add_paragraph("典型链路：用户在前端提问或触发任务 -> FastAPI 路由鉴权 -> 服务层调用 Repository/工具/RAG/模型 -> 需要长耗时则进入任务中心 -> PostgreSQL 持久化结果 -> 前端轮询或展示结果。")

    doc.add_heading("5. 项目目录结构说明", level=1)
    add_table(
        doc,
        ["路径", "类型", "作用", "是否核心", "说明"],
        [
            ("backend/", "后端目录", "FastAPI 后端、API、服务、Repository、AI 助手、任务、数据库连接。", "是", "当前 Web 平台核心后端。"),
            ("backend/app/api/v1/endpoints/", "接口目录", "assistant、auth、users、forecast、data、knowledge、task、report、model、system 等接口。", "是", "对外 REST API 入口。"),
            ("backend/app/ai_assistant/", "AI 助手模块", "意图识别、Prompt、工具、上下文、回答生成、Answer Guard、会话状态。", "是", "AI 问答核心链路。"),
            ("backend/app/services/", "服务层", "RAG、embedding、rerank、platform、stage1、数据访问、报表等服务。", "是", "业务逻辑聚合层。"),
            ("backend/app/repositories/", "数据访问层", "封装 PostgreSQL 表读写，如 forecast、knowledge、task、audit、user 等。", "是", "主事实源访问入口。"),
            ("backend/app/workers/", "任务 worker", "Celery app、任务定义、任务派发与本地线程 fallback。", "是", "任务中心核心。"),
            ("backend/alembic/", "数据库迁移", "PostgreSQL 表结构迁移。", "是", "包含 0001-0006 等迁移。"),
            ("frontend/", "前端目录", "React/Vite 工作台。", "是", "页面、API 客户端、AuthContext、业务组件。"),
            ("prediction_engine/", "预测引擎", "特征工程、模型训练、快速预测、正式预测、遗留模型封装。", "是", "电价预测核心。"),
            ("model_ops/", "模型运维", "模型注册、监控、漂移检测、自动优化等。", "是", "模型产品化支撑。"),
            ("knowledge_base/", "知识库", "项目内 Markdown 知识片段。", "是", "RAG 源知识之一。"),
            ("knowledge_pipeline/", "知识处理管道", "多格式资料转 Markdown/chunks、导入 kb 表、RAG smoke test。", "是", "批量知识库治理工具。"),
            ("scripts/", "运维/导入脚本", "PostgreSQL 检查、Redis/Celery 检查、迁移、导入、管理员创建等。", "是", "部署和验收辅助。"),
            ("tests/", "测试目录", "单元测试、评测集、产品化验收报告。", "是", "回归和验收依据。"),
            ("docker/", "容器配置", "后端/前端 Dockerfile、entrypoint。", "是", "Compose 部署支撑。"),
            ("output/", "运行输出", "预测结果、报告、图表、日志等。", "部分", "根据运行情况产生，需区分源代码与产物。"),
            ("bge-large-zh-v1.5/", "本地模型", "BGE embedding 模型目录。", "部分", "模型权重未读取；由 env 指定路径使用。"),
            ("bge-reranker-v2-m3/", "本地模型", "BGE reranker 模型目录。", "部分", "模型权重未读取。"),
        ],
    )

    doc.add_heading("6. 核心功能模块分析", level=1)
    add_table(
        doc,
        ["模块", "作用", "主要文件", "输入", "输出", "与其他模块关系"],
        [
            ("登录/RBAC/用户管理", "用户登录、JWT、角色权限、用户增删改禁用、重置密码、审计。", "auth.py、users.py、core/security.py、auth/jwt.py、auth/password.py、user_repository.py、LoginPage、UserManagementPage", "账号密码、JWT、角色权限", "token、当前用户、权限校验结果", "AUTH_REQUIRED=1 时必须 Bearer token；AUTH_REQUIRED=0 保留开发 fallback。"),
            ("数据采集与导入", "从 PJM/Open-Meteo 等外部源抓取市场、负荷、天气数据，清洗合并并可同步 PostgreSQL。", "fetch_power_market_data.py、scripts/import_stage1_raw_data_to_postgres.py", "外部 API、历史数据文件", "raw_market/raw_weather/raw_load 等表、Excel 兼容产物", "数据库优先，旧 Excel/CSV 受 DATABASE_ALLOW_LEGACY_FALLBACK 控制。"),
            ("预测与模型", "训练和加载电价预测模型，生成未来价格、风险等级、预测区间和特征快照。", "prediction_engine/model_trainer.py、fast_forecast.py、formal_forecast.py、features.py", "master_table、特征、模型配置", "forecast_results、模型指标、预测文件", "封装遗留 v4_fix1 模型，逐步产品化。"),
            ("AI 助手", "识别用户意图，调用工具/RAG/LLM，生成自然回答并隐藏调试信息。", "assistant.py、service.py、intent_router.py、prompts.py、llm_router.py、answer_guard.py", "用户问题、上下文、debug、answer_style", "answer、provider、evidence_summary、debug trace", "daily/capability fast-path 不调用 LLM；专业问题触发 RAG。"),
            ("知识库/RAG", "关键词+向量+rerank 混合检索，Top-K 证据注入 Prompt。", "rag_service.py、embedding_service.py、rerank_service.py、knowledge_repository.py、knowledge_pipeline", "用户 query、kb_chunks、embedding_json", "RAG items、证据摘要、引用来源", "BGE 1024 维 embedding；fallback 仅作为降级，正常不混用。"),
            ("报告中心", "生成、查看、下载、审核、发布报告。", "report.py、platform_services.py、report_repository.py、ReportPage", "预测/风险/AI 摘要", "报告记录、审核记录、文件路径", "报告生成已支持异步任务。"),
            ("任务中心", "统一创建、查询、重试、取消异步任务，展示进度和错误。", "task.py、task_repository.py、workers/dispatcher.py、workers/tasks.py、TaskCenterPage", "task_type、metadata", "task_id、状态、进度、result_ref", "Celery/Redis 可用时走 worker，否则开发态 local_thread。"),
            ("审计与安全", "记录敏感操作、脱敏日志和审计 metadata。", "audit_repository.py、security/audit endpoints、secret masking 工具", "用户操作、资源、结果", "audit_logs", "禁止记录密码、token、API Key、DATABASE_URL。"),
        ],
    )

    doc.add_heading("7. 核心代码文件说明", level=1)
    add_table(
        doc,
        ["文件", "关键函数/类", "作用", "输入", "输出/调用关系"],
        [
            ("backend/app/main.py", "create_app", "创建 FastAPI 应用，配置 CORS、日志、model_gateway 和 api_router。", "Settings、路由模块", "FastAPI app。"),
            ("backend/app/core/config.py", "Settings、get_settings", "集中读取环境变量和默认配置，含 DB、Redis、LLM、RAG、JWT、任务模式。", "环境变量", "全局 settings。"),
            ("backend/app/core/security.py", "CurrentUser、require_permission、get_current_user", "RBAC 权限校验、开发 fallback、JWT 用户解析。", "Authorization header、AUTH_REQUIRED", "当前用户或 401/403。"),
            ("backend/app/api/v1/endpoints/assistant.py", "chat、analyze、traces", "AI 助手入口，控制 debug 权限和审计。", "ChatRequest", "ChatResponse、trace。"),
            ("backend/app/ai_assistant/service.py", "answer_chat_accurate", "AI 助手主链路：意图、工具、RAG、LLM、Guard、trace。", "用户问题、模型选择、answer_style、debug", "回答 payload。"),
            ("backend/app/ai_assistant/llm_router.py", "LLMRouter、DeepSeekClient、OllamaClient", "多模型路由与 fallback。", "messages、model_provider、task_type", "LLMResult。"),
            ("backend/app/services/rag_service.py", "search_knowledge、index_local_knowledge", "RAG 混合召回、评分、缓存、policy 文档降噪。", "query、top_k、debug", "RAGSearchResult。"),
            ("backend/app/services/embedding_service.py", "SentenceTransformersEmbeddingProvider", "加载 BGE embedding、本地 hash fallback、批量向量化。", "文本列表", "embedding + metadata。"),
            ("backend/app/services/rerank_service.py", "BGETransformersReranker、rerank_candidates", "BGE reranker 精排，失败回退 heuristic。", "query + candidates", "rerank score。"),
            ("backend/app/repositories/knowledge_repository.py", "upsert_document、search_keyword_chunks、list_embedded_chunks", "kb_documents/kb_chunks 的入库、关键词检索、向量候选读取。", "文档/chunk/query", "数据库记录。"),
            ("prediction_engine/model_trainer.py", "train_model_package", "训练并保存模型包、指标、manifest。", "训练数据、配置", "model_artifacts。"),
            ("prediction_engine/fast_forecast.py", "run_fast_forecast", "加载 active model，执行特征工程和未来预测。", "master_table、模型包", "预测结果、业务摘要。"),
            ("fetch_power_market_data.py", "main、fetch_*、sync_stage1_to_postgres", "抓取 PJM/天气/负荷数据，清洗合并并可入库。", "外部 API、配置", "raw 表/Excel/master_table。"),
            ("backend/app/workers/tasks.py", "knowledge_import_task、embedding_refresh_task、report_generate_task", "Celery 任务执行与状态更新。", "task_id、payload", "task_runs 状态、result_ref。"),
            ("frontend/src/api.ts", "apiRequest、业务 API 方法", "前端统一 API 客户端，自动注入 token 和处理 401。", "HTTP 请求", "JSON 响应。"),
            ("frontend/src/context/AuthContext.tsx", "AuthProvider、useAuth", "前端认证状态、登录退出、权限判断。", "token、/api/auth/me", "user、permissions、isAuthenticated。"),
            ("frontend/src/pages/assistant/AssistantPage.tsx", "AssistantPage", "AI 助手页面、风格选择、debug 模式、证据折叠、复制。", "用户问题、API 响应", "对话 UI。"),
        ],
    )

    doc.add_heading("8. 项目启动与运行流程", level=1)
    doc.add_heading("8.1 后端启动流程", level=2)
    add_numbered(
        doc,
        [
            "读取环境变量并构造 Settings。",
            "FastAPI create_app 配置 CORS、日志和异常处理。",
            "注册 model_gateway router 与 api_router。",
            "运行时路由根据 AUTH_REQUIRED 判断是否需要 JWT。",
            "数据库连接通过 backend/app/db/session.py（根据代码推断）和 repository 层访问。",
        ],
    )
    doc.add_heading("8.2 Docker Compose 启动流程", level=2)
    add_numbered(
        doc,
        [
            "postgres、redis 先启动并通过 healthcheck。",
            "backend-entrypoint 等待 PostgreSQL 与 Redis。",
            "根据配置执行 Alembic migration。",
            "backend 启动 FastAPI。",
            "celery_worker 连接同一 Redis/PostgreSQL 执行异步任务。",
            "frontend 提供静态页面并代理 /api 到后端。",
        ],
    )
    doc.add_heading("8.3 本地开发常用命令", level=2)
    add_table(
        doc,
        ["场景", "命令示例", "说明"],
        [
            ("后端开发", "python -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000", "启动 FastAPI。"),
            ("前端开发", "cd frontend && npm install && npm run dev", "启动 Vite 开发服务器。"),
            ("数据库迁移", "alembic upgrade head", "执行 PostgreSQL 迁移。"),
            ("Docker 部署", "docker compose --env-file .env.docker up -d --build", "启动 compose 服务。"),
            ("预测流程", "python main_daily_run.py --refresh-data --fast-forecast", "根据代码推断，用于刷新数据和执行快速预测。"),
            ("测试", "python -m pytest -q；npm run build", "后端测试和前端构建。"),
        ],
    )

    doc.add_heading("9. 数据来源与数据处理流程", level=1)
    add_table(
        doc,
        ["数据类型", "来源", "读取方式", "落地位置", "说明"],
        [
            ("日前/实时 LMP", "PJM Data Miner API", "requests 调用 dataminer API；自动发现订阅 key 的公共配置入口。", "raw_market、master_table、forecast 输入", "代码中包含 da_hrl_lmps、rt_hrl_lmps 等逻辑。"),
            ("实际/预测负荷", "PJM Data Miner API", "hrl_load_metered、load_frcstd_7_day/load_frcstd_hist 等接口。", "raw_load、master_table", "支持实际负荷和预测负荷对齐。"),
            ("天气", "Open-Meteo Archive API；数据目录中也有 NOAA/NWS 语义标签", "按 Richmond/DOM 场景拉取温度、湿度、风速等变量。", "raw_weather、master_table", "天气归因需关注更新时间和预测窗口。"),
            ("新能源", "根据代码和表结构预留", "raw_renewable 表已建；当前导入报告显示 0 行。", "raw_renewable", "当前项目中未发现完整新能源外部抓取闭环。"),
            ("特征重要性", "模型训练/模型分析输出", "导入 feature_importance 表。", "feature_importance", "用于模型解释和 AI 工具。"),
            ("知识库", "knowledge_base Markdown、knowledge_pipeline 转换后的 chunks.jsonl", "本地索引/JSONL 导入 PostgreSQL。", "kb_documents、kb_chunks", "embedding_json 保存 BGE 向量。"),
            ("预测/报告/任务/审计", "系统运行产生", "服务层写 repository。", "forecast_results、reports、task_runs、audit_logs 等", "PostgreSQL 主事实源。"),
        ],
    )
    doc.add_paragraph("数据处理链路：外部市场/负荷/天气数据 -> 清洗标准化字段 -> 按时间对齐合并 master_table -> 特征工程 -> 模型训练/预测 -> 预测结果入库/输出文件 -> 前端展示和 AI 助手引用。")

    doc.add_heading("10. 数据分析、特征工程与建模流程", level=1)
    add_table(
        doc,
        ["阶段", "主要逻辑", "关键文件", "风险控制"],
        [
            ("数据清洗", "统一时间字段、价格字段、节点/区域字段、天气和负荷变量；处理缺失和异常格式。", "fetch_power_market_data.py、database_utils.py", "外部 API 字段变动会影响清洗；需持续维护字段映射。"),
            ("数据合并", "将 DA/RT LMP、实际负荷、预测负荷、天气和节假日/时间特征按小时合并。", "fetch_power_market_data.py", "时间区间和时区对齐是关键风险。"),
            ("特征工程", "时间特征、lag 特征、rolling/EWM、历史统计、业务特征、尖峰特征。", "prediction_engine/features.py、legacy v4_fix1", "通过 leakage_check_report 避免未来信息泄露。"),
            ("训练", "基于目标 da_price 训练多个回归模型、尖峰/高峰模型和分类器，保存模型包。", "prediction_engine/model_trainer.py、pipeline.py", "过拟合、样本偏移和训练/预测 schema 不一致是主要风险。"),
            ("预测", "加载 active model artifacts，对未来小时生成预测、风险等级、P10/P90 区间和业务摘要。", "prediction_engine/fast_forecast.py、formal_forecast.py", "若未来负荷/天气缺失，只能做有限预测或降级。"),
            ("模型运维", "模型注册、监控、漂移检测、自动优化和指标对比。", "model_ops/", "需要长期运行数据验证。"),
        ],
    )

    doc.add_heading("11. 模型与算法逻辑说明", level=1)
    add_table(
        doc,
        ["模型/算法", "所在文件", "输入", "输出", "作用", "评估方式", "备注"],
        [
            ("回归模型集合", "高峰尖刺增强版...v4_fix1.py、model_trainer.py", "特征矩阵、目标 da_price", "电价预测", "基础电价预测", "RMSE、MAE、R2、残差分析（根据代码推断）", "包含 RandomForest、ExtraTrees、HistGradientBoosting、Ridge、Huber、LinearRegression，且可选 XGBoost/LightGBM/CatBoost。"),
            ("高峰专用模型", "legacy v4_fix1、model_trainer.py", "高峰时段样本和增强特征", "高峰时段价格预测", "改善晚高峰和尖峰价格表现", "高峰样本误差和尖峰命中（根据代码推断）", "用于解决平均模型压低尖峰的问题。"),
            ("尖峰分类器", "legacy v4_fix1、model_trainer.py", "特征、尖峰标签", "尖峰概率/风险", "识别价格冲高风险", "分类指标和概率校准（根据代码推断）", "支持 RandomForest/HGB/Logistic 及可选 XGB/LGBM/CatBoost 分类器。"),
            ("预测区间", "formal_forecast.py", "预测值、测试残差", "P10/P90 或上下界", "表达不确定性", "残差分位数", "若历史残差不足，区间可靠性下降。"),
            ("偏差校正", "fast_forecast.py", "预测结果和历史/测试偏差", "校正后预测", "减少系统偏差", "校正前后误差对比（根据代码推断）", "需要防止过度校正。"),
            ("RAG 向量相似度", "embedding_service.py、rag_service.py", "query embedding、chunk embedding", "vector_score", "语义召回知识片段", "Top-K 命中率", "BGE large zh 1024 维。"),
            ("BGE reranker", "rerank_service.py", "query 与候选 chunk 对", "rerank_score", "精排证据", "评测集 RAG 命中率", "失败时回退 heuristic。"),
        ],
    )
    doc.add_paragraph("模型风险：当前项目仍保留遗留大脚本封装，虽然已通过 prediction_engine 分层封装，但核心算法模块化程度仍可继续提升。生产环境需要持续监控模型漂移、外部 API 数据质量和未来特征可用性。")

    doc.add_heading("12. AI 大模型接入与调用逻辑", level=1)
    add_kv_table(
        doc,
        [
            ("使用的大模型", "DeepSeek API、Ollama 本地模型（默认 qwen3:4b 语义）、deterministic fast-path。"),
            ("调用方式", "DeepSeek 通过 OpenAI-compatible chat completions API；Ollama 通过本地 HTTP API；fast-path 直接规则返回，不调用 LLM。"),
            ("配置方式", "通过环境变量配置 LLM_PROVIDER、LLM_ROUTER_MODE、DEEPSEEK_BASE_URL、DEEPSEEK_API_KEY、DEEPSEEK_MODEL、OLLAMA_BASE_URL、OLLAMA_MODEL 等；文档不包含真实 key。"),
            ("调用入口", "前端 AssistantPage -> POST /api/ai/chat -> endpoints/assistant.py -> answer_chat_accurate。"),
            ("Prompt 位置", "backend/app/ai_assistant/prompts.py。"),
            ("上下文构造", "context_pack_builder、工具结果、RAG evidence、会话状态、answer_style 共同构造 Prompt。"),
            ("知识库/RAG", "专业问题触发 RAG：query rewrite -> 关键词召回 -> BGE 向量召回 -> 合并候选 -> BGE reranker -> Top-K evidence 注入 Prompt。"),
            ("流式输出", "当前项目中未发现正式流式输出实现。"),
            ("function/tool calling", "当前不是 OpenAI function calling 形式，而是后端内置工具机制：如预测、数据新鲜度、知识库、报告等工具。"),
            ("多轮记忆", "conversation_store/state_resolver 支持会话状态和追问解析。"),
            ("异常与 fallback", "LLMRouter 支持 DeepSeek/Ollama 互相 fallback，错误信息脱敏；RAG/embedding/rerank 失败不应导致主问答崩溃。"),
            ("token/成本控制", "RAG Top-K、candidate limit、prompt style、fast-path 和缓存用于控制成本；仍建议后续增加更明确的 token 预算统计。"),
        ],
    )
    doc.add_paragraph("完整 AI 调用链路：用户输入 -> 意图识别/fast-path 判断 -> 工具与 RAG 触发判断 -> 构造事实包与 Prompt -> LLMRouter 选择 DeepSeek/Ollama/兜底 -> Answer Guard 安全校验 -> 审计/trace 入库 -> 前端以普通模式展示自然回答，debug 模式展示 RAG/trace/tools。")

    doc.add_heading("13. 接口与 API 调用说明", level=1)
    add_table(
        doc,
        ["接口路径/方法", "请求方式", "所在文件", "入参", "出参", "作用", "调用方"],
        [
            ("/health", "GET", "system.py", "无", "健康状态", "后端存活检查", "Docker/前端/运维"),
            ("/api/db/health", "GET", "system.py", "无", "DB 连接摘要", "PostgreSQL 健康检查", "运维/部署脚本"),
            ("/api/tasks/health", "GET", "task.py", "无", "任务中心状态", "Redis/Celery/执行模式检查", "运维/前端"),
            ("/api/auth/login", "POST", "auth.py", "username/password", "access_token/user", "用户登录", "LoginPage"),
            ("/api/auth/me", "GET", "auth.py", "Bearer token", "当前用户", "AuthContext"),
            ("/api/auth/logout", "POST", "auth.py", "token", "成功状态", "前端退出"),
            ("/api/users", "GET/POST", "users.py", "查询条件/用户信息", "用户列表/新用户", "用户管理页面"),
            ("/api/users/{id}", "PATCH/DELETE", "users.py", "用户更新/软禁用", "用户信息", "用户管理页面"),
            ("/api/users/{id}/reset-password", "POST", "users.py", "new_password", "成功状态", "管理员重置密码", "用户管理页面"),
            ("/api/ai/chat", "POST", "assistant.py", "question、model_provider、answer_style、debug", "answer、provider、evidence、debug", "AI 助手"),
            ("/api/ai/traces", "GET", "assistant.py", "trace 查询", "trace 列表", "开发者调试", "AI 助手调试区"),
            ("/api/knowledge/search", "POST/GET（根据代码推断）", "knowledge.py", "query/top_k", "RAG 结果", "知识库检索", "知识库页面/AI 工具"),
            ("/api/knowledge/index-local", "POST", "knowledge.py", "索引参数", "task_id", "异步索引本地知识库", "知识库页面"),
            ("/api/knowledge/embedding-refresh", "POST", "knowledge.py", "刷新参数", "task_id", "异步刷新 embedding", "知识库页面"),
            ("/api/data/status", "GET", "data.py", "无", "数据状态", "数据中心状态卡", "DataCenterPage"),
            ("/api/data/tables", "GET", "data.py", "无", "表清单", "数据库表浏览", "DataCenterPage/AI 工具"),
            ("/api/data/tables/{table}/rows", "GET", "data.py", "分页/筛选", "表行数据", "数据中心表格", "DataCenterPage"),
            ("/api/data/sync-core", "POST", "data.py", "同步参数", "task_id", "同步核心数据到 PostgreSQL", "数据中心"),
            ("/api/forecast/run", "POST", "forecast.py", "预测参数", "task_id", "触发预测任务", "预测中心"),
            ("/api/forecast/latest", "GET", "forecast.py", "无", "最新预测", "预测中心/AI 工具"),
            ("/api/reports/generate", "POST", "report.py", "报告参数", "task_id/报告记录", "生成报告", "报告中心"),
            ("/api/tasks", "GET/POST", "task.py", "任务创建/查询", "任务列表/任务 ID", "任务中心", "TaskCenterPage"),
            ("/api/tasks/{task_id}/retry", "POST", "task.py", "task_id", "任务状态", "重试失败/取消任务", "TaskCenterPage"),
            ("/api/tasks/{task_id}/cancel", "POST", "task.py", "task_id", "取消状态", "取消任务", "TaskCenterPage"),
            ("/api/model/*", "GET/POST", "model.py", "模型查询/操作参数", "模型状态/指标", "模型中心", "ModelCenterPage"),
        ],
    )

    doc.add_heading("14. 数据库与存储设计", level=1)
    doc.add_paragraph("数据库以 PostgreSQL 为主事实源，Alembic 迁移覆盖 0001 到 0006。当前仍保留部分旧文件 fallback，但生产环境通过 DATABASE_ALLOW_LEGACY_FALLBACK=0 禁用旧路径。")
    add_table(
        doc,
        ["表/对象", "来源迁移", "用途", "关键字段/说明"],
        [
            ("forecast_runs / forecast_results / prediction_results（根据迁移真实命名）", "0001", "保存预测运行和预测结果。", "run_id、datetime、predicted_price、risk_level 等。"),
            ("model_versions / model_metrics / model_registry（根据代码和迁移）", "0001", "模型版本、指标和注册信息。", "model_id、metrics_json、is_active。"),
            ("task_runs / task_logs", "0001、0004、0005", "任务中心主表和日志。", "task_id、task_type、status、progress、execution_mode、cancel_requested、celery_task_id。"),
            ("ai_traces", "0001、0002", "AI 调用 trace、工具、模型和上下文记录。", "trace_id、question、answer、provider、metadata。"),
            ("report_runs / reports / report_reviews", "0001", "报告生成和审核。", "report_id、status、reviewer、comments。"),
            ("kb_documents / kb_chunks", "0002", "知识库文档和片段。", "doc_id、chunk_id、content、embedding_json、metadata_json。"),
            ("roles / users / user_roles", "0002、0006", "RBAC、用户和角色关联。", "username、password_hash、role、is_active、last_login_at。"),
            ("audit_logs", "0002", "敏感操作审计。", "user_id、action、resource_type、resource_id、result、metadata_json。"),
            ("raw_market", "0003", "市场原始/标准化数据。", "时间、区域、DA/RT 价格、节点等字段（根据代码推断）。"),
            ("raw_weather", "0003", "天气原始/标准化数据。", "时间、温度、湿度、风速等字段。"),
            ("raw_load", "0003", "负荷数据。", "实际负荷、预测负荷、区域、时间。"),
            ("raw_renewable", "0003", "新能源数据预留。", "当前导入报告显示 0 行，业务链路待补。"),
            ("feature_importance", "0003", "模型特征重要性。", "feature、importance、model/run 信息。"),
        ],
    )
    doc.add_paragraph("文件存储：模型工件保存在 model_artifacts/，预测和报告运行产物在 output/，知识库转换结果在 knowledge_pipeline/，本地模型目录为 bge-large-zh-v1.5 与 bge-reranker-v2-m3。生产部署不应把大模型权重打入镜像，应通过卷挂载或外部模型服务配置。")

    doc.add_heading("15. 配置文件与环境变量说明", level=1)
    add_table(
        doc,
        ["配置类别", "关键变量/文件", "说明"],
        [
            ("应用环境", "APP_ENV、LOG_LEVEL", "区分 development/production；生产环境对 fallback、JWT secret 等有安全提示。"),
            ("数据库", "DATABASE_URL、POSTGRES_*、DATABASE_ALLOW_LEGACY_FALLBACK", "PostgreSQL 连接与旧数据源 fallback 控制。文档不记录真实密码。"),
            ("Redis/Celery", "REDIS_URL、CELERY_BROKER_URL、CELERY_RESULT_BACKEND、TASK_EXECUTION_MODE", "任务中心执行模式：auto/celery/local_thread。"),
            ("认证", "AUTH_REQUIRED、JWT_SECRET_KEY、JWT_ALGORITHM、JWT_ACCESS_TOKEN_EXPIRE_MINUTES", "生产环境应 AUTH_REQUIRED=true 并使用非默认 JWT_SECRET。"),
            ("LLM", "LLM_PROVIDER、LLM_ROUTER_MODE、DEEPSEEK_BASE_URL、DEEPSEEK_API_KEY、DEEPSEEK_MODEL、OLLAMA_BASE_URL、OLLAMA_MODEL", "统一模型网关配置。所有 Key 仅通过环境变量。"),
            ("RAG", "RAG_PROFILE、RAG_KEYWORD_LIMIT、RAG_VECTOR_LIMIT、RAG_RERANK_CANDIDATE_LIMIT、RAG_TOP_K、RAG_CACHE_TTL_SECONDS", "支持 quality/balanced/fast 检索参数分档和缓存。"),
            ("Embedding/Rerank", "RAG_EMBEDDING_PROVIDER、RAG_EMBEDDING_MODEL_PATH、RAG_RERANK_PROVIDER、RAG_RERANK_MODEL_PATH", "BGE 本地模型路径从环境变量读取，不硬编码。"),
            ("前端", "VITE_API_BASE_URL、VITE_AUTH_REQUIRED", "控制 API 地址和前端登录门禁。"),
            ("Docker", ".env.docker.example、docker-compose.yml", "容器部署示例，真实 .env 不应提交。"),
        ],
    )

    doc.add_heading("16. 项目部署与运行方式", level=1)
    doc.add_heading("16.1 本地开发", level=2)
    add_numbered(
        doc,
        [
            "准备 Python/Node/PostgreSQL/Redis 环境。",
            "复制 .env.example 为本地 .env，并填写数据库、Redis、LLM、RAG 模型路径等变量；本次分析未读取真实 .env。",
            "执行 alembic upgrade head 初始化数据库。",
            "启动后端、前端和必要的 Celery worker。",
            "如需正式账号，使用 scripts/create_admin_user.py 从环境变量创建管理员。",
        ],
    )
    doc.add_heading("16.2 Docker Compose", level=2)
    add_table(
        doc,
        ["服务", "职责", "备注"],
        [
            ("postgres", "数据库主事实源。", "compose healthcheck 验证。"),
            ("redis", "Celery broker/result backend。", "任务中心依赖。"),
            ("backend", "FastAPI API 服务。", "entrypoint 等待 DB/Redis 并可执行迁移。"),
            ("celery_worker", "异步任务执行。", "知识库、embedding、报告、数据同步等任务。"),
            ("frontend", "Web 工作台。", "静态前端和 /api 代理。"),
            ("ollama", "可选本地模型服务。", "optional profile，不强制启动。"),
        ],
    )
    doc.add_heading("16.3 已有运行态报告", level=2)
    doc.add_paragraph("tests/productization/output/docker_runtime_report.md 显示 Docker Compose 运行态曾完成 backend/frontend/postgres/redis/celery_worker 全链路验证，健康检查和任务中心 smoke test 通过。需要注意该报告生成时间早于部分后续 auth/user 变更，生产上线前应重新跑一次完整 compose 验收。")

    doc.add_heading("17. 当前项目完成度评估", level=1)
    add_table(
        doc,
        ["领域", "完成度判断", "证据", "备注"],
        [
            ("Web 工作台", "较完整", "前端有首页、数据、预测、AI、报告、模型、知识库、任务、设置/用户管理页面。", "需继续做端到端 UI 回归。"),
            ("PostgreSQL 主事实源", "基本完成", "迁移覆盖核心表；Stage1 raw 表已建设；repository 层覆盖主要模块。", "仍有 legacy fallback，需要生产禁用和监控。"),
            ("AI 助手/RAG", "较成熟", "v2.7/2.8 评测曾达到 93-94/100，RAG 命中率 100%，debug 默认隐藏。", "需继续优化数据库查询工具和线上耗时。"),
            ("知识库治理", "较完整", "批处理、chunk 校验、导入、BGE embedding、smoke test 已形成链路。", "仍有失败文件需人工处理/OCR。"),
            ("预测模型", "可用但需重构", "已有遗留模型+prediction_engine 封装+模型工件。", "核心算法仍依赖遗留大脚本，需持续模块化。"),
            ("任务中心", "开发/演示可用", "Celery/Redis/local_thread、任务状态、重试/取消、前端页面已具备。", "生产需长稳验证和 worker 运维。"),
            ("安全/RBAC", "产品化基础完成", "JWT、用户管理、角色权限、审计、脱敏已接入。", "需企业级密钥管理、SSO、密码策略。"),
            ("部署", "Compose 可用", "docker-compose、Dockerfile、entrypoint、部署脚本、README_DEPLOY。", "需生产 HTTPS、备份、监控、模型卷挂载。"),
        ],
    )

    doc.add_heading("18. 存在问题与风险点", level=1)
    add_table(
        doc,
        ["风险类型", "问题", "影响", "建议"],
        [
            ("编码与文档", "部分旧中文文件/README 在当前终端读取时出现编码显示异常。", "影响命令行审查和自动文档抽取。", "统一源文件 UTF-8 编码，避免历史乱码扩散。"),
            ("旧数据源 fallback", "Excel/CSV/本地 JSON 兼容逻辑仍存在。", "生产中可能绕过 PostgreSQL 主事实源。", "生产环境 DATABASE_ALLOW_LEGACY_FALLBACK=0，并完善告警。"),
            ("模型模块化", "核心预测算法仍封装遗留大脚本。", "后续维护、测试和模型治理成本较高。", "逐步拆分为 feature、trainer、evaluator、predictor 独立模块。"),
            ("AI 数据库查询", "根据近期人工测试，AI 助手对 raw_market/raw_load/raw_weather 等表的新鲜度查询仍可能漏查或回答不完整。", "影响“AI 直接查询数据库数据”的可信度。", "增强 AI 工具层表目录、字段映射、SQL 只读查询和错误解释。"),
            ("Docker 模型挂载", "compose 验收中 BGE 模型目录可能未真实挂载，embedding refresh 可能降级。", "容器内 RAG 质量与本地不同。", "用 volume 显式挂载 BGE 模型目录，并在 health 中检查 provider/dim。"),
            ("任务中心生产稳定性", "Celery 长稳、取消 revoke、重试幂等仍需验证。", "长任务失败恢复和资源控制风险。", "补运行监控、死信/超时、任务幂等和日志留存策略。"),
            ("安全合规", "当前已具备 JWT/RBAC/审计，但企业级密钥管理、SSO、密码策略仍不足。", "生产合规和权限治理风险。", "接入 KMS/Secret Manager、HTTPS、SSO、密码复杂度和审计报表。"),
            ("外部 API 依赖", "PJM/天气数据源接口、订阅 key 和字段可能变化。", "数据刷新失败导致预测和 AI 解释失真。", "增加数据源 SLA、重试、断点续拉和字段变更告警。"),
            ("评测覆盖", "已有 AI 评测集和产品化脚本，但缺少全链路业务验收清单。", "上线前风险难量化。", "补 E2E 测试、性能基线、用户验收用例。"),
        ],
    )

    doc.add_heading("19. 优化建议与后续开发方案", level=1)
    add_table(
        doc,
        ["阶段", "建议事项", "优先级", "说明"],
        [
            ("短期", "修复 AI 数据库查询工具的表发现、字段映射和只读 SQL 查询能力。", "高", "让 AI 在需要时可查询 raw_market/raw_load/raw_weather/forecast_results 等表，并清楚说明查不到的原因。"),
            ("短期", "统一前端 AUTH_REQUIRED 与后端 AUTH_REQUIRED 的部署配置。", "高", "避免生产打开后不弹登录或 dev fallback 暴露。"),
            ("短期", "重新运行最新 Docker Compose + auth/user 版本验收。", "高", "docker_runtime_report 需要覆盖 0006 auth/users。"),
            ("短期", "完成失败知识库文件人工转换/OCR 清单处理。", "中", "特别是扫描 PDF、DOC/WPS/OFD。"),
            ("短期", "增加 RAG/LLM token 和耗时观测报表。", "中", "持续优化响应耗时和成本。"),
            ("中期", "将遗留预测大脚本拆解为标准训练、评估、预测、特征注册模块。", "高", "降低模型迭代和回归测试成本。"),
            ("中期", "引入模型治理：模型版本审批、灰度、回滚、漂移报警。", "高", "支撑生产电价预测可靠性。"),
            ("中期", "任务中心生产化：超时、并发限制、队列分层、幂等、失败补偿。", "中", "保障导入、报告和预测长任务稳定。"),
            ("中期", "pgvector 或专用向量库迁移评估。", "中", "当前可用 embedding_json + 内存向量检索；数据量增大后需要索引化。"),
            ("长期", "企业级部署：HTTPS、备份恢复、Prometheus/Grafana、集中日志、KMS。", "高", "满足企业上线和运维要求。"),
            ("长期", "多租户/多项目隔离。", "中", "支持多个市场区域、客户、项目同时运营。"),
            ("长期", "AI Agent 编排和自动日报。", "中", "在基础稳定后再引入多 Agent，避免过早复杂化。"),
        ],
    )

    doc.add_heading("20. 总结", level=1)
    doc.add_paragraph("智能运营分析项目已经从单点预测脚本演进为包含数据、模型、AI、知识库、任务中心、安全审计和 Docker 部署的准产品化平台。其核心价值在于把复杂的电力市场数据和预测模型结果转化为业务人员可理解、可追溯、可行动的运营建议。")
    doc.add_paragraph("当前项目优势包括：数据链路和模型链路已有基础，AI 助手已经接入 DeepSeek/Ollama/RAG，PostgreSQL 主事实源和任务中心已形成骨架，前端工作台覆盖主要业务场景，部署闭环基本建立。")
    doc.add_paragraph("后续重点不应继续盲目扩功能，而应优先夯实生产可靠性：数据库查询工具、模型模块化、Docker 最新运行态验收、Celery 长稳、密钥治理、外部数据 SLA 和全链路 E2E 测试。完成这些后，项目具备进一步进入企业级试运行和交付的基础。")

    doc.add_page_break()
    doc.add_heading("附录 A：本次分析的未确认事项", level=1)
    add_bullets(
        doc,
        [
            "未读取真实 .env、.env.docker、API Key、数据库密码、JWT_SECRET 或 Authorization。",
            "未读取模型权重、大型图片、视频、二进制文件和 node_modules 等依赖产物。",
            "部分旧中文文件在当前终端输出中存在编码显示异常，文档仅引用可确认的代码路径和功能，不引用乱码内容。",
            "未对外部 PJM/天气 API 执行实时联网抓取；数据源说明基于代码和配置。",
            "未重新执行完整 100 题 AI 评测和最新 Docker Compose 验收；文档引用的是项目中已有测试/验收报告。",
            "数据库表字段的部分细节根据 Alembic 迁移和 repository 代码推断，最终以实际数据库 schema 为准。",
        ],
    )

    doc.add_heading("附录 B：建议验收命令", level=1)
    add_table(
        doc,
        ["类别", "命令", "目的"],
        [
            ("Python 编译", "python -m compileall backend tests knowledge_pipeline scripts", "检查语法。"),
            ("后端测试", "python -m pytest -q", "执行单元与回归测试。"),
            ("前端构建", "cd frontend && npm run build", "验证前端可构建。"),
            ("数据库运行态", "python tests/productization/check_postgres_runtime.py", "验证 PostgreSQL 和 Alembic。"),
            ("任务中心", "python tests/productization/check_task_center_runtime.py", "验证任务状态流转。"),
            ("Docker", "docker compose --env-file .env.docker config --quiet && docker compose --env-file .env.docker up -d --build", "验证容器部署。"),
            ("RAG", "python knowledge_pipeline/test_rag_search_after_import.py", "验证知识库检索。"),
        ],
    )

    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.text = "智能运营分析项目全面分析说明文档 | 生成日期：2026-06-10 | 不包含真实密钥"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    out = OUT
    if out.exists():
        out = ROOT / f"《智能运营分析项目》项目全面分析说明文档_v20260610_{datetime.now().strftime('%H%M%S')}.docx"
    doc.save(str(out))
    return out


if __name__ == "__main__":
    print(build_doc())
