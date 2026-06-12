from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pymysql
from docx import Document
from docx.shared import Pt

from automation_common import load_config, now_text
from database_utils import CORE_DATASET_TABLES, RESULT_FILE_TABLES, get_database_config


OUTPUT_DIR = Path("output")

TABLE_PURPOSE_OVERRIDES = {
    "raw_da_price": "存储当前最新维护窗口内的日前电价原始快照数据，用于主表构建和后续回溯。",
    "raw_rt_price": "存储当前最新维护窗口内的实时电价原始快照数据，用于价差分析和增强特征。",
    "raw_actual_load": "存储当前最新维护窗口内的实际负荷原始快照数据。",
    "raw_forecast_load_history": "存储 PJM 历史负荷预测全量快照，用于筛选固定版本预测。",
    "raw_forecast_load_selected": "存储供预测引擎直接使用的固定版本负荷预测表。",
    "raw_weather": "存储天气原始小时级快照数据。",
    "model_master_table": "存储当前预测引擎使用的建模主表快照，是后续建模和推理的直接输入。",
    "meta_data_dictionary_fields": "存储字段字典说明。",
    "meta_data_dictionary_coverage": "存储各数据表时间覆盖范围、缺失量和记录数统计。",
    "meta_data_dictionary_sources": "存储各数据源来源网址和取数口径。",
    "ai_input_summary_runs": "存储每次运行生成的 AI 输入摘要，便于回看每次报告到底基于什么事实生成。",
    "ai_report_runs": "存储每次生成的 AI 报告正文、结构化内容、生成方式和 Word 文件路径。",
    "dispatch_manifest_runs": "存储每次归档派发的输出目录、归档清单和派发结果。",
    "pipeline_run_events": "存储整条自动化链路的步骤级事件日志，可用于追踪每次运行进度与失败位置。",
    "result_table_import_registry": "存储每次结果表入库登记信息，对应文件名、目标表和记录数。",
}

FIELD_DESCRIPTION_MAP = {
    "run_id": "一次完整运行的唯一标识，用于串联所有结果表与 AI 报告。",
    "run_mode": "本次运行模式，例如 full、refresh_data、skip_prediction、prediction_report_only。",
    "run_started_at": "本次总控开始时间。",
    "generated_at": "该条摘要或报告记录生成时间。",
    "summary_signature": "基于本次摘要内容计算的签名，用于区分报告是否对应同一批结果。",
    "forecast_mode": "未来24小时预测口径，例如正式前瞻版或演示版。",
    "forecast_start": "未来24小时预测起始时刻。",
    "forecast_end": "未来24小时预测结束时刻。",
    "next_24h_avg_price": "未来24小时预测均价。",
    "next_24h_max_price": "未来24小时预测最高价。",
    "next_24h_max_hour": "未来24小时预测最高价对应时段。",
    "next_24h_min_price": "未来24小时预测最低价。",
    "next_24h_min_hour": "未来24小时预测最低价对应时段。",
    "peak_valley_spread": "预测最高价与最低价之间的峰谷价差。",
    "final_model": "本次结果对应的最终模型名称。",
    "rmse": "误差指标 RMSE。",
    "mae": "误差指标 MAE。",
    "anomaly_count": "异常点数量。",
    "anomaly_ratio_pct": "异常点占比百分比。",
    "summary_json": "完整 AI 输入摘要 JSON。",
    "report_json": "完整 AI 报告 JSON。",
    "full_report_text": "完整报告拼接文本，方便其它应用直接检索或调用。",
    "risk_level": "报告风险等级。",
    "key_hours_json": "报告中提取的重点时段 JSON。",
    "executive_summary": "执行摘要正文。",
    "market_overview": "市场概览正文。",
    "next_24h_trend": "未来24小时趋势正文。",
    "peak_risk": "高峰风险正文。",
    "operation_advice_json": "运营建议 JSON 数组。",
    "management_summary": "管理层摘要正文。",
    "alert_message": "告警摘要正文。",
    "limitations": "局限说明正文。",
    "generator_provider": "报告生成服务提供方。",
    "generator_model": "报告生成模型名称。",
    "generator_mode": "报告生成模式，例如 llm_dynamic、llm_hybrid_dynamic、fallback_dynamic。",
    "word_report_path": "Word 报告文件路径。",
    "word_report_size_bytes": "Word 报告文件大小。",
    "output_folder": "本次归档输出总目录。",
    "manifest_json": "归档清单 JSON。",
    "stage_name": "流程步骤名称。",
    "status": "步骤状态。",
    "message": "步骤说明或错误信息。",
    "event_time": "步骤事件时间。",
    "source_file_name": "对应的原始结果文件名。",
    "target_table_name": "写入到数据库中的目标表名。",
    "row_count": "写入记录数。",
    "imported_at": "结果入库时间。",
    "datetime": "小时级时间戳。",
    "datetime_utc": "UTC 时间戳。",
    "node_id": "节点 ID。",
    "node_name": "节点名称。",
    "da_price": "日前电价。",
    "rt_price": "实时电价。",
    "actual_load": "实际负荷。",
    "forecast_load": "预测负荷。",
    "forecast_evaluated_at": "负荷预测发布时间。",
    "forecast_evaluated_at_utc": "负荷预测 UTC 发布时间。",
    "temperature": "气温。",
    "wind_speed": "风速。",
    "precipitation": "降水。",
    "region_id": "区域标识。",
    "load_area": "负荷区域。",
    "forecast_area": "预测区域。",
    "weather_source": "天气来源标记。",
    "weather_model": "天气模型来源。",
    "weather_point_name": "天气代表点名称。",
}


def connect_db():
    config = load_config()
    db = get_database_config(config)
    return pymysql.connect(
        host=db.host,
        user=db.user,
        password=db.password,
        port=db.port,
        database=db.database,
        charset=db.charset,
    )


def classify_table(table_name: str) -> str:
    if table_name.startswith("raw_"):
        return "原始快照表"
    if table_name.startswith("model_"):
        return "建模主表"
    if table_name.startswith("result_"):
        return "预测结果历史表"
    if table_name.startswith("ai_"):
        return "AI 报告与摘要表"
    if table_name.startswith("dispatch_"):
        return "归档派发表"
    if table_name.startswith("pipeline_"):
        return "流程事件表"
    if table_name.startswith("meta_"):
        return "元数据说明表"
    return "其它表"


def derive_purpose(table_name: str) -> str:
    if table_name in TABLE_PURPOSE_OVERRIDES:
        return TABLE_PURPOSE_OVERRIDES[table_name]

    inverse_result_map = {v: k for k, v in RESULT_FILE_TABLES.items()}
    if table_name in inverse_result_map:
        return f"存储结果文件《{inverse_result_map[table_name]}》的历史入库记录，按 run_id 追踪每次运行。"

    inverse_core_map = {v: k for k, v in CORE_DATASET_TABLES.items()}
    if table_name in inverse_core_map:
        return f"存储文件《{inverse_core_map[table_name]}》对应的数据快照。"

    return f"用于存储 {table_name} 相关数据。"


def field_description(column_name: str) -> str:
    if column_name in FIELD_DESCRIPTION_MAP:
        return FIELD_DESCRIPTION_MAP[column_name]
    if any("\u4e00" <= ch <= "\u9fff" for ch in column_name):
        return f"该字段沿用结果表原始列名《{column_name}》，含义通常与列名一致。"
    return f"字段《{column_name}》为系统生成或原始表保留字段。"


def build_usage_sql(table_name: str, columns: list[dict]) -> list[str]:
    column_names = [col["column_name"] for col in columns]
    sqls: list[str] = []

    if "run_id" in column_names:
        sqls.append(f"SELECT * FROM `{table_name}` WHERE run_id = '20260330_185124' LIMIT 100;")
    if "datetime" in column_names:
        sqls.append(f"SELECT * FROM `{table_name}` ORDER BY datetime DESC LIMIT 24;")
    if "generated_at" in column_names:
        sqls.append(f"SELECT * FROM `{table_name}` ORDER BY generated_at DESC LIMIT 10;")
    if "event_time" in column_names:
        sqls.append(f"SELECT * FROM `{table_name}` WHERE run_id = '20260330_185124' ORDER BY event_time;")
    if not sqls:
        sqls.append(f"SELECT * FROM `{table_name}` LIMIT 20;")
    return sqls


def fetch_database_metadata() -> tuple[list[dict], dict[str, list[dict]], dict[str, int]]:
    conn = connect_db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT DATABASE()")
            schema_name = cur.fetchone()[0]

            cur.execute(
                """
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = %s
                ORDER BY table_name
                """,
                (schema_name,),
            )
            tables = [row[0] for row in cur.fetchall()]

            table_rows: dict[str, int] = {}
            for table_name in tables:
                cur.execute(f"SELECT COUNT(*) FROM `{table_name}`")
                table_rows[table_name] = int(cur.fetchone()[0])

            cur.execute(
                """
                SELECT table_name, column_name, ordinal_position, column_type, is_nullable, column_key, column_comment
                FROM information_schema.columns
                WHERE table_schema = %s
                ORDER BY table_name, ordinal_position
                """,
                (schema_name,),
            )
            rows = cur.fetchall()
    finally:
        conn.close()

    grouped: dict[str, list[dict]] = defaultdict(list)
    for table_name, column_name, ordinal_position, column_type, is_nullable, column_key, column_comment in rows:
        grouped[table_name].append(
            {
                "column_name": column_name,
                "ordinal_position": ordinal_position,
                "column_type": column_type,
                "is_nullable": is_nullable,
                "column_key": column_key,
                "column_comment": column_comment or "",
                "description": field_description(column_name),
            }
        )

    table_infos = []
    for table_name in tables:
        columns = grouped.get(table_name, [])
        table_infos.append(
            {
                "table_name": table_name,
                "category": classify_table(table_name),
                "purpose": derive_purpose(table_name),
                "row_count": table_rows.get(table_name, 0),
                "columns": columns,
                "usage_sql": build_usage_sql(table_name, columns),
            }
        )
    return table_infos, grouped, table_rows


def write_txt_doc(path: Path, table_infos: list[dict]) -> None:
    lines = [
        "数据库表说明文档",
        f"生成时间：{now_text()}",
        "",
        "说明：",
        "1. 本文档基于当前 MySQL 数据库 dianli-yuce 的真实表结构自动生成。",
        "2. 查询时建议优先使用 run_id 追踪某一次完整运行。",
        "3. raw_ 和 model_ 开头的表是当前快照表；result_、ai_、dispatch_、pipeline_ 开头的表是历史记录表。",
        "",
    ]

    for info in table_infos:
        lines.append("=" * 100)
        lines.append(f"表名：{info['table_name']}")
        lines.append(f"类别：{info['category']}")
        lines.append(f"用途：{info['purpose']}")
        lines.append(f"当前记录数：{info['row_count']}")
        lines.append("常用调用 SQL：")
        for sql in info["usage_sql"]:
            lines.append(sql)
        lines.append("字段说明：")
        for col in info["columns"]:
            lines.append(
                f"- {col['column_name']} | {col['column_type']} | 可空={col['is_nullable']} | 键={col['column_key'] or '-'} | {col['description']}"
            )
        lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")


def write_word_doc(path: Path, table_infos: list[dict]) -> None:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style.font.size = Pt(10.5)

    title = document.add_heading("数据库表说明文档", level=0)
    title.runs[0].font.name = "Microsoft YaHei"
    title.runs[0].font.size = Pt(18)

    intro = document.add_paragraph()
    intro.add_run(f"生成时间：{now_text()}\n")
    intro.add_run("数据库：dianli-yuce\n")
    intro.add_run("说明：当前文档基于数据库真实表结构自动生成。")

    category_counts: dict[str, int] = defaultdict(int)
    for info in table_infos:
        category_counts[info["category"]] += 1

    document.add_heading("总体概览", level=1)
    document.add_paragraph(f"当前共识别 {len(table_infos)} 张表。")
    for category, count in sorted(category_counts.items()):
        document.add_paragraph(f"{category}：{count} 张", style="List Bullet")

    document.add_heading("使用原则", level=1)
    usage_rules = [
        "按运行历史追踪时，优先使用 run_id 作为主查询条件。",
        "raw_ 和 model_ 表用于当前预测输入快照，通常直接读取最新整表。",
        "result_ 表用于追踪每次预测引擎输出结果，必须带 run_id 查询。",
        "ai_ 表用于调用摘要、报告正文、风险等级和 Word 报告路径。",
        "pipeline_ 表用于排查某次运行卡在哪个步骤。",
    ]
    for item in usage_rules:
        document.add_paragraph(item, style="List Bullet")

    for info in table_infos:
        document.add_heading(info["table_name"], level=1)
        document.add_paragraph(f"类别：{info['category']}")
        document.add_paragraph(f"用途：{info['purpose']}")
        document.add_paragraph(f"当前记录数：{info['row_count']}")

        document.add_paragraph("常用调用 SQL：")
        for sql in info["usage_sql"]:
            document.add_paragraph(sql, style="List Bullet")

        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = table.rows[0].cells
        headers[0].text = "字段名"
        headers[1].text = "类型"
        headers[2].text = "可空"
        headers[3].text = "键"
        headers[4].text = "说明"

        for col in info["columns"]:
            row = table.add_row().cells
            row[0].text = str(col["column_name"])
            row[1].text = str(col["column_type"])
            row[2].text = str(col["is_nullable"])
            row[3].text = str(col["column_key"] or "-")
            row[4].text = str(col["description"])

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    table_infos, _, _ = fetch_database_metadata()

    txt_path = OUTPUT_DIR / "数据库表说明文档_调用版.txt"
    docx_path = OUTPUT_DIR / "数据库表说明文档_调用版.docx"

    write_txt_doc(txt_path, table_infos)
    write_word_doc(docx_path, table_infos)

    print(txt_path.resolve())
    print(docx_path.resolve())


if __name__ == "__main__":
    main()
