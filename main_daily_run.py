from __future__ import annotations

import argparse
import os
import time

from openpyxl import load_workbook

from automation_common import format_duration, get_pipeline_paths, load_config, now_compact, now_text, setup_run_logger
from database_utils import save_pipeline_event
from services.config_service import disable_database_if_unavailable
from services.operations_service import run_model_auto_optimize
from services.pipeline_steps import run_pipeline_step


SCRIPT_SEQUENCE = [
    "00_local_model_inventory.py",
    "01_run_prediction.py",
    "02_build_ai_summary.py",
    "03_llm_generate_report.py",
    "04_dispatch_report.py",
]
DATA_REFRESH_SCRIPT = "fetch_power_market_data.py"

AI_SUMMARY_REQUIRED_FILES = [
    "10_模型评估结果.xlsx",
    "11_高峰尖刺专项评估.xlsx",
    "13_特征重要性.xlsx",
    "14_异常波动明细.xlsx",
    "15_按小时误差统计.xlsx",
    "16_滚动回测结果.xlsx",
    "19_业务统计摘要.xlsx",
]


def _xlsx_data_rows(path) -> int:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        return max(int(worksheet.max_row or 0) - 1, 0)
    finally:
        workbook.close()


def build_run_mode(args: argparse.Namespace) -> str:
    if args.model_auto_optimize:
        return "model_auto_optimize"
    if args.retrain_model:
        return "retrain_model"
    if args.refresh_data and args.fast_forecast:
        return "refresh_fast_forecast"
    if args.fast_forecast:
        return "fast_forecast"
    if args.refresh_data:
        return "refresh_data"
    if args.skip_prediction:
        return "skip_prediction"
    if args.prediction_report_only:
        return "prediction_report_only"
    return "full"


def run_script(script_name: str, root_dir, env: dict[str, str], log, config: dict) -> None:
    start = time.perf_counter()
    log(f"开始执行：{script_name}")
    save_pipeline_event(config, script_name, "started", "脚本开始执行")
    try:
        run_pipeline_step(script_name, config=config, env=env, log=log)
    except SystemExit as exc:
        save_pipeline_event(config, script_name, "failed", f"服务步骤退出：{exc.code}")
        raise
    except Exception as exc:
        save_pipeline_event(config, script_name, "failed", f"服务步骤异常：{exc}")
        raise
    elapsed = format_duration(time.perf_counter() - start)
    log(f"执行完成：{script_name}，耗时：{elapsed}")
    save_pipeline_event(config, script_name, "completed", f"执行成功，耗时：{elapsed}")


def validate_skip_prediction_inputs(paths, log) -> None:
    if not paths.result_table_dir.exists():
        raise SystemExit(f"skip-prediction 前置检查失败：结果表目录不存在：{paths.result_table_dir}")

    forecast_path = paths.result_table_dir / "18_未来24小时预测结果_正式版.xlsx"
    if not forecast_path.exists():
        raise SystemExit(
            f"skip-prediction 前置检查失败：未找到正式预测表 {forecast_path.name}。"
            "请先执行 --prediction-report-only 或完整预测流程。"
        )
    forecast_rows = _xlsx_data_rows(forecast_path)
    if forecast_rows < 24:
        raise SystemExit(
            f"skip-prediction 前置检查失败：{forecast_path.name} 只有 {forecast_rows} 行，"
            "不足 24 行。请先重新执行预测。"
        )

    missing = [name for name in AI_SUMMARY_REQUIRED_FILES if not (paths.result_table_dir / name).exists()]
    master_path = paths.data_dir / "master_table.xlsx"
    if not master_path.exists():
        missing.append(str(master_path))
    if missing:
        raise SystemExit(
            "skip-prediction 前置检查失败：AI 报告依赖的结果文件不完整："
            + "、".join(missing)
            + "。请先执行 --prediction-report-only 或完整预测流程。"
        )
    log("skip-prediction 前置检查通过：正式预测表和 AI 摘要依赖文件齐全。")


def build_sequence(args: argparse.Namespace) -> list[str]:
    if args.skip_prediction:
        return [item for item in SCRIPT_SEQUENCE if item != "01_run_prediction.py"]
    sequence = SCRIPT_SEQUENCE.copy()
    if args.prediction_report_only:
        sequence = ["01_run_prediction.py", "02_build_ai_summary.py", "03_llm_generate_report.py", "04_dispatch_report.py"]
    if args.refresh_data and not args.skip_prediction and not args.prediction_report_only:
        sequence.insert(1, DATA_REFRESH_SCRIPT)
    return sequence


def main() -> None:
    parser = argparse.ArgumentParser(description="电价预测 + AI 日报自动化总控")
    parser.add_argument("--skip-prediction", action="store_true", help="跳过预测，只基于现有结果生成 AI 报告")
    parser.add_argument("--refresh-data", action="store_true", help="先刷新数据，再执行后续流程")
    parser.add_argument("--prediction-report-only", action="store_true", help="执行预测并生成 AI 报告，但不刷新数据")
    parser.add_argument("--fast-forecast", action="store_true", help="加载 Active 模型快速预测，不重新训练")
    parser.add_argument("--retrain-model", action="store_true", help="完整重训并登记 candidate 模型")
    parser.add_argument("--model-auto-optimize", action="store_true", help="执行误差记忆、退化判断和必要时自动重训")
    args = parser.parse_args()

    config = load_config()
    paths = get_pipeline_paths(config)
    log, log_path = setup_run_logger(paths.log_dir, "main_daily_run")

    run_id = now_compact()
    run_mode = build_run_mode(args)
    run_started_at = now_text()
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    env["PIPELINE_RUN_ID"] = run_id
    env["PIPELINE_RUN_MODE"] = run_mode
    env["PIPELINE_RUN_STARTED_AT"] = run_started_at
    if args.fast_forecast:
        env["PIPELINE_FORECAST_MODE"] = "fast"

    context = {
        "run_id": run_id,
        "run_mode": run_mode,
        "run_started_at": run_started_at,
    }

    log(f"总控日志：{log_path}")
    log(f"本次运行 ID：{run_id}")
    log(f"本次运行模式：{run_mode}")

    if args.skip_prediction:
        validate_skip_prediction_inputs(paths, log)

    db_ok = disable_database_if_unavailable(config, log=log)
    if not db_ok:
        log("数据库同步、模型注册和追踪落库将被跳过。")

    total_start = time.perf_counter()
    save_pipeline_event(config, "main_daily_run", "started", f"总控开始，模式：{run_mode}")

    if args.model_auto_optimize:
        result = run_model_auto_optimize(config, run_context=context, env=env, log=log)
        log(f"模型自动优化完成：{result}")
        total_elapsed = format_duration(time.perf_counter() - total_start)
        save_pipeline_event(config, "main_daily_run", "completed", f"模型自动优化完成，总耗时：{total_elapsed}")
        return

    for script_name in build_sequence(args):
        run_script(script_name, paths.root_dir, env, log, config)

    total_elapsed = format_duration(time.perf_counter() - total_start)
    log(f"本次自动化任务全部完成，总耗时：{total_elapsed}。")
    save_pipeline_event(config, "main_daily_run", "completed", f"总控完成，总耗时：{total_elapsed}")


if __name__ == "__main__":
    main()
