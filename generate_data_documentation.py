from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


OUTPUT_DIR = Path("output")
DOCX_PATH = OUTPUT_DIR / "数据说明文档_建模版.docx"


DATASET_META = {
    "da_price_raw": {
        "title": "1. da_price_raw.xlsx：日前小时电价原始表",
        "role": "预测目标原始表，提供 DOM 区域日前小时电价及其分解项。",
        "grain": "一行 = DOM 节点 1 个小时",
        "primary_key": "datetime + node_id",
        "join_key": "与其他表按 datetime 左连接；空间上固定为 DOM。",
        "modeling_use": [
            "作为主预测目标 y，字段为 da_price。",
            "可以单独分析价格峰值、波动率、尖峰小时、分量贡献。",
            "da_congestion_price、da_marginal_loss_price、da_system_energy_price 可用于解释价格形成机制，但若预测未来 24 小时日前电价，严格来说这些分量本身通常不应直接作为未来特征输入，而更适合做历史诊断、标签拆解或衍生统计特征。",
        ],
        "risks": [
            "存在夏令时切换，因此相邻时间差大多数为 1 小时，少数为 2 小时，这是正常现象，不是缺失。",
            "价格极值较大，存在尖峰价格，建模前建议做异常值策略评估，例如保留原值、对数变换、winsorize 对比。",
        ],
        "field_meanings": {
            "datetime": "PJM EPT 本地时间，对建模建议作为主时间键。",
            "datetime_utc": "对应 UTC 时间，便于跨源核对时区。",
            "node_id": "PJM 定价节点 ID，本项目固定为 34964545。",
            "node_name": "节点名称，本项目固定为 DOM。",
            "da_price": "日前总 LMP，单位 USD/MWh，是主预测目标。",
            "da_congestion_price": "日前拥塞分量。",
            "da_marginal_loss_price": "日前边际损耗分量。",
            "da_system_energy_price": "日前系统能量分量。",
            "type": "节点类型，本项目固定为 ZONE。",
            "row_is_current": "是否当前版本，本项目均为 True。",
            "version_nbr": "版本号，本项目均为 1。",
        },
    },
    "actual_load_raw": {
        "title": "2. actual_load_raw.xlsx：小时实际负荷原始表",
        "role": "提供 DOM 负荷区小时实际负荷，是最重要的解释变量之一。",
        "grain": "一行 = DOM 负荷区 1 个小时",
        "primary_key": "datetime + load_area",
        "join_key": "按 datetime 与主表合并；空间口径固定为 DOM。",
        "modeling_use": [
            "适合构造负荷滞后项、滚动均值、滚动标准差、同比/环比特征。",
            "适合做价格解释分析，例如价格对负荷的弹性、负荷高位时的风险分层。",
            "若要严格做未来 24 小时预测，实际负荷不能直接作为未来时点输入，应只用于历史滞后特征和训练阶段标签对齐。",
        ],
        "risks": [
            "该表比日前电价晚 1 天落地，因此在主表中已裁掉最后一天，保证核心字段齐全。",
            "is_verified 既有 True 也有 False，说明部分小时仍可能处于未最终核验状态。",
        ],
        "field_meanings": {
            "datetime": "PJM EPT 本地时间。",
            "datetime_utc": "对应 UTC 时间。",
            "load_area": "负荷区域，本项目固定为 DOM。",
            "zone": "传输分区，本项目固定为 DOM。",
            "mkt_region": "市场区域，本项目固定为 SOUTH。",
            "nerc_region": "NERC 区域，本项目固定为 SERC。",
            "actual_load": "小时实际负荷，单位 MW。",
            "is_verified": "是否经公司核验，布尔值。",
        },
    },
    "forecast_load_raw": {
        "title": "3. forecast_load_raw.xlsx：历史负荷预测全量原始表",
        "role": "保留 PJM 历史负荷预测的完整快照，可用于追溯不同发布时间版本。",
        "grain": "一行 = 1 次预测发布时间 + 1 个被预测小时",
        "primary_key": "forecast_evaluated_at + datetime + forecast_area",
        "join_key": "不建议直接与主表一对一合并，应先做版本筛选。",
        "modeling_use": [
            "适合研究不同预测发布时间的版本差异。",
            "适合构造“最近一次预测”“前一日最后版本预测”“预测修正幅度”等特征。",
            "是 forecast_load_selected 的上游原始来源，若未来要调整预测版本规则，应从本表重算。",
        ],
        "risks": [
            "同一个目标小时会出现多个版本，因此 datetime 大量重复，这是正常现象。",
            "个别目标小时在 PJM 历史源中没有找到完整的前一日版本，这也是后续筛选时需要处理的边界情况。",
        ],
        "field_meanings": {
            "forecast_evaluated_at": "预测发布时间，EPT。",
            "forecast_evaluated_at_utc": "预测发布时间，UTC。",
            "datetime": "被预测的目标小时，EPT。",
            "datetime_utc": "被预测的目标小时，UTC。",
            "forecast_area": "预测区域，本项目固定为 DOM。",
            "forecast_load": "对应版本下对目标小时的负荷预测值，单位 MW。",
        },
    },
    "forecast_load_selected": {
        "title": "4. forecast_load_selected.xlsx：建模版负荷预测表",
        "role": "从历史负荷预测全量表中筛选出的建模版本，用于与价格主表一对一对齐。",
        "grain": "一行 = DOM 区域 1 个小时",
        "primary_key": "datetime + forecast_area",
        "join_key": "按 datetime 合并到主表。",
        "modeling_use": [
            "这是未来 24 小时预测场景下最关键的前瞻性特征之一。",
            "当前筛选规则是：优先取目标日前一自然日的最后一个可用预测版本；若不存在，再取目标小时前最后一个可用版本。",
            "后续可基于本表继续构造 forecast_load_lag24、forecast_change_vs_previous_issue 等特征。",
        ],
        "risks": [
            "由于 PJM 历史源边界问题，原本有极少数小时无法找到合规版本，最终主表已将这些小时剔除。",
            "若未来你希望更贴近真实交易口径，可以把筛选规则固定为某个发布时间，如前一天 23:45 或当日 05:45，再重跑脚本。",
        ],
        "field_meanings": {
            "datetime": "目标小时，EPT。",
            "datetime_utc": "目标小时，UTC。",
            "forecast_area": "预测区域，本项目固定为 DOM。",
            "forecast_load": "筛选后的负荷预测值，单位 MW。",
            "forecast_evaluated_at": "该预测值对应的发布时间，EPT。",
            "forecast_evaluated_at_utc": "该预测值对应的发布时间，UTC。",
        },
    },
    "weather_raw": {
        "title": "5. weather_raw.xlsx：小时天气原始表",
        "role": "提供 DOM 区域代表点的小时天气特征，包括温度、风速、降水。",
        "grain": "一行 = 代表点 1 个小时",
        "primary_key": "datetime",
        "join_key": "按 datetime 合并到主表；空间口径为 Richmond, Virginia 代表点。",
        "modeling_use": [
            "温度通常与负荷高度相关，适合构造冷暖偏离、温度平方项、冷负荷/热负荷分段特征。",
            "风速和降水更适合做解释类辅助特征，也可以与季节、小时交互。",
            "如后续扩展到多节点或更精细区域，建议引入多个天气点后做加权汇总，而不是一直固定单点。",
        ],
        "risks": [
            "当前天气口径是单代表点近似，不等于 DOM 全区域平均天气。",
            "最近几天 ERA5 返回缺口，已通过同接口的 best_match 做回填，因此 weather_model 存在两个取值。",
        ],
        "field_meanings": {
            "datetime": "本地时间戳，已按 America/New_York 返回。",
            "temperature": "2 米气温，单位 °C。",
            "wind_speed": "10 米风速，单位 km/h。",
            "precipitation": "逐小时降水量，单位 mm。",
            "latitude": "天气点纬度。",
            "longitude": "天气点经度。",
            "weather_model": "天气模型标识，主要为 era5，少量为 best_match_backfill。",
            "weather_point_name": "天气代表点名称，本项目为 Richmond, Virginia。",
        },
    },
    "rt_price_raw": {
        "title": "6. rt_price_raw.xlsx：实时小时电价原始表",
        "role": "提供 DOM 区域实时小时电价，是增强分析表，不是 MVP 必要输入。",
        "grain": "一行 = DOM 节点 1 个小时",
        "primary_key": "datetime + node_id",
        "join_key": "按 datetime 与主表合并。",
        "modeling_use": [
            "适合构造价格偏差分析，例如 rt_price - da_price。",
            "适合做市场波动分析、交易复盘、模型解释增强。",
            "若你的目标是纯日前价格预测，实时价不能作为未来时点输入，但可用于训练后诊断、误差解释或历史统计特征。",
        ],
        "risks": [
            "与实际负荷一样，该表比日前电价晚 1 天落地，因此主表截止到 2026-03-26 23:00。",
            "实时价波动显著高于日前价，且可出现负价，建模时要注意尺度稳定性。",
        ],
        "field_meanings": {
            "datetime": "PJM EPT 本地时间。",
            "datetime_utc": "对应 UTC 时间。",
            "node_id": "PJM 定价节点 ID，本项目固定为 34964545。",
            "node_name": "节点名称，本项目固定为 DOM。",
            "rt_price": "实时总 LMP，单位 USD/MWh。",
            "rt_congestion_price": "实时拥塞分量。",
            "rt_marginal_loss_price": "实时边际损耗分量。",
            "rt_system_energy_price": "实时系统能量分量。",
            "type": "节点类型，本项目固定为 ZONE。",
            "row_is_current": "是否当前版本。",
            "version_nbr": "版本号。",
        },
    },
    "master_table": {
        "title": "7. master_table.xlsx：建模主表",
        "role": "最终建模主表，已经完成时间统一、数据对齐、去重和核心缺失清理。",
        "grain": "一行 = DOM 区域 1 个小时",
        "primary_key": "datetime",
        "join_key": "这是最终训练输入表，不需要再与其他表联接即可开始建模。",
        "modeling_use": [
            "推荐把 da_price 作为目标变量，使用其他解释变量与日期特征做监督学习。",
            "推荐从本表继续生成 lag 特征、rolling 特征、节假日与天气交互特征。",
            "对于严格的未来 24 小时日前预测，建议训练时只使用在预测时点可得的变量，例如 forecast_load、天气预报替代变量、历史价格滞后、历史负荷滞后、日历特征。",
            "如果你只做解释分析或回顾性建模，则可同时使用 actual_load 与 rt_price 相关派生指标。",
        ],
        "risks": [
            "本表为了保证核心字段完整，已经裁掉最后 1 天和 6 个缺少合规负荷预测版本的小时，因此行数少于理论全量小时数。",
            "weather_source 当前是来源说明字段，不是用于回归的有效信号，建模时通常应剔除。",
            "node_id、node_name、region_id、load_area、forecast_area 等在本项目中几乎为常量，建模时信息量有限，可不作为数值特征输入。",
        ],
        "field_meanings": {
            "datetime": "主表 EPT 时间键。",
            "datetime_utc": "主表 UTC 时间键。",
            "region_id": "区域标识，本项目固定为 DOM。",
            "node_id": "价格节点 ID。",
            "node_name": "价格节点名。",
            "da_price": "预测目标变量，日前总 LMP。",
            "da_congestion_price": "日前拥塞分量。",
            "da_marginal_loss_price": "日前边际损耗分量。",
            "da_system_energy_price": "日前系统能量分量。",
            "rt_price": "实时总 LMP。",
            "rt_congestion_price": "实时拥塞分量。",
            "rt_marginal_loss_price": "实时边际损耗分量。",
            "rt_system_energy_price": "实时系统能量分量。",
            "price_spread_rt_minus_da": "实时价减日前价。",
            "actual_load": "小时实际负荷。",
            "forecast_load": "筛选后的负荷预测。",
            "forecast_evaluated_at": "该预测对应的发布时间。",
            "forecast_evaluated_at_utc": "该预测对应的 UTC 发布时间。",
            "load_area": "负荷区。",
            "forecast_area": "预测区。",
            "mkt_region": "市场区域。",
            "nerc_region": "NERC 区域。",
            "temperature": "小时气温。",
            "wind_speed": "小时风速。",
            "precipitation": "小时降水。",
            "hour": "小时序号 0-23。",
            "day_of_week": "星期序号，0=周一。",
            "month": "月份。",
            "year": "年份。",
            "season": "季节标签。",
            "is_weekend": "是否周末。",
            "is_holiday": "是否美国联邦节假日。",
            "is_verified": "实际负荷是否核验。",
            "weather_source": "天气来源说明。",
        },
    },
    "data_dictionary": {
        "title": "8. data_dictionary.xlsx：字段字典与覆盖说明",
        "role": "元数据工作簿，不直接参与建模，但用于快速查字段口径、时间覆盖和来源链接。",
        "grain": "按工作表分别记录字段定义、覆盖范围、来源信息",
        "primary_key": "无单一主键",
        "join_key": "不参与主表合并。",
        "modeling_use": [
            "建模前快速核对字段含义与单位。",
            "复盘数据更新时间范围与样本行数。",
            "查找来源链接，便于后续扩展抓取。",
        ],
        "risks": [
            "这是说明型文件，不是训练样本。",
        ],
        "field_meanings": {},
    },
}


def set_default_style(document: Document) -> None:
    styles = document.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "黑体"
    styles["Heading 1"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["Heading 2"].font.name = "黑体"
    styles["Heading 2"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def add_page_number(section) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    run = footer.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def load_dataset(name: str) -> pd.DataFrame:
    if name == "data_dictionary":
        return pd.read_excel(OUTPUT_DIR / "data_dictionary.xlsx", sheet_name="fields")
    return pd.read_excel(OUTPUT_DIR / f"{name}.xlsx")


def dataset_overview_table(document: Document, coverage_df: pd.DataFrame) -> None:
    document.add_heading("一、数据交付总览", level=1)
    document.add_paragraph(
        "本项目当前交付 7 张数据表和 1 个元数据说明工作簿。核心建模表为 master_table.xlsx，其他表保留原始或半加工口径，便于后续重算特征、回溯问题和调整建模规则。"
    )
    table = document.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["数据集", "行数", "列数", "开始时间", "结束时间", "总缺失"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value

    mapping = {
        "da_price_raw": "da_price_raw.xlsx",
        "rt_price_raw": "rt_price_raw.xlsx",
        "actual_load_raw": "actual_load_raw.xlsx",
        "forecast_load_raw": "forecast_load_raw.xlsx",
        "forecast_load_selected": "forecast_load_selected.xlsx",
        "weather_raw": "weather_raw.xlsx",
        "master_table": "master_table.xlsx",
    }
    for _, row in coverage_df.iterrows():
        new_row = table.add_row().cells
        new_row[0].text = mapping.get(row["dataset"], row["dataset"])
        new_row[1].text = str(int(row["rows"]))
        new_row[2].text = str(int(row["columns"]))
        new_row[3].text = "" if pd.isna(row["start_datetime"]) else str(row["start_datetime"])
        new_row[4].text = "" if pd.isna(row["end_datetime"]) else str(row["end_datetime"])
        new_row[5].text = str(int(row["total_missing_cells"]))

    document.add_paragraph(
        "说明：master_table 已做核心字段完整性约束，因此行数少于各原始表。少量小时被剔除的原因主要是 PJM 历史负荷预测边界版本缺失，以及实时价/实际负荷比日前价晚一天发布。"
    )


def add_bullet_list(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_field_table(document: Document, field_meanings: dict[str, str], dataset_columns: list[str]) -> None:
    if not dataset_columns:
        return
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    header[0].text = "字段名"
    header[1].text = "字段说明"
    header[2].text = "建模提示"

    for col in dataset_columns:
        row = table.add_row().cells
        row[0].text = str(col)
        row[1].text = field_meanings.get(col, "未单独补充说明，建议结合字段字典查看。")
        if col in {"datetime", "datetime_utc", "forecast_evaluated_at", "forecast_evaluated_at_utc"}:
            hint = "时间字段，适合派生小时、星期、节假日、滚动窗口和滞后特征。"
        elif col in {"da_price", "rt_price"}:
            hint = "价格字段，注意尖峰与负值。"
        elif col in {"actual_load", "forecast_load"}:
            hint = "负荷字段，适合做滞后、滚动和差值特征。"
        elif col in {"temperature", "wind_speed", "precipitation"}:
            hint = "天气字段，建议与季节、小时做交互。"
        else:
            hint = "根据是否为常量列、解释列或标签列决定是否纳入模型。"
        row[2].text = hint


def add_quality_summary(document: Document, name: str, df: pd.DataFrame) -> None:
    document.add_heading("数据质量与统计摘要", level=3)
    p = document.add_paragraph()
    p.add_run("表大小：").bold = True
    p.add_run(f"{len(df)} 行，{len(df.columns)} 列。")
    if "datetime" in df.columns:
        dt = pd.to_datetime(df["datetime"])
        p = document.add_paragraph()
        p.add_run("时间覆盖：").bold = True
        p.add_run(f"{dt.min()} 到 {dt.max()}。")

        gap_counts = dt.sort_values().diff().dropna().value_counts()
        top_gaps = [f"{str(idx)} x {int(val)}" for idx, val in gap_counts.head(3).items()]
        p = document.add_paragraph()
        p.add_run("相邻时间间隔：").bold = True
        p.add_run("；".join(top_gaps) if top_gaps else "无。")

    missing = df.isna().sum()
    missing = missing[missing > 0].sort_values(ascending=False)
    p = document.add_paragraph()
    p.add_run("缺失情况：").bold = True
    if missing.empty:
        p.add_run("当前表无缺失。")
    else:
        pairs = [f"{col}={int(val)}" for col, val in missing.items()]
        p.add_run("；".join(pairs))

    numeric = df.select_dtypes(include="number")
    if not numeric.empty:
        desc = numeric.describe().T[["mean", "std", "min", "50%", "max"]].round(4)
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["数值字段", "均值", "标准差", "最小值", "中位数", "最大值"]
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for idx, row in desc.head(12).iterrows():
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = str(row["mean"])
            cells[2].text = str(row["std"])
            cells[3].text = str(row["min"])
            cells[4].text = str(row["50%"])
            cells[5].text = str(row["max"])


def add_dataset_section(document: Document, dataset_name: str, df: pd.DataFrame) -> None:
    meta = DATASET_META[dataset_name]
    document.add_heading(meta["title"], level=1)
    p = document.add_paragraph()
    p.add_run("用途定位：").bold = True
    p.add_run(meta["role"])
    p = document.add_paragraph()
    p.add_run("粒度：").bold = True
    p.add_run(meta["grain"])
    p = document.add_paragraph()
    p.add_run("建议主键：").bold = True
    p.add_run(meta["primary_key"])
    p = document.add_paragraph()
    p.add_run("与主表关系：").bold = True
    p.add_run(meta["join_key"])

    add_quality_summary(document, dataset_name, df)

    document.add_heading("字段说明", level=3)
    add_field_table(document, meta["field_meanings"], list(df.columns))

    document.add_heading("建模使用建议", level=3)
    add_bullet_list(document, meta["modeling_use"])

    document.add_heading("建模风险与注意事项", level=3)
    add_bullet_list(document, meta["risks"])


def add_modeling_recommendations(document: Document, master_df: pd.DataFrame) -> None:
    document.add_heading("二、基于当前数据的建模建议", level=1)
    document.add_paragraph(
        "如果你当前目标是“未来 24 小时日前电价预测”，建议把 master_table.xlsx 作为唯一训练入口，再在其基础上追加衍生特征。核心原则是区分预测时点可得信息和不可得信息，避免未来信息泄漏。"
    )

    document.add_heading("推荐目标变量", level=2)
    add_bullet_list(
        document,
        [
            "主目标：da_price。",
            "若希望提高稳定性，可同时保留对数变换目标，如 log1p(da_price)。",
            "若后续做分类辅助任务，可构造是否高于某阈值的尖峰价标签。",
        ],
    )

    document.add_heading("推荐直接可用特征", level=2)
    add_bullet_list(
        document,
        [
            "日历类：hour、day_of_week、month、season、is_weekend、is_holiday。",
            "历史价格类：da_price 的 lag_1、lag_2、lag_24、lag_48、lag_168，以及 rolling_mean_24、rolling_std_24。",
            "负荷类：forecast_load 以及其滞后项、滚动均值、与历史实际负荷差值。",
            "天气类：temperature、wind_speed、precipitation，以及温度平方项、冷暖分段特征。",
            "增强类：price_spread_rt_minus_da 的历史滞后统计，不能直接把未来 rt_price 当作输入。",
        ],
    )

    document.add_heading("建议谨慎使用或剔除的列", level=2)
    add_bullet_list(
        document,
        [
            "常量列：region_id、node_id、node_name、load_area、forecast_area、mkt_region、nerc_region，在当前单区域项目中信息量很低。",
            "说明列：weather_source 通常不进入模型。",
            "可能泄漏未来信息的列：若预测未来 24 小时日前价，不要直接用未来时点的 actual_load、rt_price、rt_congestion_price 等实时后验变量。",
            "价格分解项：da_congestion_price 等属于与目标同刻的分解结果，训练未来预测模型时不应把未来时刻分解项直接作为输入，但可用其历史滞后统计值。",
        ],
    )

    document.add_heading("推荐验证方式", level=2)
    add_bullet_list(
        document,
        [
            "不要随机切分训练集和测试集，应采用时间顺序切分。",
            "建议至少保留最近 1 至 2 个月作为验证/测试窗口。",
            "如果做滚动验证，建议使用 expanding window 或 rolling window backtest。",
        ],
    )

    document.add_heading("当前样本特征概览", level=2)
    p = document.add_paragraph()
    p.add_run("主表样本量：").bold = True
    p.add_run(f"{len(master_df)} 行。")
    p = document.add_paragraph()
    p.add_run("目标变量范围：").bold = True
    p.add_run(
        f"da_price 最小值 {master_df['da_price'].min():.4f}，最大值 {master_df['da_price'].max():.4f}，均值 {master_df['da_price'].mean():.4f}。"
    )
    p = document.add_paragraph()
    p.add_run("核心解释变量范围：").bold = True
    p.add_run(
        f"actual_load 均值 {master_df['actual_load'].mean():.2f} MW，forecast_load 均值 {master_df['forecast_load'].mean():.2f} MW，temperature 均值 {master_df['temperature'].mean():.2f} °C。"
    )


def add_sources(document: Document, sources_df: pd.DataFrame) -> None:
    document.add_heading("三、数据来源", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "数据集"
    table.rows[0].cells[1].text = "来源链接"
    table.rows[0].cells[2].text = "当前选用区域/节点"

    for _, row in sources_df.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(row["dataset"])
        cells[1].text = str(row["source_url"])
        cells[2].text = str(row["selected_region_or_node"])


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    coverage_df = pd.read_excel(OUTPUT_DIR / "data_dictionary.xlsx", sheet_name="coverage")
    sources_df = pd.read_excel(OUTPUT_DIR / "data_dictionary.xlsx", sheet_name="sources")
    master_df = pd.read_excel(OUTPUT_DIR / "master_table.xlsx")

    document = Document()
    set_default_style(document)
    add_page_number(document.sections[0])

    document.add_heading("电力市场项目数据说明文档", level=0)
    document.add_paragraph(
        "本文档针对当前项目目录中已经整理完成的所有数据文件做逐表说明，重点覆盖数据含义、粒度主键、质量情况、字段口径、建模可用性和注意事项，便于后续直接开展时间序列建模与特征工程。"
    )
    document.add_paragraph(
        "项目范围：DOM 区域，小时级数据，覆盖最近约 24 个月。主表已完成对齐，可直接用于建模。"
    )

    dataset_overview_table(document, coverage_df)

    document.add_page_break()
    for name in [
        "da_price_raw",
        "actual_load_raw",
        "forecast_load_raw",
        "forecast_load_selected",
        "weather_raw",
        "rt_price_raw",
        "master_table",
        "data_dictionary",
    ]:
        df = load_dataset(name)
        add_dataset_section(document, name, df)
        document.add_page_break()

    add_modeling_recommendations(document, master_df)
    document.add_page_break()
    add_sources(document, sources_df)

    document.save(DOCX_PATH)
    print(DOCX_PATH.resolve())


if __name__ == "__main__":
    main()
