from __future__ import annotations

import json
import os
import time
from pathlib import Path

from docx import Document
from docx.shared import Pt

from automation_common import (
    format_duration,
    get_pipeline_paths,
    get_run_context,
    load_config,
    now_text,
    read_json,
    setup_run_logger,
    write_json,
)
from database_utils import save_ai_report, save_pipeline_event
from llm_client import LLMClient


LLM_REQUIRED_KEYS = {
    "executive_summary",
    "market_overview",
    "next_24h_trend",
    "peak_risk",
    "operation_advice",
    "management_summary",
    "alert_message",
    "limitations",
}

DOMAIN_KEYWORDS = ["电价", "价格", "价差", "市场", "负荷", "高峰", "尖峰", "晨峰", "晚峰", "预测", "时段", "风险", "运营", "调度", "回测", "模型", "波动"]
IRRELEVANT_KEYWORDS = ["供应链", "库存", "JIT", "员工", "心理健康", "团队协作", "客户", "销售", "原材料"]

LLM_REPORT_SCHEMA = {
    "type": "object",
    "properties": {
        "executive_summary": {"type": "string"},
        "market_overview": {"type": "string"},
        "next_24h_trend": {"type": "string"},
        "peak_risk": {"type": "string"},
        "operation_advice": {"type": "array", "items": {"type": "string"}},
        "management_summary": {"type": "string"},
        "alert_message": {"type": "string"},
        "limitations": {"type": "string"},
    },
    "required": list(LLM_REQUIRED_KEYS),
}


def build_subset_schema(section_names: list[str]) -> dict:
    properties = {key: LLM_REPORT_SCHEMA["properties"][key] for key in section_names}
    return {
        "type": "object",
        "properties": properties,
        "required": section_names,
    }


def build_llm_report_input(summary: dict) -> dict:
    forecast = summary["forecast_summary"]
    anomaly = summary["anomaly_summary"]
    feature = summary["feature_summary"]
    comparison = summary["comparison_summary"]
    backtest = summary["backtest_summary"]
    return {
        "run_id": summary["run_id"],
        "summary_signature": summary["summary_signature"],
        "forecast_mode": summary["forecast_mode"],
        "report_facts": summary["report_facts"],
        "forecast_summary": {
            "forecast_start": forecast.get("forecast_start"),
            "forecast_end": forecast.get("forecast_end"),
            "next_24h_avg_price": forecast.get("next_24h_avg_price"),
            "next_24h_max_price": forecast.get("next_24h_max_price"),
            "next_24h_max_hour": forecast.get("next_24h_max_hour"),
            "next_24h_min_price": forecast.get("next_24h_min_price"),
            "next_24h_min_hour": forecast.get("next_24h_min_hour"),
            "peak_valley_spread": forecast.get("peak_valley_spread"),
            "top_risk_hours": forecast.get("top_risk_hours", [])[:3],
            "weather_sources": forecast.get("weather_sources", []),
            "forecast_load_sources": forecast.get("forecast_load_sources", []),
        },
        "comparison_summary": comparison,
        "anomaly_summary": {
            "anomaly_count": anomaly.get("anomaly_count"),
            "anomaly_ratio_pct": anomaly.get("anomaly_ratio_pct"),
            "main_risk_hours": anomaly.get("main_risk_hours", [])[:4],
            "top_anomalies": anomaly.get("top_anomalies", [])[:3],
        },
        "feature_summary": {
            "top_features": feature.get("top_features", [])[:5],
        },
        "peak_model_summary": summary.get("peak_model_summary", {}),
        "backtest_summary": {
            "avg_rmse": backtest.get("avg_rmse"),
            "avg_mae": backtest.get("avg_mae"),
            "avg_r2": backtest.get("avg_r2"),
            "worst_fold": backtest.get("worst_fold"),
        },
    }


def build_prompts(summary: dict) -> tuple[str, str]:
    facts = summary["report_facts"]
    report_input = build_llm_report_input(summary)
    system_prompt = (
        "你是电力市场预测结果解读助手。"
        "你只能基于输入摘要中的事实生成中文简体报告。"
        "你必须保留关键数值和关键时段，不得编造数据，不得引入与电力项目无关的内容。"
        "报告中的价格单位统一使用 USD/MWh。"
        "如果事实中没有合同、库存、资金、收益金额、客户、供应链等信息，就绝不能扩展出这些内容。"
        "运营建议只能围绕监控、复核、调度关注、风险提示、数据跟踪等动作。"
        "输出必须是严格 JSON。"
    )

    user_prompt = f"""
请根据以下事实生成“本次运行专属”的电价智能分析报告。

硬性要求：
1. 只能基于输入事实输出，不得编造不存在的数值或时段。
2. 报告必须体现本次运行的具体数值，不要写成任何一次都能复用的通用套话。
3. 请优先引用以下事实：
   - 运行ID：{summary['run_id']}
   - 摘要签名：{summary['summary_signature']}
   - 数据窗口：{facts['data_window_start']} 至 {facts['data_window_end']}
   - 预测窗口：{facts['forecast_start']} 至 {facts['forecast_end']}
   - 未来24小时均价：{facts['avg_price']}
   - 最高价与时段：{facts['max_price']} / {facts['max_hour']}
   - 最低价与时段：{facts['min_price']} / {facts['min_hour']}
   - 峰谷价差：{facts['peak_valley_spread']}
   - 最终模型：{facts['final_model']}
   - 测试集 RMSE：{facts['rmse']}
   - 测试集 MAE：{facts['mae']}
   - 测试集 R2：{facts['r2']}
   - 平均滚动回测 RMSE：{facts['backtest_avg_rmse']}
   - 异常点数量：{facts['anomaly_count']}
   - 异常占比：{facts['anomaly_ratio_pct']}
   - 重点风险时段：{facts['main_risk_hours']}
   - 与昨日均价变化：{facts['vs_yesterday_avg_pct']}
   - 与上周均价变化：{facts['vs_lastweek_avg_pct']}
4. operation_advice 需要 3 到 5 条，每条都要基于本次结果给出。
4.1 operation_advice 不要写“储备资金”“签约采购”“客户交付”“库存”这类超出输入事实的建议。
4.2 请尽量使用输入中的准确时段、准确价格和准确指标。
5. 不要输出 markdown，不要额外解释，只返回 JSON。

本次报告可用的精简输入如下：
{json.dumps(report_input, ensure_ascii=False, indent=2)}
""".strip()
    return system_prompt, user_prompt


def build_regeneration_prompts(summary: dict, missing_sections: list[str]) -> tuple[str, str]:
    forecast = summary["forecast_summary"]
    report_input = build_llm_report_input(summary)
    system_prompt = (
        "你是电力市场报告补写助手。"
        "你只需要重写指定字段，必须严格基于输入事实，必须使用中文简体，必须输出 JSON。"
        "价格单位统一使用 USD/MWh。"
    )
    user_prompt = f"""
请只补写以下字段：{missing_sections}

强约束：
1. executive_summary 必须同时出现最高价时段 {forecast['next_24h_max_hour']} 和最低价时段 {forecast['next_24h_min_hour']}。
2. next_24h_trend 必须同时出现最高价时段 {forecast['next_24h_max_hour']} 和最低价时段 {forecast['next_24h_min_hour']}。
3. peak_risk 或 alert_message 若被要求补写，至少要出现以下高风险时段中的一个：{[item.get('hour') for item in forecast.get('top_risk_hours', [])[:3]]}
4. 不要缩写成标题，不要只写一句空泛概括，要写成完整业务句子。
5. 只返回 JSON，不要解释。

输入事实如下：
{json.dumps(report_input, ensure_ascii=False, indent=2)}
""".strip()
    return system_prompt, user_prompt


def build_fallback_report(summary: dict) -> dict:
    forecast = summary["forecast_summary"]
    model = summary["model_summary"]
    anomaly = summary["anomaly_summary"]
    comparison = summary["comparison_summary"]
    facts = summary["report_facts"]
    risk_hours = forecast.get("top_risk_hours", [])
    key_hours = [item["hour"] for item in risk_hours[:3]] or anomaly.get("main_risk_hours", [])
    max_risk_prob = max((item.get("spike_risk_probability") or 0) for item in risk_hours) if risk_hours else 0
    anomaly_ratio = anomaly.get("anomaly_ratio_pct") or 0

    if max_risk_prob >= 0.35 or anomaly_ratio >= 1.0:
        risk_level = "高"
    elif max_risk_prob >= 0.15 or anomaly_ratio >= 0.5:
        risk_level = "中"
    else:
        risk_level = "低"

    return {
        "executive_summary": (
            f"本次运行基于 {summary['engine_version']} 完成，摘要签名为 {summary['summary_signature']}。"
            f"当前最终模型为 {model['final_model']}，测试集 RMSE 为 {model['rmse']:.4f}，MAE 为 {model['mae']:.4f}。"
            f"未来24小时{summary['forecast_mode']}预测均价为 {forecast['next_24h_avg_price']:.2f} USD/MWh，"
            f"最高价预计出现在 {forecast['next_24h_max_hour']}，最低价预计出现在 {forecast['next_24h_min_hour']}。"
        ),
        "market_overview": (
            f"从本次结果看，模型整体精度维持在可用区间，测试集 R2 为 {model['r2']:.4f}，"
            f"平均滚动回测 RMSE 为 {summary['backtest_summary']['avg_rmse']:.4f}。"
            f"异常点数量为 {anomaly['anomaly_count']} 个，异常占比 {anomaly['anomaly_ratio_pct']:.2f}%，"
            f"说明当前系统适合做日常价格趋势判断和高风险时段提示。"
        ),
        "next_24h_trend": (
            f"未来24小时预测窗口为 {forecast['forecast_start']} 至 {forecast['forecast_end']}。"
            f"均价 {forecast['next_24h_avg_price']:.2f} USD/MWh，最高价 {forecast['next_24h_max_price']:.2f} "
            f"出现在 {forecast['next_24h_max_hour']}，最低价 {forecast['next_24h_min_price']:.2f} 出现在 {forecast['next_24h_min_hour']}，"
            f"峰谷价差为 {forecast['peak_valley_spread']:.2f} USD/MWh。"
            f"与昨日均价相比变动 {comparison['vs_yesterday_avg_pct'] if comparison['vs_yesterday_avg_pct'] is not None else '暂无可比值'}%，"
            f"与上周均价相比变动 {comparison['vs_lastweek_avg_pct'] if comparison['vs_lastweek_avg_pct'] is not None else '暂无可比值'}%。"
        ),
        "peak_risk": (
            f"本次高风险时段主要集中在 {', '.join(key_hours) if key_hours else '少数高峰小时'}。"
            f"其中最高风险概率约为 {max_risk_prob:.3f}。"
            f"历史异常也主要集中在 {', '.join(anomaly.get('main_risk_hours', [])[:4]) or '无明显集中时段'}，"
            "说明晨峰和局部晚峰仍然是重点关注窗口。"
        ),
        "operation_advice": [
            f"围绕 {', '.join(key_hours[:3]) if key_hours else '重点高峰时段'} 做人工复核，优先关注价格尖峰风险。",
            f"将未来24小时均价 {forecast['next_24h_avg_price']:.2f}、最高价 {forecast['next_24h_max_price']:.2f} 和峰谷价差 {forecast['peak_valley_spread']:.2f} 纳入日常运营简报。",
            f"结合异常点数量 {anomaly['anomaly_count']} 个和异常占比 {anomaly['anomaly_ratio_pct']:.2f}% 回看近一周高误差小时，优化高峰时段解释逻辑。",
            f"本次数据窗口覆盖 {facts['data_window_start']} 至 {facts['data_window_end']}，若后续继续扩充最新数据，应优先观察晨峰时段误差变化。",
        ],
        "management_summary": (
            f"本次电价预测运行已完成，使用数据窗口 {facts['data_window_start']} 至 {facts['data_window_end']}。"
            f"模型整体表现稳定，测试集 RMSE 为 {model['rmse']:.4f}。"
            f"未来24小时预计均价 {forecast['next_24h_avg_price']:.2f} USD/MWh，"
            f"高风险时段集中在 {', '.join(key_hours[:3]) if key_hours else '若干高峰时段'}。"
            "建议将该结果作为运营监控和风险提示依据，并在尖峰时段保留人工复核。"
        ),
        "alert_message": (
            f"电价告警：未来24小时最高价预计出现在 {forecast['next_24h_max_hour']}，"
            f"重点关注 {', '.join(key_hours[:3]) if key_hours else '高峰时段'} 的异常波动风险。"
        ),
        "limitations": (
            "本次报告基于本次运行摘要动态生成。"
            "正式前瞻预测优先使用真实历史数据、负荷预测和可获取天气信息；"
            "未来真实值在预测时刻不可得时，会使用历史同小时或滞后代理特征。"
        ),
        "risk_level": risk_level,
        "key_hours": key_hours,
        "summary_signature": summary["summary_signature"],
    }


def write_word_report(report: dict, output_path: Path) -> None:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style.font.size = Pt(10.5)

    title = document.add_heading("电价智能分析综合报告", level=0)
    title.runs[0].font.name = "Microsoft YaHei"
    title.runs[0].font.size = Pt(18)

    meta = document.add_paragraph()
    meta.add_run(f"生成时间：{now_text()}\n")
    meta.add_run(f"风险等级：{report['risk_level']}\n")
    meta.add_run(f"重点时段：{', '.join(report['key_hours']) if report['key_hours'] else '无'}\n")
    meta.add_run(f"摘要签名：{report.get('summary_signature', '')}\n")
    meta.add_run(f"报告生成方式：{report['generator']['mode']}")

    sections = [
        ("执行摘要", report["executive_summary"]),
        ("市场概览", report["market_overview"]),
        ("未来24小时趋势", report["next_24h_trend"]),
        ("高峰风险", report["peak_risk"]),
        ("运营建议", report["operation_advice"]),
        ("管理层摘要", report["management_summary"]),
        ("告警摘要", report["alert_message"]),
        ("局限说明", report["limitations"]),
    ]

    for title_text, content in sections:
        document.add_heading(title_text, level=1)
        if isinstance(content, list):
            for item in content:
                document.add_paragraph(str(item), style="List Bullet")
        else:
            document.add_paragraph(str(content))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def write_report_files(report: dict, current_dir: Path) -> Path:
    write_json(current_dir / "ai_report_structured.json", report)
    for legacy_name in ["daily_report.md", "daily_report.txt", "management_summary.txt", "alert_message.txt"]:
        legacy_path = current_dir / legacy_name
        if legacy_path.exists():
            legacy_path.unlink()
    word_path = current_dir / "电价智能分析综合报告.docx"
    write_word_report(report, word_path)
    return word_path


def write_llm_attempt_meta(current_dir: Path, client: LLMClient, success: bool, error: str = "") -> None:
    write_json(
        current_dir / "llm_report_attempt_meta.json",
        {
            "success": success,
            "provider": client.provider,
            "model": client.model,
            "retry_count": client.last_retry_count,
            "raw_response": client.last_raw_response,
            "parse_error": client.last_parse_error,
            "schema_error": client.last_schema_error,
            "error": error,
            "recorded_at": now_text(),
        },
    )


def is_valid_llm_payload(report: dict) -> bool:
    if not isinstance(report, dict):
        return False
    missing = LLM_REQUIRED_KEYS - set(report.keys())
    if missing:
        return False
    if not isinstance(report.get("operation_advice"), list):
        return False
    return True


def is_domain_relevant_text(text: str) -> bool:
    text = (text or "").strip()
    if not text:
        return False
    if any(keyword in text for keyword in IRRELEVANT_KEYWORDS):
        return False
    return any(keyword in text for keyword in DOMAIN_KEYWORDS)


def _combined_text_from_report(report: dict) -> str:
    return "\n".join(
        [
            str(report.get("executive_summary", "")),
            str(report.get("market_overview", "")),
            str(report.get("next_24h_trend", "")),
            str(report.get("peak_risk", "")),
            str(report.get("management_summary", "")),
            str(report.get("alert_message", "")),
            " ".join(report.get("operation_advice", [])),
        ]
    )


def contains_key_facts(summary: dict, report: dict) -> bool:
    combined_text = _combined_text_from_report(report)
    forecast = summary["forecast_summary"]
    max_hour = str(forecast.get("next_24h_max_hour", ""))
    min_hour = str(forecast.get("next_24h_min_hour", ""))
    top_risk_hours = [str(item.get("hour", "")) for item in forecast.get("top_risk_hours", [])[:3] if str(item.get("hour", ""))]

    if max_hour and max_hour not in combined_text:
        return False
    if min_hour and min_hour not in combined_text:
        return False
    if top_risk_hours and not any(hour in combined_text for hour in top_risk_hours):
        return False
    return True


def is_valid_section(section_name: str, value, summary: dict) -> bool:
    if section_name == "operation_advice":
        if not isinstance(value, list) or not (3 <= len(value) <= 5):
            return False
        return all(is_domain_relevant_text(str(item)) for item in value)

    text = str(value or "").strip()
    if not text or not is_domain_relevant_text(text):
        return False

    forecast = summary["forecast_summary"]
    max_hour = str(forecast.get("next_24h_max_hour", ""))
    min_hour = str(forecast.get("next_24h_min_hour", ""))
    top_risk_hours = [str(item.get("hour", "")) for item in forecast.get("top_risk_hours", [])[:3] if str(item.get("hour", ""))]

    if section_name in {"executive_summary", "next_24h_trend"}:
        return max_hour in text and min_hour in text
    if section_name in {"peak_risk", "alert_message"} and top_risk_hours:
        return any(hour in text for hour in top_risk_hours)
    return True


def merge_llm_with_fallback(summary: dict, llm_payload: dict, fallback_report: dict) -> tuple[dict, list[str], list[str]]:
    merged = fallback_report.copy()
    used_sections: list[str] = []
    fallback_sections: list[str] = []

    for key in LLM_REQUIRED_KEYS:
        if key in llm_payload and is_valid_section(key, llm_payload[key], summary):
            merged[key] = llm_payload[key]
            used_sections.append(key)
        else:
            fallback_sections.append(key)
    return merged, used_sections, fallback_sections


def main() -> None:
    start = time.perf_counter()
    config = load_config()
    paths = get_pipeline_paths(config)
    run_context = get_run_context()
    log, _ = setup_run_logger(paths.log_dir, "03_llm_generate_report")

    summary = read_json(paths.current_dir / "ai_input_summary.json")
    client = LLMClient(config)
    system_prompt, user_prompt = build_prompts(summary)
    report = build_fallback_report(summary)

    try:
        if os.environ.get("LLM_DISABLE", "0") == "1":
            raise RuntimeError("LLM_DISABLE=1，跳过大模型调用并使用模板回退。")
        log(f"开始调用本地模型生成报告：{client.model}")
        llm_payload = client.generate_json(system_prompt, user_prompt, schema=LLM_REPORT_SCHEMA)
        write_llm_attempt_meta(paths.current_dir, client, success=True)
        write_json(paths.current_dir / "llm_report_raw.json", llm_payload)
        if not is_valid_llm_payload(llm_payload):
            raise ValueError(f"LLM 返回字段不完整：{sorted(set(llm_payload.keys()))}")
        merged_report, used_sections, fallback_sections = merge_llm_with_fallback(summary, llm_payload, report)
        if fallback_sections:
            retry_system_prompt, retry_user_prompt = build_regeneration_prompts(summary, fallback_sections)
            retry_schema = build_subset_schema(fallback_sections)
            retry_payload = client.generate_json(retry_system_prompt, retry_user_prompt, schema=retry_schema)
            write_llm_attempt_meta(paths.current_dir, client, success=True)
            write_json(paths.current_dir / "llm_report_retry_raw.json", retry_payload)
            for key in fallback_sections[:]:
                if key in retry_payload and is_valid_section(key, retry_payload[key], summary):
                    merged_report[key] = retry_payload[key]
                    used_sections.append(key)
                    fallback_sections.remove(key)
        if not used_sections:
            raise ValueError("LLM 各分段均未通过校验，已回退到动态模板。")
        if not contains_key_facts(summary, merged_report):
            raise ValueError("LLM 合并结果未体现本次关键时段事实，已回退到动态模板。")

        report = merged_report
        mode = "llm_dynamic" if not fallback_sections else "llm_hybrid_dynamic"
        report["generator"] = {
            "provider": client.provider,
            "model": client.model,
            "mode": mode,
            "llm_sections": used_sections,
            "fallback_sections": fallback_sections,
        }
    except Exception as exc:
        log(f"LLM 动态生成失败，改用数据驱动模板回退：{exc}")
        write_llm_attempt_meta(paths.current_dir, client, success=False, error=str(exc))
        report["generator"] = {"provider": "fallback", "model": "template", "mode": "fallback_dynamic", "llm_failed": True}

    report["summary_signature"] = summary["summary_signature"]
    word_path = write_report_files(report, paths.current_dir)
    save_ai_report(config, report, word_path, run_context=run_context, log=log)

    elapsed = format_duration(time.perf_counter() - start)
    log(f"已生成 Word 综合报告与结构化报告，耗时：{elapsed}。")
    save_pipeline_event(config, "ai_report_generation", "completed", f"AI 报告已生成，耗时：{elapsed}", run_context)


if __name__ == "__main__":
    main()
